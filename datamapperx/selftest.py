"""``DataMapperX.exe --smoke-test``: prove a freshly built exe actually works.

PyInstaller can produce an exe that starts but is missing a bundled file (a theme, a codec, a module that is only
imported when a rarely used feature runs). This exercises everything the app needs, including the features that
load code late, and reports through the exit code because a windowed exe has no console. The release workflow
runs it before publishing anything.
"""

from __future__ import annotations

import logging
import shutil
import tempfile
import time
from pathlib import Path

logger = logging.getLogger(__name__)


class SmokeTestError(Exception):
    """Something the application needs is missing or gives a wrong answer."""


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise SmokeTestError(message)


def run() -> int:
    try:
        with tempfile.TemporaryDirectory(prefix="dmx_smoke_") as tmp:
            sample = _check_logic(Path(tmp))
            _check_window(sample)
    except Exception:
        logger.exception("Smoke test failed")
        return 1
    logger.info("Smoke test passed")
    return 0


def _check_logic(folder: Path) -> Path:
    """Write, scan (fast reader and the openpyxl fallback), detect a profile and export. Returns a sample workbook."""
    from openpyxl import Workbook, load_workbook

    from .exporter import generate_excel
    from .profiles import DEFAULT_PROFILES
    from .sampling import detect_excel
    from .scanner import scan_folder
    from .xlsx import OpenpyxlReader, XlsxReader

    for run_no in (1, 2):
        wb = Workbook()
        wb.active.title = "Session"
        wb.active["B2"] = "VSE000SMOKE"
        results = wb.create_sheet("Results")
        results.append(["#", "Id", "Test", "Type", "Min", "Measured", "Max", "Unit"])
        results.append([1, 1, "R1", "MinMax", 0, 1.5 * run_no, 3, "V"])
        wb.save(folder / f"VSE000SMOKE_{run_no}.xlsx")
    sample = folder / "VSE000SMOKE_1.xlsx"

    result = scan_folder(folder, dict(DEFAULT_PROFILES["Function Test Log"]))
    _require(result.total_sns == 1 and result.total_files == 2, "scan found the wrong files")

    with XlsxReader(sample) as fast, OpenpyxlReader(sample) as slow:  # the fallback is rarely used, so test it here
        _require(fast.sheet_names == slow.sheet_names == ["Session", "Results"], "the two readers disagree on sheets")
        _require(fast.grid("Results", 5, 8)[1][:6] == slow.grid("Results", 5, 8)[1][:6], "the two readers disagree on cells")

    detected = detect_excel(sample).profile
    _require(detected["data_sheet"] == "Results" and detected["measure_col"] == "F", "profile detection failed")

    output = folder / "out.xlsx"
    generate_excel(result.grouped, ["R1"], output, expected_repeated=2)
    sheet = load_workbook(output)["Raw Data"]
    _require(sheet["C2"].value == 1.5 and sheet["C3"].value == 3, "export wrote the wrong values")
    _check_statistical_reports(folder / "statistical-reports")
    return sample


def _check_statistical_reports(folder: Path) -> None:
    """Exercise late-loaded PDF/Excel dependencies in the built executable too."""
    from openpyxl import load_workbook

    from .stat_reports import ALL_SECTIONS, ReportContext, export_statistical_report
    from .statistics import AnalysisConfig, Observation, analyze

    folder.mkdir(exist_ok=True)
    studies = (
        ("capability", [Observation("SMOKE_CAPABILITY", str(index), value, 0, 6)
                        for index, value in enumerate((1, 2, 3, 4, 5))], AnalysisConfig(), "Cpk", 2, 1.128, 1),
        ("grr", [Observation("SMOKE_GRR", str(part), value, 0, 20, trial=str(trial))
                 for part, values in enumerate(((1, 2, 3), (4, 5, 6), (7, 8, 9)))
                 for trial, value in enumerate(values)], AnalysisConfig(method="grr"), "Total Gage R&R", 5, 30.0, 4),
    )
    for method, rows, config, label, column, expected, chart_count in studies:
        results = analyze(rows, config)
        _require(len(results) == 1 and results[0].valid, f"{method} smoke dataset was not analyzable")
        pdf_path = export_statistical_report(folder / f"{method}.pdf", results)
        payload = pdf_path.read_bytes()
        _require(payload.startswith(b"%PDF-") and b"%%EOF" in payload[-1024:],
                 f"{method} PDF report was not generated correctly")
        workbook_path = export_statistical_report(folder / f"{method}.xlsx", results,
                                                  context=ReportContext(included=1, total=1, sections=ALL_SECTIONS))
        workbook = load_workbook(workbook_path)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            found = [row for row in sheet.iter_rows(max_col=column) if str(row[0].value).strip() == label]
            actual = found[-1][column - 1].value
            _require(isinstance(actual, (int, float)) and abs(actual - expected) < 1e-9,
                     f"{method} report wrote the wrong metric: {actual!r}")
            _require(len(sheet._charts) == chart_count, f"{method} Excel report is missing study charts")
            _require(workbook["Observations"].max_row == len(rows) + 1,
                     f"{method} report did not retain its observations")
        finally:
            workbook.close()
    logger.info("Statistical report self-test passed (capability and GRR PDF/Excel)")
    _check_customer_criteria()


def _check_customer_criteria() -> None:
    """Per-customer Gage R&R acceptance: the same study gives each customer's own verdict."""
    from dataclasses import replace

    from .statistics import BUILTIN_CUSTOMERS, AnalysisConfig, Observation, analyze

    rows = [Observation("SMOKE_CUSTOMER", str(part), value, 0, 20, trial=str(trial))
            for part, values in enumerate(((1, 2, 3), (4, 5, 6), (7, 8, 9)))
            for trial, value in enumerate(values)]  # %SV/Tol 30, %Study Var 32.2, NDC 4
    expected = {"SolarEdge": "CONDITIONAL PASS", "Zebra": "CONDITIONAL PASS", "SES": "FAIL", "Jabil": "FAIL"}
    for name, status in expected.items():
        result, = analyze(rows, AnalysisConfig(method="grr", profile=name, criteria=BUILTIN_CUSTOMERS[name]))
        _require(result.status == status, f"customer {name} gave {result.status}, expected {status}")
    # A customer that applies only NDC and %SV/Tol, with its own limits, can accept the same study.
    relaxed = tuple(replace(c, enabled=c.metric != "percent_study_var",
                            pass_limit=40 if c.metric == "percent_tolerance" else 4,
                            fail_limit=50 if c.metric == "percent_tolerance" else 4) for c in BUILTIN_CUSTOMERS["Jabil"])
    result, = analyze(rows, AnalysisConfig(method="grr", profile="Smoke", criteria=relaxed))
    _require(result.status == "PASS" and result.details["acceptance_conflicts"] == ("percent_study_var",),
             f"custom customer limits gave {result.status}")
    logger.info("Customer acceptance self-test passed")


def _check_window(sample: Path) -> None:
    from .ui.main_window import MainWindow
    from .ui.profile_dialog import ProfileManager
    from .ui.sample_wizard import SampleWizard

    app = MainWindow(background_checks=False)  # the build's check must not depend on GitHub, or on the owner's off switch
    try:
        app.update()
        _require(bool(app.winfo_exists()), "the main window did not open")
        _check_steps(app, sample.parent)
        _check_customer_dialog(app)
        manager = ProfileManager(app, app.profiles, lambda _name: None)
        wizard = SampleWizard(app, lambda _name, _profile: None, sample=sample)
        for _ in range(5):
            app.update()
        _require(wizard.kind == "excel", "the sample wizard did not read the sample")
        wizard.destroy()
        manager.destroy()
    finally:
        app.destroy()


def _check_customer_dialog(app) -> None:
    """Open the customer criteria dialog, add a customer, apply a metric choice and save; then put everything back."""
    from .ui.validation import CustomersDialog

    app._show_workspace("validation")
    panel = app.validation
    original_customers, original_choice = dict(panel.customers), panel.profile_var.get()
    try:
        dialog = CustomersDialog(panel)
        app.update()
        dialog.new_var.set("Smoke customer")
        dialog.add_customer()
        for metric, applied in (("percent_tolerance", True), ("percent_study_var", False), ("ndc", True)):
            dialog.apply_vars[metric].set(applied)
        dialog.pass_vars["percent_tolerance"].set("15")
        dialog.fail_vars["percent_tolerance"].set("25")
        app.update()
        _require("PASS < 15%" in dialog.bands["percent_tolerance"].get(), "the dialog did not show the PASS/FAIL bands")
        dialog.save()
        app.update()
        saved = panel.customers.get("Smoke customer")
        _require(saved is not None and [c.enabled for c in saved] == [True, False, True] and saved[0].pass_limit == 15,
                 "the customer criteria were not saved")
        _require(panel.profile_var.get() == "Smoke customer" and "Smoke customer" in app.settings.customer_criteria,
                 "the customer was not selected and remembered")
        _require(panel._evaluation_for("grr").criteria == saved, "the analysis does not use the customer's criteria")
    finally:
        panel.set_customers(original_customers, selected=original_choice)
        app._show_workspace("mapping")


def _wait(app, condition, what: str, seconds: float = 30.0) -> None:
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        app.update()
        if condition():
            return
        time.sleep(0.01)
    raise SmokeTestError(f"timed out waiting for {what}")


def _check_steps(app, folder: Path) -> None:
    """Scan the sample files, pick the component and walk through both steps of the window."""
    from tkinter import messagebox

    def refuse(*args, **kwargs):  # a dialog nobody can answer would hang the run
        raise SmokeTestError(f"the window showed a dialog: {args[1:2]}")

    logs = folder / "steps"  # only the two sample logs: the export written by the logic check must not be scanned
    logs.mkdir(exist_ok=True)
    for source in folder.glob("VSE000SMOKE_*.xlsx"):
        shutil.copy(source, logs / source.name)
    saved = {name: getattr(messagebox, name) for name in ("showinfo", "showwarning", "showerror", "askyesno")}
    for name in saved:
        setattr(messagebox, name, refuse)
    try:
        app.profile_var.set("Function Test Log")
        app.folder_var.set(str(logs))
        app.expected_var.set("1")
        app.repeats_var.set("2")
        app.scan_logs()
        _wait(app, lambda: not app.is_busy and app.scan_result is not None, "the scan")
        _require(app.scan_result.total_sns == 1, "the scan found the wrong files")
        app.clear_selected()  # the data folder may remember the choice from an earlier run
        _require("R1" in app.available_list.items(), "the window did not list the components")
        _require(app.available_list._values[0] != ("–", "–"), "the component list did not show the limits")
        app.add_all_shown()
        _require(app.selected_components == ["R1"] and app.review_button.instate(["!disabled"]), "the component was not selected")
        app._show_step("review")
        app.update()
        _require(app.step == "review" and len(app.preview_grid.rows()) == 2, "the review did not show the data")
        _require(len(app.limits_grid.rows()) == 1 and app.check_chips.chips(), "the review did not show the limits")
        app._show_workspace("validation")  # Validation & Reports analyses a copy of what was scanned
        workspace = app.validation
        _wait(app, lambda: workspace.dataset is not None and not workspace.busy, "the validation data")
        workspace.run_button.invoke()
        _wait(app, lambda: workspace.evaluated is not None and not workspace.running, "the validation analysis")
        _require(len(workspace.results) == 1 and workspace.results[0].test_item == "R1", "validation did not analyse the scan")
        app._show_workspace("mapping")
        for _ in range(2):  # dark and back again, so the setting it saves is the one it started with
            app.dark_var.set(not app.dark_var.get())
            app._on_theme_toggle()
        app._show_step("select")
        app.update()
        _require(app.step == "select", "could not go back to the first step")
        app.reset_workflow()
    finally:
        for name, original in saved.items():
            setattr(messagebox, name, original)
