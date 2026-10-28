import json
import logging
import os
import platform
import re
import shutil
import subprocess
import sys
import threading
from collections import Counter, deque
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
from zipfile import BadZipFile

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


# =============================================================================
# USER-EDITABLE CONFIGURATION
# =============================================================================
APP_TITLE = "DataMapperX"
APP_VERSION = "v1.1.0"
APP_BUILD = "Build 2026.05.19"
APP_DEVELOPER_LINE = "Developed by MAN_TRAN8118@jabil.com"
HEADER_LOGO_SUBSAMPLE = 2

# Keep blank by default. Change this to "N/A" or "-" if you want missing values
# to be visible in the exported Excel file.
MISSING_TEXT = ""

# Default file types scanned by each profile type.
TEXT_FILE_EXTENSIONS = (".log", ".txt", ".csv")
EXCEL_FILE_EXTENSIONS = (".xlsx", ".xlsm")

DEFAULT_OUTPUT_FILE = "RawData.xlsx"
BUILT_IN_PROFILE_NAMES = ("ICT_I3070 Test Log", "Function Test Log")

# All major UI labels/messages are centralized here so you can rename the app,
# translate it, or adjust instructions without hunting through the code.
UI_TEXT = {
    "workspace_setup": "⚙️ Workspace Setup",
    "profile": "📋 Profile",
    "profile_file_hint": "Input: {input_type} | Accepted files: {extensions}",
    "profile_type_text": "Text log",
    "profile_type_excel": "Excel log",
    "manage_profiles": "⚙️ Manage Profiles",
    "source_folder": "📁 Source Folder",
    "browse": "🔍 Browse",
    "scan_logs": "▶ Scan Logs",
    "expected_samples": "🎯 Expected Samples",
    "repeats_per_sample": "🔄 Repeats per Sample",
    "detected_count": "Total: {count}",
    "selected_count": "Selected: {count}",

    "available_components": "📋 Available Components",
    "selected_components": "✅ Selected Components",
    "add": "➕ Add",
    "remove": "➖ Remove",
    "move_up": "⬆️ Move Up",
    "move_down": "⬇️ Move Down",
    "clear_all": "🗑️ Clear All",

    "preview_export": "👁️ Preview and Export",
    "preview": "👁️ Preview",
    "export_excel": "💾 Export Excel",
    "tree_no": "No.",
    "tree_sn": "Serial Number",
    "tree_count": "Files",
    "tree_files": "Sorted File Names",

    "status_ready": "Ready.",
    "status_scanning": "Scanning log files...",
    "status_scan_failed": "Scan failed.",
    "status_preview_loaded": "Preview loaded.",
    "status_generating": "Generating Excel...",
    "status_export_failed": "Export failed.",
    "status_done": "Done. Exported: {output}",

    "manage_title": "Manage Profiles",
    "profiles": "Profiles",
    "edit_profile": "Edit Profile",
    "new": "New",
    "delete": "Delete",
    "save_profile": "Save Profile",
    "profile_name": "Profile Name:",
    "input_type": "Input Type:",
    "sn_regex_cell": "SN_Regex / Cell:",
    "sn_tip": "(For Excel you can enter a cell like B2, or leave blank to search filename/cells)",
    "regex_help": "📖 View Regex Guide & Rules",
    "name_regex": "Name_Regex:",
    "measure_regex": "Measure_Regex:",
    "minmax_regex1": "Min_Max_Regex1:",
    "minmax_regex2": "Min_Max_Regex2:",
    "regex_tester": "Live Data Regex Tester",
    "paste_data_row": "Paste Data Row:",
    "result_empty": "Result: No data",
    "result_unrecognized": "Result: Format not recognized with current Regexes",
    "excel_sn_sheet": "SN_Sheet:",
    "excel_data_sheet": "Data_Sheet:",
    "excel_row_start": "Row_Start:",
    "excel_component_col": "Component_Col:",
    "excel_measure_col": "Measure_Col:",
    "excel_min_col": "Min_Col:",
    "excel_max_col": "Max_Col:",
    "excel_unit_col": "Unit_Col (Leave blank if none):",

    "msg_choose_folder": "Please choose a folder first.",
    "msg_invalid_folder": "Folder path is not valid.",
    "msg_invalid_profile": "Please select a valid profile.",
    "msg_invalid_folder_detail": "Folder path is not valid:\n{folder}",
    "msg_invalid_profile_detail": "Selected profile was not found:\n{profile}",
    "msg_scan_no_data": "No valid log data found.",
    "msg_scan_success": "Scan Successful!\n\nSuccessfully scanned {samples} samples across {files} total files.",
    "msg_expected_samples": "Expected {expected} samples, but found {actual} samples.",
    "msg_invalid_repeats": "Some SN do not have exactly {expected} files:",
    "msg_parse_warnings": "Some files could not be read and were skipped:",
    "msg_scan_first": "Please scan a folder first.",
    "msg_select_component": "Please select at least one component.",
    "msg_export_success": "Excel exported successfully:\n{output}",
    "msg_export_open_folder": "Excel exported successfully.\n\nOpen the output folder now?",
    "msg_export_mismatch_confirm": (
        "Expected {expected} samples, but found {actual} samples.\n\n"
        "Do you still want to export?"
    ),
}


DEFAULT_PROFILES = {
    "ICT_I3070 Test Log": {
        "type": "text",
        "file_extensions": ".log,.txt,.csv",
        "sn_regex": r"(VSE[A-Z0-9]+)",
        "block_regex": r"\{@BLOCK\|(?P<component>[^|]+)\|00\b",
        "measure_regex": r"\{@A-(?P<type>[A-Z0-9]+)\|[^|]*\|(?P<measured>[+-]?\d+(?:\.\d+)?E[+-]?\d+)(?P<tail>.*)",
        "lim2_regex": r"\{@LIM2\|(?P<max>[+-]?\d+(?:\.\d+)?E[+-]?\d+)\|(?P<min>[+-]?\d+(?:\.\d+)?E[+-]?\d+)\}\}",
        "lim3_regex": r"\{@LIM3\|(?P<normal>[+-]?\d+(?:\.\d+)?E[+-]?\d+)\|(?P<max>[+-]?\d+(?:\.\d+)?E[+-]?\d+)\|(?P<min>[+-]?\d+(?:\.\d+)?E[+-]?\d+)\}\}",
    },
    "Function Test Log": {
        "type": "excel",
        "file_extensions": ".xlsx,.xlsm",
        "sn_regex": r"(VSE[A-Z0-9]+)",
        "sn_sheet": "Session",
        "data_sheet": "Results",
        "row_start": "2",
        "component_col": "C",
        "measure_col": "F",
        "min_col": "E",
        "max_col": "G",
        "unit_col": "H",
    },
}


# =============================================================================
# APP STORAGE / LOGGING
# =============================================================================
def is_frozen_app() -> bool:
    return bool(getattr(sys, "frozen", False))


def app_dir() -> Path:
    if is_frozen_app():
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def resource_dir() -> Path:
    """Read-only resource root. In PyInstaller one-file mode this is sys._MEIPASS."""
    if is_frozen_app() and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return app_dir()


def resource_path(*parts: str) -> Path:
    return resource_dir().joinpath(*parts)


def user_data_dir() -> Path:
    if not is_frozen_app():
        return app_dir()

    base = os.getenv("APPDATA") or os.getenv("LOCALAPPDATA")
    if base:
        return Path(base) / APP_TITLE
    return app_dir()


def ensure_user_data_dir() -> Path:
    path = user_data_dir()
    try:
        path.mkdir(parents=True, exist_ok=True)
    except OSError:
        logging.getLogger(__name__).warning("Could not create user data directory: %s", path, exc_info=True)
        return app_dir()
    return path


def profile_path() -> Path:
    return ensure_user_data_dir() / "profiles.json"


def seed_profiles_if_needed(path: Path) -> None:
    if path.exists():
        return

    for candidate in (
        app_dir() / "profiles.json",
        resource_path("profiles.json"),
    ):
        if candidate.exists() and candidate.resolve() != path.resolve():
            try:
                shutil.copy2(candidate, path)
                return
            except OSError:
                logger.warning("Could not seed profiles from %s", candidate, exc_info=True)

    with open(path, "w", encoding="utf-8") as f:
        json.dump(DEFAULT_PROFILES, f, indent=4, ensure_ascii=False)


def setup_logging() -> None:
    root = logging.getLogger()
    if root.handlers:
        return

    handlers: List[logging.Handler] = [logging.StreamHandler()]
    try:
        log_dir = user_data_dir()
        log_dir.mkdir(parents=True, exist_ok=True)
        handlers.insert(0, logging.FileHandler(log_dir / "app_debug.log", encoding="utf-8"))
    except OSError:
        pass

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
        handlers=handlers,
    )


setup_logging()
logger = logging.getLogger(__name__)


def t(key: str, **kwargs) -> str:
    text = UI_TEXT.get(key, key)
    return text.format(**kwargs) if kwargs else text


def configure_process_dpi_awareness() -> None:
    if platform.system() != "Windows":
        return

    try:
        import ctypes

        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        logging.getLogger(__name__).debug("Could not set DPI awareness", exc_info=True)


def set_window_icon(window: tk.Tk | tk.Toplevel) -> None:
    icon = resource_path("assets", "DataMapperX.ico")
    if icon.exists():
        try:
            window.iconbitmap(default=str(icon))
        except tk.TclError:
            logger.debug("Could not apply icon: %s", icon, exc_info=True)


# =============================================================================
# PROFILE LOADING / SAVING
# =============================================================================
def load_profiles() -> Dict[str, dict]:
    path = profile_path()
    seed_profiles_if_needed(path)

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)

        migrated = False
        for old_name, new_name in [
            ("HP3070 (Built-in)", "ICT_I3070 Test Log"),
            ("FT build", "Function Test Log"),
            ("ICT_3070", "ICT_I3070 Test Log"),
        ]:
            if old_name in data:
                data[new_name] = data.pop(old_name)
                migrated = True

        # Always refresh built-in profiles so the app remains predictable.
        for profile_name, profile_data in DEFAULT_PROFILES.items():
            if data.get(profile_name) != profile_data:
                data[profile_name] = profile_data
                migrated = True

        if migrated:
            save_profiles(data)

        return data
    except Exception as exc:
        logger.exception("Could not load profiles.json: %s", exc)
        return json.loads(json.dumps(DEFAULT_PROFILES))


def save_profiles(profiles_data: Dict[str, dict]) -> None:
    path = profile_path()
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profiles_data, f, indent=4, ensure_ascii=False)


# =============================================================================
# PARSING HELPERS
# =============================================================================
SN_PATTERN = re.compile(r"(VSE[A-Z0-9]+)", re.IGNORECASE)


@dataclass
class Measurement:
    component: str
    measured: str
    min_value: str
    max_value: str
    unit: str
    source_row: str = ""


def safe_text(value) -> str:
    return "" if value is None else str(value).strip()


def normalize_component_name(name: str) -> str:
    return safe_text(name)


def measurement_choice_label(measurement: Measurement) -> str:
    details = [f"row {measurement.source_row}" if measurement.source_row else "first match"]
    if measurement.measured:
        details.append(f"value={measurement.measured}")
    if measurement.min_value or measurement.max_value:
        details.append(f"min={measurement.min_value or MISSING_TEXT}, max={measurement.max_value or MISSING_TEXT}")
    if measurement.unit:
        details.append(f"unit={measurement.unit}")
    return " | ".join(details)


def parse_positive_int(value, label: str, min_value: int = 1) -> int:
    try:
        parsed = int(str(value).strip())
    except (TypeError, ValueError):
        raise ValueError(f"{label} must be a number.") from None

    if parsed < min_value:
        raise ValueError(f"{label} must be at least {min_value}.")

    return parsed


def looks_like_cell_coordinate(value: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z]{1,3}\d{1,7}", safe_text(value)))


def normalize_excel_column(value: str, label: str, allow_blank: bool = False) -> str:
    col = safe_text(value).upper()
    if not col and allow_blank:
        return ""
    if not re.fullmatch(r"[A-Z]{1,3}", col):
        raise ValueError(f"{label} must be an Excel column letter, for example A, B, AA.")
    return col


def compile_profile_regex(pattern: str, label: str, flags: int = 0) -> Optional[re.Pattern]:
    pattern = safe_text(pattern)
    if not pattern:
        return None
    try:
        return re.compile(pattern, flags)
    except re.error as exc:
        raise ValueError(f"Invalid Regex in {label}: {exc}") from exc


def sn_from_match(match: re.Match) -> str:
    if "sn" in match.groupdict():
        value = match.group("sn")
    elif match.groups():
        value = match.group(1)
    else:
        value = match.group(0)
    return safe_text(value).upper()


def extract_sn_from_text(source: str, sn_regex: str = "") -> str:
    source = safe_text(source)
    if not source:
        return ""

    pattern = compile_profile_regex(sn_regex, "SN_Regex", re.IGNORECASE) if safe_text(sn_regex) else SN_PATTERN
    match = pattern.search(source)
    return sn_from_match(match) if match else ""


def iter_excel_cell_text(ws, max_rows: int = 80, max_cols: int = 30) -> Iterable[str]:
    row_limit = min(ws.max_row or 1, max_rows)
    col_limit = min(ws.max_column or 1, max_cols)
    for row in ws.iter_rows(max_row=row_limit, max_col=col_limit, values_only=True):
        for value in row:
            text = safe_text(value)
            if text:
                yield text


def get_profile_extensions(profile: dict) -> Tuple[str, ...]:
    raw = profile.get("file_extensions")
    if isinstance(raw, str) and raw.strip():
        parts = re.split(r"[,;\s]+", raw.strip())
    elif isinstance(raw, (list, tuple)):
        parts = list(raw)
    else:
        profile_type = profile.get("type", "text")
        parts = EXCEL_FILE_EXTENSIONS if profile_type == "excel" else TEXT_FILE_EXTENSIONS

    cleaned = []
    for item in parts:
        item = safe_text(item).lower()
        if not item:
            continue
        cleaned.append(item if item.startswith(".") else f".{item}")

    return tuple(dict.fromkeys(cleaned)) or TEXT_FILE_EXTENSIONS


# Compile text regexes once so validation and parsing share the same setup.
def compile_text_profile_regexes(profile: dict) -> dict:
    """Compile and return all regexes for a text profile. Raises ValueError on bad patterns."""
    return {
        "sn_re":      compile_profile_regex(profile.get("sn_regex", ""),    "SN_Regex",      re.IGNORECASE),
        "block_re":   compile_profile_regex(profile.get("block_regex", ""), "Name_Regex",    re.IGNORECASE),
        "measure_re": compile_profile_regex(profile.get("measure_regex", ""), "Measure_Regex", re.IGNORECASE),
        "lim2_re":    compile_profile_regex(profile.get("lim2_regex", ""),  "Min_Max_Regex1"),
        "lim3_re":    compile_profile_regex(profile.get("lim3_regex", ""),  "Min_Max_Regex2"),
    }


def validate_profile(profile: dict) -> None:
    profile_type = profile.get("type", "text")
    if profile_type not in ("text", "excel"):
        raise ValueError("Profile type must be either 'text' or 'excel'.")

    if profile_type == "text":
        if not safe_text(profile.get("block_regex")):
            raise ValueError("Text profile is missing Name_Regex.")
        if not safe_text(profile.get("measure_regex")):
            raise ValueError("Text profile is missing Measure_Regex.")
        # Compile once to validate — parsers will compile again via compile_text_profile_regexes()
        compile_text_profile_regexes(profile)

    else:
        parse_positive_int(profile.get("row_start", 2), "Row_Start")
        normalize_excel_column(profile.get("component_col", "A"), "Component_Col")
        normalize_excel_column(profile.get("measure_col", "B"), "Measure_Col")
        normalize_excel_column(profile.get("min_col", "C"), "Min_Col", allow_blank=True)
        normalize_excel_column(profile.get("max_col", "D"), "Max_Col", allow_blank=True)
        normalize_excel_column(profile.get("unit_col", "E"), "Unit_Col", allow_blank=True)

        sn_rule = safe_text(profile.get("sn_regex", ""))
        if sn_rule and not looks_like_cell_coordinate(sn_rule):
            compile_profile_regex(sn_rule, "SN_Regex", re.IGNORECASE)


def parse_excel_file(path: Path, profile: dict) -> Tuple[str, Dict[str, Measurement], List[dict]]:
    try:
        wb = load_workbook(str(path), data_only=True)
    except PermissionError:
        raise PermissionError(
            f"Cannot access '{path.name}'.\n\n"
            "Microsoft Excel locks files when they are open. "
            "Please CLOSE the file in Excel and try scanning again."
        ) from None
    except BadZipFile:
        raise ValueError(f"'{path.name}' is not a valid Excel workbook.") from None

    entries: Dict[str, Measurement] = {}
    duplicate_groups_by_component: Dict[str, dict] = {}
    sn = ""

    sn_sheet_name = safe_text(profile.get("sn_sheet", ""))
    data_sheet_name = safe_text(profile.get("data_sheet", ""))

    ws_sn = wb[sn_sheet_name] if sn_sheet_name in wb.sheetnames else wb.worksheets[0]
    if data_sheet_name in wb.sheetnames:
        ws_data = wb[data_sheet_name]
    elif len(wb.worksheets) > 1:
        ws_data = wb.worksheets[1]
    else:
        ws_data = wb.active

    sn_rule = safe_text(profile.get("sn_regex", ""))
    if looks_like_cell_coordinate(sn_rule):
        cell_text = safe_text(ws_sn[sn_rule.upper()].value)
        sn = extract_sn_from_text(cell_text) or cell_text.upper()

    if not sn:
        sn = extract_sn_from_text(path.name, sn_rule if not looks_like_cell_coordinate(sn_rule) else "")

    if not sn:
        for cell_text in iter_excel_cell_text(ws_sn):
            sn = extract_sn_from_text(cell_text, sn_rule if not looks_like_cell_coordinate(sn_rule) else "")
            if sn:
                break

    row_start = parse_positive_int(profile.get("row_start", 2), "Row_Start")
    col_comp = normalize_excel_column(profile.get("component_col", "A"), "Component_Col")
    col_meas = normalize_excel_column(profile.get("measure_col", "B"), "Measure_Col")
    col_min = normalize_excel_column(profile.get("min_col", "C"), "Min_Col", allow_blank=True)
    col_max = normalize_excel_column(profile.get("max_col", "D"), "Max_Col", allow_blank=True)
    col_unit = normalize_excel_column(profile.get("unit_col", "E"), "Unit_Col", allow_blank=True)

    for row in range(row_start, (ws_data.max_row or 0) + 1):
        compn = normalize_component_name(ws_data[f"{col_comp}{row}"].value)
        if not compn:
            continue

        meas = safe_text(ws_data[f"{col_meas}{row}"].value) if col_meas else ""
        mi = safe_text(ws_data[f"{col_min}{row}"].value) if col_min else ""
        ma = safe_text(ws_data[f"{col_max}{row}"].value) if col_max else ""
        unit = safe_text(ws_data[f"{col_unit}{row}"].value) if col_unit else ""

        if compn not in entries:
            entries[compn] = Measurement(
                component=compn,
                measured=meas,
                min_value=mi,
                max_value=ma,
                unit=unit,
                source_row=str(row),
            )
        else:
            duplicate_group = duplicate_groups_by_component.setdefault(
                compn,
                {
                    "component": compn,
                    "row_label": "Rows",
                    "options": [entries[compn]],
                },
            )
            duplicate_group["options"].append(
                Measurement(
                    component=compn,
                    measured=meas,
                    min_value=mi,
                    max_value=ma,
                    unit=unit,
                    source_row=str(row),
                )
            )
            logger.debug("Duplicate component '%s' in %s — keeping first occurrence.", compn, path.name)

    return sn, entries, list(duplicate_groups_by_component.values())


def parse_text_file(path: Path, profile: dict, compiled: Optional[dict] = None) -> Tuple[str, Dict[str, Measurement], List[dict]]:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except PermissionError:
        raise PermissionError(
            f"Cannot access '{path.name}'.\n\n"
            "The file is locked by another program. Please CLOSE it before scanning."
        ) from None

    entries: Dict[str, Measurement] = {}
    duplicate_groups_by_component: Dict[str, dict] = {}

    # Compile on demand only when called standalone.
    if compiled is None:
        compiled = compile_text_profile_regexes(profile)

    sn_re      = compiled["sn_re"]
    block_re   = compiled["block_re"]
    measure_re = compiled["measure_re"]
    lim2_re    = compiled["lim2_re"]
    lim3_re    = compiled["lim3_re"]

    sn_regex = safe_text(profile.get("sn_regex", ""))
    # Extract SN from filename first, then from file content
    sn = ""
    if sn_re:
        m = sn_re.search(path.name)
        sn = sn_from_match(m).upper() if m else ""
    if not sn:
        sn = extract_sn_from_text(path.name, sn_regex)
    if not sn:
        sn = extract_sn_from_text(text, sn_regex)

    if not block_re or not measure_re:
        raise ValueError("Text profile must include Name_Regex and Measure_Regex.")

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        block_match = block_re.search(line)

        if block_match:
            try:
                component = block_match.group("component").strip()
            except IndexError:
                component = block_match.group(1).strip() if block_match.groups() else ""

            component = normalize_component_name(component)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1

            if component and j < len(lines):
                measure_match = measure_re.search(lines[j])
                if measure_match:
                    try:
                        unit = measure_match.group("type").upper()
                    except IndexError:
                        unit = ""

                    try:
                        measured = measure_match.group("measured")
                    except IndexError:
                        measured = measure_match.group(1) if measure_match.groups() else ""

                    try:
                        tail = measure_match.group("tail")
                    except IndexError:
                        tail = ""

                    min_value, max_value = "", ""
                    if tail and lim3_re:
                        lim3_match = lim3_re.search(tail)
                        if lim3_match:
                            try:
                                min_value, max_value = lim3_match.group("min"), lim3_match.group("max")
                            except IndexError:
                                pass

                    if tail and not min_value and lim2_re:
                        lim2_match = lim2_re.search(tail)
                        if lim2_match:
                            try:
                                min_value, max_value = lim2_match.group("min"), lim2_match.group("max")
                            except IndexError:
                                pass

                    if component not in entries:
                        entries[component] = Measurement(
                            component=component,
                            measured=safe_text(measured),
                            min_value=safe_text(min_value),
                            max_value=safe_text(max_value),
                            unit=safe_text(unit),
                            source_row=str(i + 1),
                        )
                    else:
                        duplicate_group = duplicate_groups_by_component.setdefault(
                            component,
                            {
                                "component": component,
                                "row_label": "Lines",
                                "options": [entries[component]],
                            },
                        )
                        duplicate_group["options"].append(
                            Measurement(
                                component=component,
                                measured=safe_text(measured),
                                min_value=safe_text(min_value),
                                max_value=safe_text(max_value),
                                unit=safe_text(unit),
                                source_row=str(i + 1),
                            )
                        )
                        logger.debug("Duplicate component '%s' in %s — keeping first occurrence.", component, path.name)

            i = j

        i += 1

    return sn, entries, list(duplicate_groups_by_component.values())


def parse_log_file(path: Path, profile: dict, compiled: Optional[dict] = None) -> Tuple[str, Dict[str, Measurement], List[dict]]:
    profile_type = profile.get("type", "text")
    if profile_type == "excel":
        return parse_excel_file(path, profile)
    return parse_text_file(path, profile, compiled=compiled)


def scan_folder(
    folder: Path,
    profile: dict,
    cancel_event: Optional[threading.Event] = None,
    include_diagnostics: bool = False,
):
    validate_profile(profile)

    profile_type = profile.get("type", "text")
    extensions = get_profile_extensions(profile)
    logger.info("Scanning folder: %s | profile type: %s | extensions: %s", folder, profile_type, extensions)

    # Text profiles can reuse compiled regexes across every file in this scan.
    compiled_regexes: Optional[dict] = None
    if profile_type == "text":
        compiled_regexes = compile_text_profile_regexes(profile)

    candidate_files = [
        p
        for p in folder.rglob("*")
        if p.is_file()
        and not p.name.startswith("~$")
    ]
    candidate_files = sorted(candidate_files, key=lambda p: str(p).lower())

    files = [
        p
        for p in candidate_files
        if p.suffix.lower() in extensions
    ]
    files = sorted(files, key=lambda p: str(p).lower())

    results = []
    parse_warnings: List[str] = []
    component_counter = Counter()
    duplicate_component_choices: List[dict] = []
    skipped_without_sn: List[str] = []
    files_without_components: List[str] = []
    ignored_extensions = Counter(
        p.suffix.lower() or "[no extension]"
        for p in candidate_files
        if p.suffix.lower() not in extensions
    )

    for log_path in files:
        # Stop promptly when the UI requests cancellation.
        if cancel_event and cancel_event.is_set():
            logger.info("Scan cancelled by user after %d files.", len(results))
            break

        try:
            sn, entries, duplicate_groups = parse_log_file(log_path, profile, compiled=compiled_regexes)
        except Exception as exc:
            logger.warning("Skipped %s: %s", log_path, exc)
            parse_warnings.append(f"{log_path.name}: {exc}")
            continue

        if not sn:
            skipped_without_sn.append(log_path.name)
            continue

        if not entries:
            files_without_components.append(log_path.name)

        item = {
            "path": log_path,
            "filename": log_path.name,
            "sn": sn,
            "entries": entries,
        }

        for component in entries:
            component_counter[component] += 1

        for group in duplicate_groups:
            group["item"] = item
            group["filename"] = log_path.name
            group["sn"] = sn
            duplicate_component_choices.append(group)

        results.append(item)

    grouped = {}
    for item in results:
        grouped.setdefault(item["sn"], []).append(item)

    for sn in grouped:
        grouped[sn] = sorted(grouped[sn], key=lambda x: x["filename"].lower())

    available_components = sorted(component_counter.keys(), key=lambda s: (s.lower(), s))
    diagnostics = {
        "profile_type": profile_type,
        "extensions": extensions,
        "total_file_count": len(candidate_files),
        "accepted_file_count": len(files),
        "ignored_extensions": dict(sorted(ignored_extensions.items())),
        "skipped_without_sn": skipped_without_sn,
        "files_without_components": files_without_components,
    }
    if include_diagnostics:
        return grouped, available_components, parse_warnings, duplicate_component_choices, diagnostics
    return grouped, available_components, parse_warnings, duplicate_component_choices


# =============================================================================
# EXCEL EXPORT
# =============================================================================
def build_summary_info(grouped_logs, selected_components: Sequence[str]):
    summary = {}
    for component in selected_components:
        found = None
        for sn in sorted(grouped_logs.keys()):
            for item in grouped_logs[sn]:
                entry = item["entries"].get(component)
                if entry:
                    found = {
                        "min": entry.min_value or MISSING_TEXT,
                        "max": entry.max_value or MISSING_TEXT,
                        "unit": entry.unit or MISSING_TEXT,
                    }
                    break
            if found:
                break
        summary[component] = found or {"min": MISSING_TEXT, "max": MISSING_TEXT, "unit": MISSING_TEXT}
    return summary


def set_cell_value(cell, val_str: str):
    """Set an Excel cell value while preserving scientific notation formatting."""
    val_str = safe_text(val_str)
    try:
        f_val = float(val_str)
        if "E" in val_str.upper():
            cell.value = f_val
            base = val_str.upper().split("E")[0]
            if "." in base:
                decimals = len(base.split(".")[1])
                cell.number_format = "0." + ("0" * decimals) + "E+00"
            else:
                cell.number_format = "0E+00"
        elif f_val.is_integer():
            cell.value = int(f_val)
        else:
            cell.value = f_val
    except ValueError:
        cell.value = val_str


def write_summary(ws, summary_col: int, selected_components: Sequence[str], summary_info) -> None:
    headers = ["NAME", "MIN", "MAX", "UNIT"]
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")
    right_align = Alignment(horizontal="right")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for offset, text in enumerate(headers):
        cell = ws.cell(1, summary_col + offset)
        cell.value = text
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border

    for idx, component in enumerate(selected_components, start=2):
        info = summary_info.get(component, {})
        data_row = [
            component,
            info.get("min", MISSING_TEXT),
            info.get("max", MISSING_TEXT),
            info.get("unit", MISSING_TEXT),
        ]

        for offset, value in enumerate(data_row):
            cell = ws.cell(idx, summary_col + offset)
            if offset in (1, 2) and value != MISSING_TEXT:
                set_cell_value(cell, value)
            else:
                cell.value = value

            cell.border = thin_border
            if offset > 0:
                cell.alignment = right_align


def validate_grouped_logs_for_export(grouped_logs, expected_repeated: int) -> List[str]:
    return [sn for sn, files in grouped_logs.items() if len(files) != expected_repeated]


def generate_excel_from_grouped(grouped_logs, selected_components: Sequence[str], output_path: Path, expected_repeated: int):
    if not grouped_logs:
        raise ValueError("No scanned data is available. Please scan a folder first.")

    invalid_sns = validate_grouped_logs_for_export(grouped_logs, expected_repeated)
    if invalid_sns:
        raise ValueError(
            f"Each SN must have exactly {expected_repeated} files. "
            "Please update the folder. Problem SN: " + ", ".join(invalid_sns)
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Data"

    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")
    right_align = Alignment(horizontal="right")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws.cell(1, 1, "STT").font = bold_font
    ws.cell(1, 1).alignment = center_align
    ws.cell(1, 2, "SN").font = bold_font
    ws.cell(1, 2).alignment = center_align

    for index, component in enumerate(selected_components, start=3):
        cell = ws.cell(1, index)
        cell.value = component
        cell.font = bold_font
        cell.alignment = center_align

    row = 2
    stt_index = 1
    max_col_data = 2 + len(selected_components)

    for sn in sorted(grouped_logs.keys()):
        for item in grouped_logs[sn]:
            cell_stt = ws.cell(row, 1)
            cell_stt.alignment = center_align
            cell_stt.value = stt_index
            stt_index += 1

            ws.cell(row, 2, sn)

            for comp_index, component in enumerate(selected_components, start=3):
                cell = ws.cell(row, comp_index)
                entry = item["entries"].get(component)
                if entry and entry.measured != MISSING_TEXT:
                    set_cell_value(cell, entry.measured)
                else:
                    cell.value = MISSING_TEXT

            row += 1

    for r in range(1, row):
        for c in range(1, max_col_data + 1):
            cell = ws.cell(r, c)
            cell.border = thin_border
            if r > 1 and c > 2:
                cell.alignment = right_align

    summary_info = build_summary_info(grouped_logs, selected_components)
    summary_start_col = max_col_data + 2
    write_summary(ws, summary_start_col, selected_components, summary_info)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 22
    for col in range(3, max_col_data + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    for col in range(summary_start_col, summary_start_col + 4):
        ws.column_dimensions[get_column_letter(col)].width = 16

    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
    except PermissionError:
        raise PermissionError(
            f"Cannot save '{output_path.name}'.\n\n"
            "The output file may already be open in Excel. Please close it and export again."
        ) from None


# =============================================================================
# TKINTER WIDGETS
# =============================================================================
class UndoableEntry(ttk.Entry):
    """
    ttk.Entry with safe Ctrl+Z / Ctrl+Y.

    Uses collections.deque(maxlen=100) for O(1) history management.
    _set_value() no longer attempts to sync StringVar via low-level Tcl
    calls — delete/insert on the widget is sufficient and correct.
    """
    def __init__(self, master=None, **kwargs):
        super().__init__(master, **kwargs)

        # Keep undo history bounded with O(1) append/drop behavior.
        self._history: deque = deque(maxlen=100)
        self._history_pos = -1
        self._internal_change = False

        self.bind("<FocusIn>", self._sync_history, add="+")
        self.bind("<KeyRelease>", self._record_state, add="+")
        self.bind("<Control-z>", self.undo, add="+")
        self.bind("<Control-Z>", self.undo, add="+")
        self.bind("<Control-y>", self.redo, add="+")
        self.bind("<Control-Y>", self.redo, add="+")
        self.bind("<Control-a>", self.select_all, add="+")
        self.bind("<Control-A>", self.select_all, add="+")
        self.after_idle(self.reset_undo)

    def reset_undo(self):
        value = self.get()
        self._history.clear()
        self._history.append(value)
        self._history_pos = 0

    def _sync_history(self, _event=None):
        value = self.get()

        if self._history_pos < 0:
            self.reset_undo()
            return

        # If the field was changed externally (e.g. profile switch via StringVar),
        # reset history so Ctrl+Z cannot undo back to a previous profile's content.
        history_list = list(self._history)
        if self._history_pos >= len(history_list) or history_list[self._history_pos] != value:
            self.reset_undo()

    def _record_state(self, event=None):
        if self._internal_change:
            return

        # Don't record Ctrl+Z/Y/A/C/V/X as new history entries.
        if event and (event.state & 0x4) and event.keysym.lower() in {"z", "y", "a", "c", "v", "x"}:
            return

        current = self.get()
        history_list = list(self._history)

        if self._history_pos < 0:
            self.reset_undo()
            return

        if self._history_pos < len(history_list) and current == history_list[self._history_pos]:
            return

        # Truncate any redo future, then append
        # deque doesn't support slicing, so we rebuild from the kept portion
        kept = list(self._history)[: self._history_pos + 1]
        self._history.clear()
        for item in kept:
            self._history.append(item)
        self._history.append(current)
        self._history_pos = len(self._history) - 1

    def _set_value(self, value: str):
        # The linked StringVar updates automatically through widget edits.
        self._internal_change = True
        try:
            self.delete(0, tk.END)
            self.insert(0, value)
            self.icursor(tk.END)
        finally:
            self._internal_change = False

    def undo(self, event=None):
        self._sync_history()
        history_list = list(self._history)
        if self._history_pos > 0:
            self._history_pos -= 1
            self._set_value(history_list[self._history_pos])
        return "break"

    def redo(self, event=None):
        history_list = list(self._history)
        if self._history_pos < len(history_list) - 1:
            self._history_pos += 1
            self._set_value(history_list[self._history_pos])
        return "break"

    def select_all(self, event=None):
        self.selection_range(0, tk.END)
        self.icursor(tk.END)
        return "break"


# =============================================================================
# CROSS-PLATFORM FOLDER OPENER
# =============================================================================
def open_folder_in_explorer(path: Path) -> None:
    """Open the given folder in the OS file manager (Windows / macOS / Linux)."""
    try:
        system = platform.system()
        if system == "Windows":
            os.startfile(str(path))
        elif system == "Darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception as exc:
        logger.warning("Could not open folder %s: %s", path, exc)


# =============================================================================
# MAIN APP
# =============================================================================
class MasterApp(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title(APP_TITLE)
        set_window_icon(self)
        screen_w = max(self.winfo_screenwidth(), 1024)
        screen_h = max(self.winfo_screenheight(), 720)
        window_w = min(1320, max(980, int(screen_w * 0.88)))
        window_h = min(860, max(680, int(screen_h * 0.86)))
        self.geometry(f"{window_w}x{window_h}")
        self.minsize(min(980, screen_w - 80), min(680, screen_h - 80))

        self.folder_var = tk.StringVar()
        self.status_var = tk.StringVar(value=t("status_ready"))
        self.detected_count_var = tk.StringVar(value=t("detected_count", count=0))
        self.selected_count_var = tk.StringVar(value=t("selected_count", count=0))
        self.expected_samples_var = tk.StringVar(value="10")
        self.repeated_times_var = tk.StringVar(value="3")

        self.profiles = load_profiles()

        self.selected_profile_var = tk.StringVar()
        if self.profiles:
            self.selected_profile_var.set(next(iter(self.profiles.keys())))

        self.profile_file_hint_var = tk.StringVar()

        self.grouped_logs = {}

        self.available_components: List[str] = []
        self.selected_components: List[str] = []
        self.parse_warnings: List[str] = []
        self.duplicate_component_choices: List[dict] = []
        self.reviewed_duplicate_components = set()
        self.busy_widgets: List[tk.Widget] = []

        # Shared cancellation flag for the background scan thread.
        self._scan_cancel = threading.Event()

        self._build_modern_ui()
        self.update_profile_hint()

    # -------------------------------------------------------------------------
    # UI
    # -------------------------------------------------------------------------
    def _build_modern_ui(self):
        self.configure(bg="#eef6ff")
        self._responsive_labels: List[tk.Widget] = []

        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except Exception:
            pass

        palette = {
            "bg": "#eef6ff",
            "card": "#ffffff",
            "ink": "#10243e",
            "muted": "#52677f",
            "brand": "#0b84ff",
            "brand_dark": "#0564c8",
            "line": "#b9d8ff",
            "list_bg": "#f8fbff",
            "selected_bg": "#edf6ff",
            "danger": "#b42318",
        }

        style.configure("App.TFrame", background=palette["bg"])
        style.configure("Card.TFrame", background=palette["card"])
        style.configure("Toolbar.TFrame", background=palette["card"])
        style.configure("Card.TLabelframe", background=palette["card"], borderwidth=1, relief="solid")
        style.configure(
            "Card.TLabelframe.Label",
            background=palette["card"],
            foreground=palette["brand_dark"],
            font=("Segoe UI Semibold", 11),
        )

        style.configure("TLabel", background=palette["bg"], foreground=palette["ink"], font=("Segoe UI", 10))
        style.configure("Card.TLabel", background=palette["card"], foreground=palette["ink"], font=("Segoe UI", 10))
        style.configure("HeroTitle.TLabel", background=palette["bg"], foreground=palette["ink"], font=("Segoe UI Semibold", 20))
        style.configure("HeroSub.TLabel", background=palette["bg"], foreground=palette["muted"], font=("Segoe UI", 9))
        style.configure("Meta.TLabel", background=palette["card"], foreground=palette["muted"], font=("Segoe UI", 9))
        style.configure("Stat.TLabel", background=palette["card"], foreground=palette["brand_dark"], font=("Segoe UI Semibold", 11))
        style.configure("Version.TLabel", background=palette["bg"], foreground=palette["brand_dark"], font=("Segoe UI Semibold", 9))

        style.configure("TButton", font=("Segoe UI", 10), padding=(8, 5))
        style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"), foreground="#ffffff", background=palette["brand"])
        style.map("Accent.TButton", background=[("active", palette["brand_dark"]), ("pressed", palette["brand_dark"])])
        style.configure("Export.TButton", font=("Segoe UI", 10, "bold"), foreground="#ffffff", background=palette["brand_dark"])
        style.map("Export.TButton", background=[("active", "#034f9f"), ("pressed", "#034f9f")])
        style.configure("Cancel.TButton", font=("Segoe UI", 10), foreground="#ffffff", background=palette["danger"])
        style.map("Cancel.TButton", background=[("active", "#912018"), ("pressed", "#912018")])
        style.configure("Action.TButton", font=("Segoe UI", 10), padding=(8, 5))
        style.configure("Subtle.TButton", font=("Segoe UI", 10), padding=(8, 5), foreground=palette["brand_dark"])

        style.configure("TEntry", fieldbackground="#ffffff")
        style.configure("TCombobox", fieldbackground="#ffffff")
        style.configure("Treeview", background="#ffffff", fieldbackground="#ffffff", rowheight=28, font=("Segoe UI", 9))
        style.configure("Treeview.Heading", font=("Segoe UI Semibold", 9), background="#dcebff", foreground=palette["ink"])

        hero_content = ttk.Frame(self, style="App.TFrame", padding=(16, 8, 16, 6))
        hero_content.pack(fill="x")
        hero_content.columnconfigure(0, weight=1)
        hero_content.columnconfigure(1, weight=1)

        title_block = ttk.Frame(hero_content, style="App.TFrame")
        title_block.grid(row=0, column=0, sticky="ew")
        ttk.Label(title_block, text=APP_TITLE, style="HeroTitle.TLabel").pack(anchor="w")
        ttk.Label(title_block, text=APP_DEVELOPER_LINE, style="HeroSub.TLabel").pack(anchor="w")

        self.header_logo_source_image = None
        self.header_logo_image = None
        self.header_logo_label = None
        logo_path = resource_path("assets", "Jabil_logo.png")
        if not logo_path.exists():
            logo_path = resource_path("assets", "DataMapperX.png")
        if logo_path.exists():
            try:
                self.header_logo_source_image = tk.PhotoImage(file=str(logo_path))
                self.header_logo_image = self.header_logo_source_image.subsample(HEADER_LOGO_SUBSAMPLE)
                self.header_logo_label = ttk.Label(hero_content, image=self.header_logo_image, style="TLabel")
                self.header_logo_label.place(relx=0.5, y=0, anchor="n")
            except tk.TclError:
                logger.debug("Could not load header logo: %s", logo_path, exc_info=True)

        version_block = ttk.Frame(hero_content, style="App.TFrame")
        version_block.grid(row=0, column=1, sticky="ne", padx=(16, 0))
        ttk.Label(version_block, text=APP_VERSION, style="Version.TLabel", anchor="e").pack(anchor="e")
        ttk.Label(version_block, text=APP_BUILD, style="HeroSub.TLabel", anchor="e").pack(anchor="e")

        shell = ttk.Frame(self, style="App.TFrame", padding=(12, 6, 12, 8))
        shell.pack(fill="both", expand=True)

        main_pane = ttk.Panedwindow(shell, orient=tk.VERTICAL)
        main_pane.pack(fill="both", expand=True)

        top_panel = ttk.Frame(main_pane, style="App.TFrame")
        preview_panel = ttk.Frame(main_pane, style="App.TFrame")
        main_pane.add(top_panel, weight=2)
        main_pane.add(preview_panel, weight=8)

        top_pane = tk.PanedWindow(
            top_panel,
            orient=tk.HORIZONTAL,
            bg=palette["bg"],
            bd=0,
            sashwidth=6,
            sashrelief=tk.RAISED,
            opaqueresize=True,
        )
        top_pane.pack(fill="both", expand=True)

        workspace_slot = ttk.Frame(top_pane, style="App.TFrame")
        available_slot = ttk.Frame(top_pane, style="App.TFrame")
        selected_slot = ttk.Frame(top_pane, style="App.TFrame")

        workspace_card = ttk.LabelFrame(workspace_slot, text=t("workspace_setup"), style="Card.TLabelframe", padding=10)
        available_card = ttk.LabelFrame(available_slot, text=t("available_components"), style="Card.TLabelframe", padding=10)
        selected_card = ttk.LabelFrame(selected_slot, text=t("selected_components"), style="Card.TLabelframe", padding=10)
        workspace_card.pack(fill="both", expand=True, padx=(0, 4))
        available_card.pack(fill="both", expand=True, padx=(0, 4))
        selected_card.pack(fill="both", expand=True, padx=(4, 0))

        top_pane.add(workspace_slot, minsize=360)
        top_pane.add(available_slot, minsize=210)
        top_pane.add(selected_slot, minsize=360)

        def set_initial_top_sashes():
            width = top_pane.winfo_width()
            if width <= 1:
                self.after(50, set_initial_top_sashes)
                return

            workspace_width = max(400, min(500, int(width * 0.38)))
            available_width = max(220, min(300, int(width * 0.18)))
            top_pane.sash_place(0, workspace_width, 0)
            top_pane.sash_place(1, workspace_width + available_width, 0)

        self.after_idle(set_initial_top_sashes)

        workspace_card.columnconfigure(0, weight=1)

        profile_header = ttk.Frame(workspace_card, style="Card.TFrame")
        profile_header.grid(row=0, column=0, sticky="ew")
        profile_header.columnconfigure(0, weight=1)
        ttk.Label(profile_header, text=t("profile"), style="Card.TLabel").grid(row=0, column=0, sticky="w")
        manage_button = ttk.Button(profile_header, text=t("manage_profiles"), command=self.manage_profiles, style="Subtle.TButton")
        manage_button.grid(row=0, column=1, sticky="e")
        self.busy_widgets.append(manage_button)

        self.profile_combo = ttk.Combobox(
            workspace_card,
            textvariable=self.selected_profile_var,
            values=list(self.profiles.keys()),
            state="readonly",
        )
        self.profile_combo.grid(row=1, column=0, sticky="ew", pady=(4, 4))
        self.profile_combo.bind("<<ComboboxSelected>>", self.on_profile_changed, add="+")

        self.profile_file_hint_label = ttk.Label(
            workspace_card,
            textvariable=self.profile_file_hint_var,
            style="Meta.TLabel",
            anchor="w",
        )
        self.profile_file_hint_label.grid(row=2, column=0, sticky="ew", pady=(0, 8))
        self._responsive_labels.append(self.profile_file_hint_label)

        ttk.Separator(workspace_card).grid(row=3, column=0, sticky="ew", pady=(2, 8))
        ttk.Label(workspace_card, text=t("source_folder"), style="Card.TLabel").grid(row=4, column=0, sticky="w")
        entry = ttk.Entry(workspace_card, textvariable=self.folder_var)
        entry.grid(row=5, column=0, sticky="ew", pady=(4, 6))

        folder_actions = ttk.Frame(workspace_card, style="Card.TFrame")
        folder_actions.grid(row=6, column=0, sticky="ew")
        for col in range(4):
            folder_actions.columnconfigure(col, weight=1, uniform="folder_actions")

        browse_button = ttk.Button(folder_actions, text=t("browse"), command=self.pick_folder, style="Subtle.TButton")
        browse_button.grid(row=0, column=0, sticky="ew", padx=(0, 4))
        self.refresh_button = ttk.Button(folder_actions, text="Refresh", style="Subtle.TButton", command=self.reset_workflow)
        self.refresh_button.grid(row=0, column=1, sticky="ew", padx=8)
        self.scan_button = ttk.Button(folder_actions, text=t("scan_logs"), style="Accent.TButton", command=self.scan_logs)
        self.scan_button.grid(row=0, column=2, sticky="ew", padx=4)
        self.cancel_button = ttk.Button(folder_actions, text="Cancel", style="Cancel.TButton", command=self._cancel_scan)
        self.cancel_button.grid(row=0, column=3, sticky="ew", padx=(4, 0))
        self.cancel_button.configure(state="disabled")

        self.busy_widgets.extend([browse_button, self.scan_button, self.profile_combo, self.refresh_button])

        ttk.Separator(workspace_card).grid(row=7, column=0, sticky="ew", pady=10)
        settings_grid = ttk.Frame(workspace_card, style="Card.TFrame")
        settings_grid.grid(row=8, column=0, sticky="ew")
        settings_grid.columnconfigure(1, weight=1)
        settings_grid.columnconfigure(3, weight=1)
        ttk.Label(settings_grid, text=t("expected_samples"), style="Card.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 6))
        ttk.Entry(settings_grid, textvariable=self.expected_samples_var, width=8).grid(row=0, column=1, sticky="ew", padx=(0, 10))
        ttk.Label(settings_grid, text=t("repeats_per_sample"), style="Card.TLabel").grid(row=0, column=2, sticky="w", padx=(0, 6))
        ttk.Entry(settings_grid, textvariable=self.repeated_times_var, width=8).grid(row=0, column=3, sticky="ew")

        stats_box = ttk.Frame(workspace_card, style="Card.TFrame")
        stats_box.grid(row=9, column=0, sticky="ew", pady=(10, 0))
        ttk.Label(stats_box, textvariable=self.detected_count_var, style="Stat.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 40))
        ttk.Label(stats_box, textvariable=self.selected_count_var, style="Stat.TLabel").grid(row=0, column=1, sticky="e")
        available_card.rowconfigure(1, weight=1)
        available_card.columnconfigure(0, weight=1)

        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(available_card, textvariable=self.search_var)
        search_entry.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        search_entry.bind("<KeyRelease>", lambda _event: self.refresh_available_list())

        available_frame = ttk.Frame(available_card, style="Card.TFrame")
        available_frame.grid(row=1, column=0, sticky="nsew")
        available_frame.rowconfigure(0, weight=1)
        available_frame.columnconfigure(0, weight=1)
        self.available_list = tk.Listbox(
            available_frame,
            selectmode=tk.EXTENDED,
            font=("Cascadia Mono", 9),
            activestyle="none",
            bg=palette["list_bg"],
            fg=palette["ink"],
            selectbackground=palette["brand"],
            selectforeground="#ffffff",
            highlightthickness=1,
            highlightbackground=palette["line"],
            bd=0,
        )
        av_scroll = ttk.Scrollbar(available_frame, orient="vertical", command=self.available_list.yview)
        self.available_list.configure(yscrollcommand=av_scroll.set)
        self.available_list.grid(row=0, column=0, sticky="nsew")
        av_scroll.grid(row=0, column=1, sticky="ns")
        self.available_list.bind("<Double-Button-1>", lambda _event: self.add_components())
        self.available_list.bind("<Control-a>", self._select_all_available)
        self.available_list.bind("<Control-A>", self._select_all_available)

        selected_card.rowconfigure(0, weight=1)
        selected_card.columnconfigure(0, weight=1)

        selected_frame = ttk.Frame(selected_card, style="Card.TFrame")
        selected_frame.grid(row=0, column=0, sticky="nsew")
        selected_frame.rowconfigure(0, weight=1)
        selected_frame.columnconfigure(0, weight=1)
        self.selected_list = tk.Listbox(
            selected_frame,
            selectmode=tk.EXTENDED,
            font=("Cascadia Mono", 9),
            activestyle="none",
            bg=palette["selected_bg"],
            fg=palette["ink"],
            selectbackground=palette["brand"],
            selectforeground="#ffffff",
            highlightthickness=1,
            highlightbackground=palette["line"],
            bd=0,
        )
        sel_scroll = ttk.Scrollbar(selected_frame, orient="vertical", command=self.selected_list.yview)
        self.selected_list.configure(yscrollcommand=sel_scroll.set)
        self.selected_list.grid(row=0, column=0, sticky="nsew")
        sel_scroll.grid(row=0, column=1, sticky="ns")
        self.selected_list.bind("<Double-Button-1>", lambda _event: self.remove_components())
        self.selected_list.bind("<Delete>", lambda _event: self.remove_components())

        action_wrap = ttk.Frame(selected_card, style="Card.TFrame")
        action_wrap.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        for col in range(5):
            action_wrap.columnconfigure(col, weight=1, uniform="component_actions")
        add_button = ttk.Button(action_wrap, text=t("add"), command=self.add_components, style="Accent.TButton")
        remove_button = ttk.Button(action_wrap, text=t("remove"), command=self.remove_components, style="Subtle.TButton")
        up_button = ttk.Button(action_wrap, text=t("move_up"), command=lambda: self.move_selected(-1), style="Action.TButton")
        down_button = ttk.Button(action_wrap, text=t("move_down"), command=lambda: self.move_selected(1), style="Action.TButton")
        clear_button = ttk.Button(action_wrap, text=t("clear_all"), command=self.clear_selected, style="Subtle.TButton")
        add_button.grid(row=0, column=0, sticky="ew", padx=(0, 4))
        remove_button.grid(row=0, column=1, sticky="ew", padx=4)
        up_button.grid(row=0, column=2, sticky="ew", padx=4)
        down_button.grid(row=0, column=3, sticky="ew", padx=4)
        clear_button.grid(row=0, column=4, sticky="ew", padx=(4, 0))

        preview_card = ttk.LabelFrame(preview_panel, text=t("preview_export"), style="Card.TLabelframe", padding=10)
        preview_card.pack(fill="both", expand=True, pady=(8, 0))
        preview_card.rowconfigure(1, weight=1)
        preview_card.columnconfigure(0, weight=1)

        toolbar = ttk.Frame(preview_card, style="Toolbar.TFrame")
        toolbar.grid(row=0, column=0, sticky="ew")
        toolbar.columnconfigure(0, weight=1)
        self.preview_button = ttk.Button(toolbar, text=t("preview"), command=self.preview_summary)
        self.preview_button.grid(row=0, column=1, sticky="e", padx=(0, 8))
        self.export_button = ttk.Button(toolbar, text=t("export_excel"), style="Export.TButton", command=self.export_excel)
        self.export_button.grid(row=0, column=2, sticky="e")
        self.busy_widgets.extend([self.preview_button, self.export_button])

        tree_frame = ttk.Frame(preview_card, style="Card.TFrame")
        tree_frame.grid(row=1, column=0, sticky="nsew", pady=(8, 0))
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        columns = ("stt", "sn", "count", "files")
        self.preview_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=14)
        self.preview_tree.heading("stt", text=t("tree_no"))
        self.preview_tree.heading("sn", text=t("tree_sn"))
        self.preview_tree.heading("count", text=t("tree_count"))
        self.preview_tree.heading("files", text=t("tree_files"))
        self.preview_tree.column("stt", width=64, anchor="center", stretch=False)
        self.preview_tree.column("sn", width=220, anchor="w", stretch=True)
        self.preview_tree.column("count", width=80, anchor="center", stretch=False)
        self.preview_tree.column("files", width=980, anchor="w", stretch=True)
        tree_frame.bind("<Configure>", self._resize_preview_columns, add="+")
        tree_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.preview_tree.yview)
        tree_xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.preview_tree.xview)
        self.preview_tree.configure(yscrollcommand=tree_scroll.set, xscrollcommand=tree_xscroll.set)
        self.preview_tree.grid(row=0, column=0, sticky="nsew")
        tree_scroll.grid(row=0, column=1, sticky="ns")
        tree_xscroll.grid(row=1, column=0, sticky="ew")

        status_bar = ttk.Frame(self, style="App.TFrame", padding=(12, 4, 12, 8))
        status_bar.pack(fill="x")
        self.progress = ttk.Progressbar(status_bar, mode="indeterminate", length=170)
        self.progress.pack(side="right", padx=(8, 0))
        ttk.Label(
            status_bar,
            textvariable=self.status_var,
            foreground="#30445f",
            background=palette["bg"],
            font=("Segoe UI", 10),
        ).pack(side="left", fill="x", expand=True)
        self.bind("<Configure>", self._resize_text_widgets, add="+")

    def _resize_text_widgets(self, _event=None):
        wrap = max(220, min(560, int(self.winfo_width() * 0.38)))
        for widget in getattr(self, "_responsive_labels", []):
            try:
                widget.configure(wraplength=wrap)
            except tk.TclError:
                pass

    def _resize_preview_columns(self, event=None):
        if event is None or not hasattr(self, "preview_tree"):
            return

        width = max(event.width - 42, 560)
        fixed = 64 + 80
        sn_width = max(190, min(320, int(width * 0.2)))
        files_width = max(720, width - fixed - sn_width)
        self.preview_tree.column("stt", width=64)
        self.preview_tree.column("sn", width=sn_width)
        self.preview_tree.column("count", width=80)
        self.preview_tree.column("files", width=files_width)

    def _select_all_available(self, _event=None):
        self.available_list.select_set(0, tk.END)
        return "break"

    def format_profile_file_hint(self) -> str:
        profile = self.profiles.get(self.selected_profile_var.get(), {})
        if not profile:
            return ""

        extensions = ", ".join(get_profile_extensions(profile))
        input_type = t("profile_type_excel") if profile.get("type") == "excel" else t("profile_type_text")

        return t("profile_file_hint", input_type=input_type, extensions=extensions)

    def update_profile_hint(self):
        if hasattr(self, "profile_file_hint_var"):
            self.profile_file_hint_var.set(self.format_profile_file_hint())

    def on_profile_changed(self, _event=None):
        self.update_profile_hint()

        self.grouped_logs = {}
        self.available_components = []
        self.selected_components = []
        self.parse_warnings = []
        self.duplicate_component_choices = []
        self.reviewed_duplicate_components = set()

        if hasattr(self, "available_list"):
            self.available_list.delete(0, tk.END)

        if hasattr(self, "selected_list"):
            self.selected_list.delete(0, tk.END)

        if hasattr(self, "preview_tree"):
            self.preview_tree.delete(*self.preview_tree.get_children())

        self.detected_count_var.set(t("detected_count", count=0))
        self.selected_count_var.set(t("selected_count", count=0))
        self.set_status(f"Profile changed. {self.format_profile_file_hint()}")

    def set_status(self, message: str, loading: bool = False):
        self.status_var.set(message)
        if loading:
            self.progress.start(10)
        else:
            self.progress.stop()
        self.update_idletasks()

    def set_busy(self, busy: bool):
        state = "disabled" if busy else "normal"
        for widget in self.busy_widgets:
            try:
                if isinstance(widget, ttk.Combobox):
                    widget.configure(state="disabled" if busy else "readonly")
                else:
                    widget.configure(state=state)
            except tk.TclError:
                pass

        try:
            self.cancel_button.configure(state="normal" if busy else "disabled")
        except tk.TclError:
            pass

    # -------------------------------------------------------------------------
    # PROFILE MANAGER
    # -------------------------------------------------------------------------
    def manage_profiles(self):
        top = tk.Toplevel(self)
        top.title(t("manage_title"))
        set_window_icon(top)
        top.geometry("760x640")
        top.minsize(700, 580)
        top.transient(self)

        main_frame = ttk.Frame(top, padding=12)
        main_frame.pack(fill="both", expand=True)

        left = ttk.Frame(main_frame)
        left.pack(side="left", fill="y", padx=(0, 10))

        ttk.Label(left, text=t("profiles"), font=("Segoe UI", 11, "bold")).pack(anchor="w")

        profile_listbox = tk.Listbox(left, height=15, width=27)
        profile_listbox.pack(fill="y", expand=True, pady=(5, 5))

        for profile_name in self.profiles:
            profile_listbox.insert(tk.END, profile_name)

        btn_frame = ttk.Frame(left)
        btn_frame.pack(fill="x")

        right = ttk.LabelFrame(main_frame, text=t("edit_profile"), padding=12)
        right.pack(side="left", fill="both", expand=True)

        name_var = tk.StringVar()
        original_name_var = tk.StringVar()
        type_var = tk.StringVar(value="text")
        sn_var = tk.StringVar()

        block_var = tk.StringVar()
        measure_var = tk.StringVar()
        lim2_var = tk.StringVar()
        lim3_var = tk.StringVar()

        sn_sheet_var = tk.StringVar()
        data_sheet_var = tk.StringVar()
        row_start_var = tk.StringVar(value="2")
        comp_col_var = tk.StringVar(value="A")
        meas_col_var = tk.StringVar(value="B")
        min_col_var = tk.StringVar(value="C")
        max_col_var = tk.StringVar(value="D")
        unit_col_var = tk.StringVar(value="E")
        test_input_var = tk.StringVar()
        test_result_var = tk.StringVar(value="Result: ")

        def reset_editor():
            original_name_var.set("")
            name_var.set("New Profile")
            type_var.set("text")
            sn_var.set("")
            block_var.set("")
            measure_var.set("")
            lim2_var.set("")
            lim3_var.set("")
            sn_sheet_var.set("")
            data_sheet_var.set("")
            row_start_var.set("2")
            comp_col_var.set("A")
            meas_col_var.set("B")
            min_col_var.set("C")
            max_col_var.set("D")
            unit_col_var.set("E")
            test_input_var.set("")
            test_result_var.set("Result: ")

            save_btn.grid()
            save_btn.config(state="normal")
            set_profile_readonly(False)
            update_fields_visibility()

        def on_select_profile(_event=None):
            selection = profile_listbox.curselection()
            if not selection:
                return

            profile_name = profile_listbox.get(selection[0])
            profile_data = self.profiles[profile_name]

            original_name_var.set(profile_name)
            name_var.set(profile_name)
            type_var.set(profile_data.get("type", "text"))
            sn_var.set(profile_data.get("sn_regex", ""))

            block_var.set(profile_data.get("block_regex", ""))
            measure_var.set(profile_data.get("measure_regex", ""))
            lim2_var.set(profile_data.get("lim2_regex", ""))
            lim3_var.set(profile_data.get("lim3_regex", ""))

            sn_sheet_var.set(profile_data.get("sn_sheet", ""))
            data_sheet_var.set(profile_data.get("data_sheet", ""))
            row_start_var.set(profile_data.get("row_start", "2"))
            comp_col_var.set(profile_data.get("component_col", "A"))
            meas_col_var.set(profile_data.get("measure_col", "B"))
            min_col_var.set(profile_data.get("min_col", "C"))
            max_col_var.set(profile_data.get("max_col", "D"))
            unit_col_var.set(profile_data.get("unit_col", "E"))

            update_fields_visibility()
            is_built_in = profile_name in BUILT_IN_PROFILE_NAMES

            if is_built_in:
                save_btn.grid_remove()
                set_profile_readonly(True)
            else:
                save_btn.grid()
                save_btn.config(state="normal")
                set_profile_readonly(False)

        profile_listbox.bind("<<ListboxSelect>>", on_select_profile)

        def new_profile():
            profile_listbox.selection_clear(0, tk.END)
            reset_editor()

        ttk.Button(btn_frame, text=t("new"), command=new_profile).pack(side="left", expand=True, fill="x")

        def delete_profile():
            selection = profile_listbox.curselection()
            if not selection:
                return

            profile_name = profile_listbox.get(selection[0])
            if profile_name in BUILT_IN_PROFILE_NAMES:
                messagebox.showerror(APP_TITLE, "Cannot delete built-in profile.", parent=top)
                return

            if messagebox.askyesno(t("delete"), f"Delete {profile_name}?", parent=top):
                del self.profiles[profile_name]
                save_profiles(self.profiles)
                profile_listbox.delete(selection[0])
                self.profile_combo["values"] = list(self.profiles.keys())
                if self.selected_profile_var.get() == profile_name:
                    self.selected_profile_var.set(next(iter(self.profiles.keys()), ""))

                self.update_profile_hint()
                reset_editor()

        ttk.Button(btn_frame, text=t("delete"), command=delete_profile).pack(side="left", expand=True, fill="x")

        ttk.Label(right, text=t("profile_name")).grid(row=0, column=0, sticky="w", pady=5)
        name_entry = UndoableEntry(right, textvariable=name_var)
        name_entry.grid(row=0, column=1, sticky="ew", pady=5)

        ttk.Label(right, text=t("input_type")).grid(row=1, column=0, sticky="w", pady=5)
        type_combo = ttk.Combobox(right, textvariable=type_var, values=["text", "excel"], state="readonly")
        type_combo.grid(row=1, column=1, sticky="ew", pady=5)

        ttk.Label(right, text=t("sn_regex_cell")).grid(row=2, column=0, sticky="w", pady=5)
        UndoableEntry(right, textvariable=sn_var).grid(row=2, column=1, sticky="ew", pady=5)
        ttk.Label(right, text=t("sn_tip"), font=("Segoe UI", 8, "italic"), foreground="#888").grid(
            row=3, column=1, sticky="w", pady=(0, 5)
        )

        text_frame = ttk.Frame(right)

        def show_regex_help():
            help_text = (
                "To extract data from text logs, use Python Regex.\n\n"
                "RECOMMENDED GROUPS:\n"
                "• Name_Regex: capture component as (?P<component>...)\n"
                "• Measure_Regex: capture value as (?P<measured>...)\n"
                "  Optional: (?P<type>...) for Unit, (?P<tail>...) for limits\n"
                "• Min_Max regex: capture (?P<min>...) and (?P<max>...)\n\n"
                "If you are not familiar with Regex, duplicate an existing profile in profiles.json "
                "or ask your technical lead for the correct pattern."
            )
            messagebox.showinfo("Regex Cheat Sheet", help_text, parent=top)

        ttk.Button(text_frame, text=t("regex_help"), command=show_regex_help).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))

        ttk.Label(text_frame, text=t("name_regex")).grid(row=1, column=0, sticky="w", pady=5)
        UndoableEntry(text_frame, textvariable=block_var).grid(row=1, column=1, sticky="ew", pady=5)
        ttk.Label(text_frame, text=t("measure_regex")).grid(row=2, column=0, sticky="w", pady=5)
        UndoableEntry(text_frame, textvariable=measure_var).grid(row=2, column=1, sticky="ew", pady=5)
        ttk.Label(text_frame, text=t("minmax_regex1")).grid(row=3, column=0, sticky="w", pady=5)
        UndoableEntry(text_frame, textvariable=lim2_var).grid(row=3, column=1, sticky="ew", pady=5)
        ttk.Label(text_frame, text=t("minmax_regex2")).grid(row=4, column=0, sticky="w", pady=5)
        UndoableEntry(text_frame, textvariable=lim3_var).grid(row=4, column=1, sticky="ew", pady=5)

        test_frame = ttk.LabelFrame(text_frame, text=t("regex_tester"), padding=8)
        test_frame.grid(row=5, column=0, columnspan=2, sticky="ew", pady=(10, 0))
        ttk.Label(test_frame, text=t("paste_data_row")).grid(row=0, column=0, sticky="w", pady=2)
        test_input_entry = ttk.Entry(test_frame, textvariable=test_input_var)
        test_input_entry.grid(row=1, column=0, sticky="ew", pady=2)
        ttk.Label(test_frame, textvariable=test_result_var, font=("Segoe UI", 9, "bold"), foreground="blue").grid(
            row=2, column=0, sticky="w", pady=(4, 0)
        )
        test_frame.columnconfigure(0, weight=1)
        text_frame.columnconfigure(1, weight=1)

        def evaluate_test(*_args):
            data = test_input_var.get().strip()
            if not data:
                test_result_var.set(t("result_empty"))
                return

            try:
                block_pattern = block_var.get().strip()
                measure_pattern = measure_var.get().strip()

                if block_pattern:
                    match = re.search(block_pattern, data, re.IGNORECASE)
                    if match:
                        if "component" in match.groupdict():
                            test_result_var.set(f"Result: MATCHED BLOCK - Component: {match.group('component')}")
                        elif match.groups():
                            test_result_var.set(f"Result: MATCHED BLOCK - Component: {match.group(1)}")
                        else:
                            test_result_var.set("Result: Matched Block but no component group found")
                        return

                if measure_pattern:
                    match = re.search(measure_pattern, data, re.IGNORECASE)
                    if match:
                        result = []
                        if "measured" in match.groupdict():
                            result.append(f"Measured: {match.group('measured')}")
                        elif match.groups():
                            result.append(f"Measured: {match.group(1)}")

                        tail = match.group("tail") if "tail" in match.groupdict() else data
                        limit_found = False

                        for pattern_text in (lim2_var.get().strip(), lim3_var.get().strip()):
                            if not pattern_text:
                                continue
                            limit_match = re.search(pattern_text, tail)
                            if limit_match:
                                limit_found = True
                                result.append(f"Min: {limit_match.group('min')}" if "min" in limit_match.groupdict() else "Min: ?")
                                result.append(f"Max: {limit_match.group('max')}" if "max" in limit_match.groupdict() else "Max: ?")
                                break

                        if not result:
                            result.append("Matched, but no measured group found")
                        if not limit_found:
                            result.append("Limits: not found")

                        test_result_var.set("Result: MATCHED MEASURE - " + ", ".join(result))
                        return

                test_result_var.set(t("result_unrecognized"))
            except re.error as exc:
                test_result_var.set(f"Result: Regex error - {exc}")

        for var in (test_input_var, block_var, measure_var, lim2_var, lim3_var):
            var.trace_add("write", evaluate_test)

        excel_frame = ttk.Frame(right)
        ttk.Label(excel_frame, text=t("excel_sn_sheet")).grid(row=0, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=sn_sheet_var).grid(row=0, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_data_sheet")).grid(row=1, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=data_sheet_var).grid(row=1, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_row_start")).grid(row=2, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=row_start_var).grid(row=2, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_component_col")).grid(row=3, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=comp_col_var).grid(row=3, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_measure_col")).grid(row=4, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=meas_col_var).grid(row=4, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_min_col")).grid(row=5, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=min_col_var).grid(row=5, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_max_col")).grid(row=6, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=max_col_var).grid(row=6, column=1, sticky="ew", pady=5)
        ttk.Label(excel_frame, text=t("excel_unit_col")).grid(row=7, column=0, sticky="w", pady=5)
        UndoableEntry(excel_frame, textvariable=unit_col_var).grid(row=7, column=1, sticky="ew", pady=5)
        excel_frame.columnconfigure(1, weight=1)

        def update_fields_visibility(*_args):
            if type_var.get() == "text":
                excel_frame.grid_remove()
                text_frame.grid(row=4, column=0, columnspan=2, sticky="ew")
            else:
                text_frame.grid_remove()
                excel_frame.grid(row=4, column=0, columnspan=2, sticky="ew")

        type_var.trace_add("write", update_fields_visibility)
        update_fields_visibility()

        right.columnconfigure(1, weight=1)

        def set_profile_readonly(readonly: bool):
            entry_state = "readonly" if readonly else "normal"
            combo_state = "disabled" if readonly else "readonly"

            def walk(widget):
                for child in widget.winfo_children():
                    if child is test_input_entry:
                        continue
                    if isinstance(child, ttk.Entry):
                        child.configure(state=entry_state)
                    elif isinstance(child, ttk.Combobox):
                        child.configure(state=combo_state)
                    walk(child)

            walk(right)
            name_entry.configure(state=entry_state)
            type_combo.configure(state=combo_state)

        def save_current_profile():
            profile_name = name_var.get().strip()
            if not profile_name:
                messagebox.showwarning(APP_TITLE, "Profile name cannot be empty.", parent=top)
                return

            if profile_name in BUILT_IN_PROFILE_NAMES:
                messagebox.showwarning(APP_TITLE, "Cannot override built-in profiles.", parent=top)
                return

            profile_data = {"type": type_var.get(), "sn_regex": sn_var.get()}
            if profile_data["type"] == "text":
                profile_data["file_extensions"] = ",".join(TEXT_FILE_EXTENSIONS)
                profile_data["block_regex"] = block_var.get()
                profile_data["measure_regex"] = measure_var.get()
                profile_data["lim2_regex"] = lim2_var.get()
                profile_data["lim3_regex"] = lim3_var.get()
            else:
                profile_data["file_extensions"] = ",".join(EXCEL_FILE_EXTENSIONS)
                profile_data["sn_sheet"] = sn_sheet_var.get()
                profile_data["data_sheet"] = data_sheet_var.get()
                profile_data["row_start"] = row_start_var.get()
                profile_data["component_col"] = comp_col_var.get()
                profile_data["measure_col"] = meas_col_var.get()
                profile_data["min_col"] = min_col_var.get()
                profile_data["max_col"] = max_col_var.get()
                profile_data["unit_col"] = unit_col_var.get()

            try:
                validate_profile(profile_data)
            except ValueError as exc:
                messagebox.showerror(APP_TITLE, str(exc), parent=top)
                return

            original_name = original_name_var.get().strip()
            if original_name and original_name != profile_name and original_name not in BUILT_IN_PROFILE_NAMES:
                self.profiles.pop(original_name, None)

            self.profiles[profile_name] = profile_data
            save_profiles(self.profiles)

            profile_listbox.delete(0, tk.END)
            for item in self.profiles:
                profile_listbox.insert(tk.END, item)

            self.profile_combo["values"] = list(self.profiles.keys())
            self.selected_profile_var.set(profile_name)
            self.update_profile_hint()
            original_name_var.set(profile_name)

            messagebox.showinfo("Saved", f"Profile '{profile_name}' saved.", parent=top)

        # Keep the save action below the dynamic profile field frames.
        save_frame = ttk.Frame(right)
        save_frame.grid(row=6, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        save_btn = ttk.Button(save_frame, text=t("save_profile"), style="Accent.TButton", command=save_current_profile)
        save_btn.pack()

        if self.profiles:
            profile_listbox.selection_set(0)
            on_select_profile()

    # -------------------------------------------------------------------------
    # APP ACTIONS
    # -------------------------------------------------------------------------
    def get_expected_samples(self) -> Optional[int]:
        try:
            return parse_positive_int(self.expected_samples_var.get(), "Expected Samples")
        except ValueError as exc:
            messagebox.showwarning(APP_TITLE, str(exc))
            return None

    def get_repeated_times(self) -> Optional[int]:
        try:
            return parse_positive_int(self.repeated_times_var.get(), "Repeats per Sample")
        except ValueError as exc:
            messagebox.showwarning(APP_TITLE, str(exc))
            return None

    def pick_folder(self):
        folder = filedialog.askdirectory(title="Select folder containing log files")
        if folder:
            self.folder_var.set(folder)

    def _cancel_scan(self):
        self._scan_cancel.set()
        self.set_status("Cancelling scan...")

    def _format_extension_counts(self, extension_counts: dict) -> str:
        if not extension_counts:
            return ""

        return ", ".join(
            f"{extension} ({count})"
            for extension, count in sorted(extension_counts.items(), key=lambda item: item[0])
        )

    def _format_scan_no_data_message(
        self,
        folder: Path,
        profile_name: str,
        profile: dict,
        parse_warnings: Sequence[str],
        diagnostics: Optional[dict],
    ) -> str:
        diagnostics = diagnostics or {}
        extensions = ", ".join(diagnostics.get("extensions") or get_profile_extensions(profile))
        profile_type = diagnostics.get("profile_type") or profile.get("type", "text")
        input_type = t("profile_type_excel") if profile_type == "excel" else t("profile_type_text")

        total_files = diagnostics.get("total_file_count", 0)
        accepted_files = diagnostics.get("accepted_file_count", 0)
        skipped_without_sn = diagnostics.get("skipped_without_sn") or []
        files_without_components = diagnostics.get("files_without_components") or []
        ignored_extensions = diagnostics.get("ignored_extensions") or {}

        lines = [
            t("msg_scan_no_data"),
            "",
            f"Selected profile: {profile_name} ({input_type})",
            f"Accepted files: {extensions}",
            f"Source folder: {folder}",
            "",
        ]

        if total_files == 0:
            lines.append("Reason: The selected folder does not contain any scan files.")
        elif accepted_files == 0:
            lines.append("Reason: No files match the selected profile's accepted file types.")
            found_extensions = self._format_extension_counts(ignored_extensions)
            if found_extensions:
                lines.append(f"Files found instead: {found_extensions}")
            lines.append("Action: Choose a profile that matches these files, or select a folder with the accepted file types.")
        elif parse_warnings:
            lines.append(f"Reason: {accepted_files} accepted file(s) were found, but they could not be parsed.")
            lines.append("")
            lines.append(t("msg_parse_warnings"))
            lines.extend(parse_warnings[:8])
        elif skipped_without_sn:
            lines.append(
                f"Reason: {accepted_files} accepted file(s) were scanned, but no serial number matched this profile."
            )
            lines.append("Action: Check SN_Regex / Cell in Manage Profiles, or choose the correct profile.")
            lines.append("")
            lines.append("Example files:")
            lines.extend(skipped_without_sn[:8])
        elif files_without_components:
            lines.append(
                f"Reason: Serial numbers were found, but no components matched this profile's measurement rules."
            )
            lines.append("Action: Check Name_Regex, Measure_Regex, and min/max regex settings in Manage Profiles.")
            lines.append("")
            lines.append("Example files:")
            lines.extend(files_without_components[:8])
        else:
            lines.append("Reason: Accepted files were scanned, but no valid records were produced.")
            lines.append("Action: Verify the selected profile matches the file format.")

        return "\n".join(lines)

    def scan_logs(self):
        folder_text = self.folder_var.get().strip()
        if not folder_text:
            messagebox.showwarning(APP_TITLE, t("msg_choose_folder"))
            return

        folder = Path(folder_text)
        if not folder.exists() or not folder.is_dir():
            messagebox.showerror(APP_TITLE, t("msg_invalid_folder_detail", folder=folder_text))
            return

        expected_sns = self.get_expected_samples()
        expected_repeated = self.get_repeated_times()
        if expected_sns is None or expected_repeated is None:
            return

        profile_key = self.selected_profile_var.get()
        if not profile_key or profile_key not in self.profiles:
            messagebox.showerror(APP_TITLE, t("msg_invalid_profile_detail", profile=profile_key or "(blank)"))
            return

        # Snapshot the profile dict so background scanning is isolated from UI changes.
        current_profile = dict(self.profiles[profile_key])

        self._scan_cancel.clear()
        self.set_busy(True)
        self.set_status(t("status_scanning"), loading=True)
        self.preview_tree.delete(*self.preview_tree.get_children())

        def task():
            try:
                grouped_logs, available_components, parse_warnings, duplicate_component_choices, scan_diagnostics = scan_folder(
                    folder, current_profile, cancel_event=self._scan_cancel, include_diagnostics=True
                )
                self.after(
                    0,
                    self._on_scan_success,
                    grouped_logs,
                    available_components,
                    parse_warnings,
                    expected_sns,
                    expected_repeated,
                    duplicate_component_choices,
                    folder,
                    profile_key,
                    current_profile,
                    scan_diagnostics,
                )
            except Exception as exc:
                self.after(0, self._on_scan_error, str(exc))

        threading.Thread(target=task, daemon=True).start()

    def _on_scan_success(
        self,
        grouped_logs,
        available_components,
        parse_warnings,
        expected_sns,
        expected_repeated,
        duplicate_component_choices=None,
        folder=None,
        profile_name="",
        profile=None,
        scan_diagnostics=None,
    ):
        self.grouped_logs = grouped_logs
        self.available_components = available_components
        self.parse_warnings = parse_warnings
        self.duplicate_component_choices = duplicate_component_choices or []
        self.reviewed_duplicate_components = set()

        self.selected_components = [c for c in self.selected_components if c in self.available_components]
        self.refresh_available_list()
        self.refresh_selected_list()

        self.detected_count_var.set(t("detected_count", count=len(available_components)))

        total_sns = len(grouped_logs)
        total_files = sum(len(items) for items in grouped_logs.values())

        cancelled = self._scan_cancel.is_set()
        status_suffix = " (cancelled)" if cancelled else ""
        self.set_status(
            f"Scan complete{status_suffix}. SN: {total_sns}, files: {total_files}, components: {len(available_components)}",
            loading=False,
        )
        self.set_busy(False)

        if not grouped_logs:
            message = self._format_scan_no_data_message(
                folder=folder or Path(self.folder_var.get().strip() or "."),
                profile_name=profile_name or self.selected_profile_var.get(),
                profile=profile or self.profiles.get(self.selected_profile_var.get(), {}),
                parse_warnings=parse_warnings,
                diagnostics=scan_diagnostics,
            )
            if parse_warnings:
                messagebox.showerror(APP_TITLE, message)
            else:
                messagebox.showwarning(APP_TITLE, message)
            return

        warning_parts = []
        if cancelled:
            warning_parts.append("Scan was cancelled. Results may be incomplete.")

        if total_sns != expected_sns:
            warning_parts.append(t("msg_expected_samples", expected=expected_sns, actual=total_sns))

        invalid = validate_grouped_logs_for_export(grouped_logs, expected_repeated)
        if invalid:
            warning_parts.append(t("msg_invalid_repeats", expected=expected_repeated) + "\n" + "\n".join(invalid))

        files_without_components = (scan_diagnostics or {}).get("files_without_components") or []
        if not available_components and files_without_components:
            warning_parts.append(
                "Serial numbers were found, but no components matched this profile's measurement rules.\n"
                "Check Name_Regex, Measure_Regex, and min/max regex settings in Manage Profiles."
            )

        if parse_warnings:
            warning_parts.append(t("msg_parse_warnings") + "\n" + "\n".join(parse_warnings[:8]))

        if warning_parts:
            messagebox.showwarning(APP_TITLE, "\n\n".join(warning_parts))
        else:
            messagebox.showinfo(APP_TITLE, t("msg_scan_success", samples=total_sns, files=total_files))

    def _on_scan_error(self, err_msg: str):
        logger.error("Scan error: %s", err_msg, exc_info=True)
        self.set_status(t("status_scan_failed"), loading=False)
        self.set_busy(False)
        messagebox.showerror(APP_TITLE, err_msg)

    def duplicate_components_for(self, components: Iterable[str]) -> List[str]:
        requested = set(components)
        duplicates = {
            group["component"]
            for group in getattr(self, "duplicate_component_choices", [])
            if group["component"] in requested
        }
        return sorted(duplicates, key=lambda item: item.lower())

    def show_duplicate_component_dialog(self, components: Optional[Iterable[str]] = None):
        choices = getattr(self, "duplicate_component_choices", [])
        if not choices:
            return

        component_filter = set(components) if components is not None else None
        if component_filter is not None:
            choices = [group for group in choices if group["component"] in component_filter]
            if not choices:
                return

        grouped_choices: Dict[str, dict] = {}
        for group in choices:
            component = group["component"]
            grouped = grouped_choices.setdefault(
                component,
                {
                    "component": component,
                    "row_label": group.get("row_label", "Rows"),
                    "options": group["options"],
                    "groups": [],
                },
            )
            grouped["groups"].append(group)

        component_choices = sorted(grouped_choices.values(), key=lambda item: item["component"].lower())
        dialog_components = {group["component"] for group in component_choices}

        top = tk.Toplevel(self)
        top.title("Duplicate Component Names")
        set_window_icon(top)
        top.geometry("820x460")
        top.minsize(720, 360)
        top.transient(self)
        top.grab_set()

        header = ttk.Frame(top, padding=(14, 12, 14, 6))
        header.pack(fill="x")
        ttk.Label(
            header,
            text="Duplicate component names were found. The first row/line is used by default.",
            font=("Segoe UI Semibold", 11),
        ).pack(anchor="w")
        ttk.Label(
            header,
            text="Each duplicate name is shown once. Your choice is applied to all files with the same structure.",
            foreground="#667085",
        ).pack(anchor="w", pady=(2, 0))

        canvas = tk.Canvas(top, highlightthickness=0, bg="#ffffff")
        scrollbar = ttk.Scrollbar(top, orient="vertical", command=canvas.yview)
        body = ttk.Frame(canvas, padding=(14, 8, 14, 8))
        body.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas_window = canvas.create_window((0, 0), window=body, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True, padx=(14, 0), pady=(0, 12))
        scrollbar.pack(side="right", fill="y", padx=(0, 14), pady=(0, 12))
        canvas.bind("<Configure>", lambda event: canvas.itemconfigure(canvas_window, width=event.width))

        body.columnconfigure(1, weight=1)
        body.columnconfigure(2, weight=2)
        ttk.Label(body, text="No.", font=("Segoe UI Semibold", 9)).grid(row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 6))
        ttk.Label(body, text="Duplicate Name", font=("Segoe UI Semibold", 9)).grid(row=0, column=1, sticky="w", padx=(0, 10), pady=(0, 6))
        ttk.Label(body, text="Use This Row / Value", font=("Segoe UI Semibold", 9)).grid(row=0, column=2, sticky="w", pady=(0, 6))

        selections = []
        for row_index, group in enumerate(component_choices, start=1):
            component = group["component"]
            options = group["options"]
            labels = [f"{idx}. {measurement_choice_label(option)}" for idx, option in enumerate(options, start=1)]
            var = tk.StringVar(value=labels[0])
            ttk.Label(body, text=str(row_index), anchor="w").grid(row=row_index, column=0, sticky="w", padx=(0, 10), pady=(0, 6))
            ttk.Label(body, text=component, anchor="w").grid(row=row_index, column=1, sticky="ew", padx=(0, 10), pady=(0, 6))
            combo = ttk.Combobox(body, textvariable=var, values=labels, state="readonly")
            combo.grid(row=row_index, column=2, sticky="ew", pady=(0, 6))
            combo.current(0)
            selections.append((group, combo))

        footer = ttk.Frame(top, padding=(14, 0, 14, 14))
        footer.pack(fill="x")

        def keep_defaults():
            self.reviewed_duplicate_components.update(dialog_components)
            self.set_status(f"Default duplicate choices kept for {len(dialog_components)} component name(s).")
            top.destroy()

        def apply_choices():
            for grouped, combo in selections:
                selected_index = combo.current()
                if selected_index < 0:
                    selected_index = 0

                component = grouped["component"]
                for group in grouped["groups"]:
                    item = group["item"]
                    options = group["options"]
                    option_index = selected_index if selected_index < len(options) else 0
                    item["entries"][component] = options[option_index]

            self.reviewed_duplicate_components.update(dialog_components)
            self.set_status(f"Duplicate choices applied for {len(selections)} component name(s).")
            top.destroy()

        top.protocol("WM_DELETE_WINDOW", keep_defaults)
        ttk.Button(footer, text="Keep Defaults", command=keep_defaults).pack(side="right")
        ttk.Button(footer, text="Apply Choices", style="Accent.TButton", command=apply_choices).pack(side="right", padx=(0, 8))

    def reset_workflow(self):
        self.folder_var.set("")
        self.grouped_logs = {}
        self.available_components = []
        self.selected_components = []
        self.parse_warnings = []
        self.duplicate_component_choices = []
        self.reviewed_duplicate_components = set()
        if hasattr(self, "search_var"):
            self.search_var.set("")
        if hasattr(self, "available_list"):
            self.available_list.delete(0, tk.END)
        if hasattr(self, "selected_list"):
            self.selected_list.delete(0, tk.END)
        if hasattr(self, "preview_tree"):
            self.preview_tree.delete(*self.preview_tree.get_children())
        self.detected_count_var.set(t("detected_count", count=0))
        self.selected_count_var.set(t("selected_count", count=0))
        self.set_status(t("status_ready"))

    def refresh_available_list(self):
        keyword = self.search_var.get().strip().lower() if hasattr(self, "search_var") else ""
        self.available_list.delete(0, tk.END)
        for component in self.available_components:
            if keyword and keyword not in component.lower():
                continue
            self.available_list.insert(tk.END, component)

    def refresh_selected_list(self):
        self.selected_list.delete(0, tk.END)
        for component in self.selected_components:
            self.selected_list.insert(tk.END, component)
        self.selected_count_var.set(t("selected_count", count=len(self.selected_components)))

    def add_components(self):
        selected = [self.available_list.get(index) for index in self.available_list.curselection()]
        newly_added = []
        for item in selected:
            if item not in self.selected_components:
                self.selected_components.append(item)
                newly_added.append(item)
        self.refresh_selected_list()

        duplicate_components = [
            component
            for component in self.duplicate_components_for(newly_added)
            if component not in self.reviewed_duplicate_components
        ]
        if duplicate_components:
            self.show_duplicate_component_dialog(duplicate_components)

    def remove_components(self):
        indices = list(self.selected_list.curselection())
        for index in reversed(indices):
            if 0 <= index < len(self.selected_components):
                self.selected_components.pop(index)
        self.refresh_selected_list()
        return "break"

    def clear_selected(self):
        self.selected_components = []
        self.refresh_selected_list()

    def move_selected(self, direction: int):
        indices = self.selected_list.curselection()
        if len(indices) != 1:
            return

        index = indices[0]
        new_index = index + direction
        if new_index < 0 or new_index >= len(self.selected_components):
            return

        item = self.selected_components.pop(index)
        self.selected_components.insert(new_index, item)
        self.refresh_selected_list()
        self.selected_list.selection_set(new_index)
        self.selected_list.activate(new_index)

    def preview_summary(self):
        self.preview_tree.delete(*self.preview_tree.get_children())
        if not self.grouped_logs:
            messagebox.showwarning(APP_TITLE, t("msg_scan_first"))
            return

        for idx, sn in enumerate(sorted(self.grouped_logs.keys()), start=1):
            files = self.grouped_logs[sn]
            file_names = ", ".join(item["filename"] for item in files)
            self.preview_tree.insert("", tk.END, values=(idx, sn, len(files), file_names))

        self.set_status(t("status_preview_loaded"))

    def export_excel(self):
        if not self.grouped_logs:
            messagebox.showwarning(APP_TITLE, t("msg_scan_first"))
            return

        if not self.selected_components:
            messagebox.showwarning(APP_TITLE, t("msg_select_component"))
            return

        expected_sns = self.get_expected_samples()
        expected_repeated = self.get_repeated_times()
        if expected_sns is None or expected_repeated is None:
            return

        total_sns = len(self.grouped_logs)
        if total_sns != expected_sns:
            if not messagebox.askyesno(APP_TITLE, t("msg_export_mismatch_confirm", expected=expected_sns, actual=total_sns)):
                return

        invalid = validate_grouped_logs_for_export(self.grouped_logs, expected_repeated)
        if invalid:
            messagebox.showerror(APP_TITLE, t("msg_invalid_repeats", expected=expected_repeated) + "\n" + "\n".join(invalid))
            return

        output = filedialog.asksaveasfilename(
            title="Save output Excel",
            defaultextension=".xlsx",
            initialfile=DEFAULT_OUTPUT_FILE,
            filetypes=[("Excel file", "*.xlsx")],
        )
        if not output:
            return

        output_path = Path(output)

        # Snapshot selected components for the export thread
        components_snapshot = list(self.selected_components)
        grouped_snapshot = dict(self.grouped_logs)

        self.set_busy(True)
        self.set_status(t("status_generating"), loading=True)

        def task():
            try:
                generate_excel_from_grouped(grouped_snapshot, components_snapshot, output_path, expected_repeated)
                self.after(0, self._on_export_success, str(output_path))
            except Exception as exc:
                self.after(0, self._on_export_error, str(exc))

        threading.Thread(target=task, daemon=True).start()

    def _on_export_success(self, output: str):
        self.set_status(t("status_done", output=output), loading=False)
        self.set_busy(False)

        if messagebox.askyesno(APP_TITLE, t("msg_export_open_folder")):
            open_folder_in_explorer(Path(output).parent)

    def _on_export_error(self, err_msg: str):
        logger.error("Export error: %s", err_msg, exc_info=True)
        self.set_status(t("status_export_failed"), loading=False)
        self.set_busy(False)
        messagebox.showerror(APP_TITLE, err_msg)


if __name__ == "__main__":
    configure_process_dpi_awareness()
    app = MasterApp()
    app.mainloop()
