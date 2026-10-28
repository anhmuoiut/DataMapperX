"""Working out a profile from a sample, previewing it, and checking it against a folder."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
from openpyxl import Workbook

from datamapperx.parsing import parse_log_file, read_text_file
from datamapperx.profiles import DEFAULT_PROFILES, patterns_from_profile, text_profile_from_patterns
from datamapperx.sampling import (
    I3070_PATTERNS,
    WorkbookSketch,
    annotate_text,
    check_folder,
    detect_excel,
    detect_profile,
    detect_text,
    guess_serial_example,
    preview_profile,
    sibling_files,
    suggest_profile,
)
from datamapperx.scanner import scan_folder

from .conftest import ict_block, write_ict_log


# ----------------------------------------------------------------------------- builders
def write_results_workbook(
    path: Path,
    sn: str,
    rows: list[tuple],
    *,
    labelled_sn: bool = True,
    extra_graph_sheets: int = 0,
) -> Path:
    """A tester workbook: a Session sheet with the serial number and a Results sheet with a headed table."""
    wb = Workbook()
    session = wb.active
    session.title = "Session"
    session["A1"] = "Serial number" if labelled_sn else "Unit"
    session["B1"] = sn
    session["C1"] = "Serial number 2"
    session["D1"] = sn[-6:]
    session["A2"] = "Start"
    session["A3"] = "Family"
    session["B3"] = "Inverter EOL"
    results = wb.create_sheet("Results")
    results.append(["#", "Id", "Test", "Type", "Min", "Result", "Max", "Units", "Status", "Duration"])
    for number, (test, result, low, high, unit) in enumerate(rows, start=1):
        results.append([number, 1000 + number if result is not None else None, test, "MinMax", low, result, high, unit,
                        "Pass" if result is not None else None, "00:00:01"])
    for n in range(extra_graph_sheets):
        graph = wb.create_sheet(f"Graph {n + 1}")
        graph.append(["Test", f"Dump {n}"])
        graph.append([])
        graph.append(["#", "X", "0 - "])
        for i in range(1, 30):
            graph.append([i, i - 1, i * 3])
    wb.save(path)
    return path


EOL_ROWS = [
    ("INITIALIZATION (Init)", None, None, None, None),
    ("Init, Test SIM SN", "1662036166", None, None, None),
    ("Setup Golden Inverter", None, None, None, None),
    ("Config, Get Portia SW Version", "6.25.35", None, None, None),
    ("DC bus voltage", 399.5, 380, 420, "V"),
]


def make_ict_folder(folder: Path, serials=("VSE2C00000112", "VSE2C00000113"), repeats: int = 2) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for s_index, sn in enumerate(serials):
        for r in range(repeats):
            blocks = [
                ict_block("1%r114", "+4.26E+04", "+9.99E+99", "+9.10E+03", kind="JUM"),
                ict_block("c23", "+1.0E-07", "+1.1E-07", "+9.0E-08", kind="CAP"),
                ict_block("u1107", "+1.68E+00|Output_1_high", None, None, kind="MEA"),
            ]
            name = f"AP1020K-DG-08-FW4_F_{sn}_26090816{s_index}{r}50_1_P.log"
            paths.append(write_ict_log(folder / name, sn, blocks))
    return paths


# ----------------------------------------------------------------------------- Excel detection
def test_excel_layout_and_serial_cell_are_found_from_the_headings(tmp_path):
    path = write_results_workbook(tmp_path / "SV3826-06310A8C6-E1 1.xlsx", "SV3826-06310A8C6-E1", EOL_ROWS, extra_graph_sheets=2)
    detection = detect_excel(path)
    profile = detection.profile

    assert profile["data_sheet"] == "Results" and profile["sn_sheet"] == "Session"
    assert profile["sn_regex"] == "B1"  # the cell whose text is in the file name; the shorter "Serial number 2" loses
    assert (profile["row_start"], profile["component_col"], profile["measure_col"]) == ("2", "C", "F")
    assert (profile["min_col"], profile["max_col"], profile["unit_col"]) == ("E", "G", "H")
    assert detection.suggested_name == tmp_path.name
    assert any("cell B1" in note for note in detection.notes) and not detection.warnings


def test_a_detected_excel_profile_reads_the_sample_and_skips_section_headings(tmp_path):
    path = write_results_workbook(tmp_path / "SV3826-06310A8C6-E1 1.xlsx", "SV3826-06310A8C6-E1", EOL_ROWS)
    parsed = parse_log_file(path, detect_excel(path).profile)

    assert parsed.sn == "SV3826-06310A8C6-E1"
    assert list(parsed.entries) == ["Init, Test SIM SN", "Config, Get Portia SW Version", "DC bus voltage"]
    voltage = parsed.entries["DC bus voltage"]
    assert (voltage.measured, voltage.min_value, voltage.max_value, voltage.unit) == ("399.5", "380", "420", "V")


def test_section_headings_can_be_kept_when_a_profile_asks_for_them(tmp_path):
    path = write_results_workbook(tmp_path / "SV3826-06310A8C6-E1 1.xlsx", "SV3826-06310A8C6-E1", EOL_ROWS)
    profile = {**detect_excel(path).profile, "skip_empty_rows": False}
    assert "INITIALIZATION (Init)" in parse_log_file(path, profile).entries


def test_serial_is_found_next_to_its_label_when_the_file_name_does_not_contain_it(tmp_path):
    path = write_results_workbook(tmp_path / "report 7.xlsx", "ZX9981-AB", EOL_ROWS)
    detection = detect_excel(path)
    assert detection.profile["sn_regex"] == "B1"
    assert parse_log_file(path, detection.profile).sn == "ZX9981-AB"


def test_a_workbook_without_a_recognisable_table_asks_for_help(tmp_path):
    wb = Workbook()
    wb.active.append(["alpha", "beta"])
    wb.save(tmp_path / "odd 1.xlsx")
    detection = detect_excel(tmp_path / "odd 1.xlsx")
    assert any("could not find a table" in w for w in detection.warnings)


def test_sheet_sketch_lists_columns_with_their_headings_and_examples(tmp_path):
    path = write_results_workbook(tmp_path / "SV3826-06310A8C6-E1 1.xlsx", "SV3826-06310A8C6-E1", EOL_ROWS)
    sketch = WorkbookSketch.load(path)
    assert sketch.names == ["Session", "Results"]
    assert sketch.cell("Session", "B1") == "SV3826-06310A8C6-E1" and sketch.cell("Session", "Z99") == ""
    labels = {c.letter: c.label for c in sketch.columns("Results", 2)}
    assert labels["C"] == "C  -  Test" and labels["F"] == "F  -  Result"
    headingless = {c.letter: c.label for c in sketch.columns("Results", 1)}
    assert headingless["C"].startswith("C  -  (e.g. ")


# ----------------------------------------------------------------------------- text detection
def test_i3070_logs_are_recognised_and_read_exactly_like_the_built_in_profile(tmp_path):
    files = make_ict_folder(tmp_path / "ict")
    detection = detect_profile(files[0])
    profile = detection.profile

    assert patterns_from_profile(profile) == (I3070_PATTERNS, [])
    assert profile["sn_example"] == "VSE2C00000112"  # differs between files and is written inside them
    assert any("i3070" in note for note in detection.notes) and not detection.warnings

    built_in = copy.deepcopy(DEFAULT_PROFILES["ICT_I3070 Test Log"])
    for path in files:
        mine, theirs = parse_log_file(path, profile), parse_log_file(path, built_in)
        assert mine.sn == theirs.sn
        assert mine.entries == theirs.entries and mine.entries


def test_an_unknown_text_format_is_reported_instead_of_guessed(tmp_path):
    path = tmp_path / "SN12345_run.txt"
    path.write_text("R1 = 5.0 ohm\nR2 = 7.1 ohm\n", encoding="utf-8")
    detection = detect_text(path)
    assert detection.profile["block_regex"] == "" and detection.profile["measure_regex"] == ""
    assert any("do not know this log format" in w for w in detection.warnings)


def test_serial_guess_prefers_the_part_that_varies_and_is_written_inside_the_file(tmp_path):
    files = make_ict_folder(tmp_path / "ict")
    text = read_text_file(files[0])
    assert guess_serial_example(files[0], sibling_files(files[0]), text) == "VSE2C00000112"
    assert guess_serial_example(files[0], (), text) == "VSE2C00000112"  # written inside the file, so still preferred
    # With nothing to compare against, the longest mixed token is a weak guess, but better than nothing.
    assert guess_serial_example(files[0], ()) == "AP1020K-DG-08-FW4"
    assert guess_serial_example(tmp_path / "ab.log") == ""


def test_siblings_are_the_other_files_of_the_same_kind(tmp_path):
    folder = tmp_path / "f"
    folder.mkdir()
    for name in ("a.log", "b.log", "c.txt", "~$d.log"):
        (folder / name).write_text("x", encoding="utf-8")
    assert [p.name for p in sibling_files(folder / "a.log")] == ["b.log"]


# ----------------------------------------------------------------------------- previews and highlights
def test_preview_reports_values_or_the_reason_nothing_was_read(tmp_path):
    files = make_ict_folder(tmp_path / "ict")
    profile = detect_profile(files[0]).profile
    preview = preview_profile(files[0], profile, limit=2)
    assert preview.sn == "VSE2C00000112" and preview.total == 3 and len(preview.rows) == 2
    assert preview.rows[0].component == "1%r114" and preview.rows[0].min_value == "+9.10E+03"

    broken = {**profile, "measure_regex": "(unclosed"}
    assert "Invalid Regex" in preview_profile(files[0], broken).error
    assert preview_profile(tmp_path / "missing.log", profile).error


def test_annotations_mark_exactly_the_text_the_profile_reads(tmp_path):
    files = make_ict_folder(tmp_path / "ict", serials=("VSE2C00000112",), repeats=1)
    text = read_text_file(files[0])
    lines = text.splitlines()
    marks = annotate_text(text, detect_profile(files[0]).profile)

    def covered(role: str) -> list[str]:
        return [lines[m.line][m.start : m.end] for m in marks if m.role == role]

    assert covered("serial") == ["VSE2C00000112"]
    assert covered("component") == ["1%r114", "c23", "u1107"]
    assert covered("measured") == ["+4.26E+04", "+1.0E-07", "+1.68E+00"]
    assert covered("unit") == ["JUM", "CAP", "MEA"]
    assert covered("max") == ["+9.99E+99", "+1.1E-07"] and covered("min") == ["+9.10E+03", "+9.0E-08"]
    assert {m.role for m in marks} >= {"block", "measure"}


def test_annotating_with_a_broken_pattern_raises_a_clear_error():
    with pytest.raises(ValueError, match="Invalid Regex"):
        annotate_text("x", {"type": "text", "block_regex": "(", "measure_regex": "x"})


# ----------------------------------------------------------------------------- whole folders
def test_folder_check_counts_files_that_can_and_cannot_be_read(tmp_path):
    files = make_ict_folder(tmp_path / "ict")
    (tmp_path / "ict" / "VSE2C00000199_empty.log").write_text("nothing here", encoding="utf-8")
    result = check_folder(tmp_path / "ict", detect_profile(files[0]).profile)
    assert (result.tried, result.readable) == (5, 4)
    assert not result.all_ok and "VSE2C00000199_empty.log: no components found" in result.problems


def test_the_profile_that_reads_a_folder_is_suggested(tmp_path):
    make_ict_folder(tmp_path / "ict")
    (tmp_path / "eol").mkdir()
    eol_file = write_results_workbook(tmp_path / "eol" / "SV3826-06310A8C6-E1 1.xlsx", "SV3826-06310A8C6-E1", EOL_ROWS)
    eol_profile = detect_excel(eol_file).profile
    profiles = {n: p for n, p in copy.deepcopy(DEFAULT_PROFILES).items() if n != "EOL Test Log"} | {"EOL": eol_profile}

    assert suggest_profile(tmp_path / "ict", profiles) == "ICT_I3070 Test Log"
    assert suggest_profile(tmp_path / "eol", profiles) == "EOL"
    (tmp_path / "other").mkdir()
    (tmp_path / "other" / "x.log").write_text("nothing", encoding="utf-8")
    assert suggest_profile(tmp_path / "other", profiles) is None


def test_a_profile_built_from_patterns_scans_a_folder(tmp_path):
    make_ict_folder(tmp_path / "ict")
    profile = text_profile_from_patterns(I3070_PATTERNS, "VSE2C00000112", ".log")
    result = scan_folder(tmp_path / "ict", profile)
    assert sorted(result.grouped) == ["VSE2C00000112", "VSE2C00000113"]
    assert sorted(result.components) == ["1%r114", "c23", "u1107"]


# ----------------------------------------------------------------------------- the real sample logs (local only)
SAMPLES = Path(__file__).resolve().parent.parent / "samples"
needs_samples = pytest.mark.skipif(not SAMPLES.is_dir(), reason="the samples folder holds company data and is not published")


def first_file(family: str, pattern: str) -> Path:
    return sorted((SAMPLES / family).glob(pattern))[0]


@needs_samples
def test_real_ict_logs_are_detected_and_match_the_built_in_profile():
    detection = detect_profile(first_file("ICT logfile", "*.log"))
    built_in = DEFAULT_PROFILES["ICT_I3070 Test Log"]
    mine, theirs = scan_folder(SAMPLES / "ICT logfile", detection.profile), scan_folder(SAMPLES / "ICT logfile", built_in)
    assert not detection.warnings
    assert mine.components == theirs.components and len(mine.components) == 397
    for sn, records in theirs.grouped.items():
        assert [r.entries for r in mine.grouped[sn]] == [r.entries for r in records]


@needs_samples
def test_real_function_test_logs_are_detected():
    detection = detect_profile(first_file("Function test logfile", "*.xlsx"))
    result = scan_folder(SAMPLES / "Function test logfile", detection.profile)
    assert not detection.warnings
    assert len(result.grouped) == 10 and not result.diagnostics.skipped_without_sn
    assert len(result.components) == 104  # 114 names minus 10 section headings


@needs_samples
def test_real_eol_logs_that_no_built_in_profile_could_read_are_detected():
    detection = detect_profile(first_file("EOL logfile", "*.xlsx"))
    assert detection.profile["sn_regex"] == "B1" and not detection.warnings
    result = scan_folder(SAMPLES / "EOL logfile", detection.profile)
    assert len(result.grouped) == 10 and not result.diagnostics.skipped_without_sn
    assert result.components and "INITIALIZATION (Init)" not in result.components
    assert check_folder(SAMPLES / "EOL logfile", detection.profile).all_ok
