"""The fast workbook reader must return exactly what openpyxl returns, and fall back when it cannot cope."""

from __future__ import annotations

import datetime as dt
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_timedelta_format
from openpyxl.styles.numbers import is_date_format as openpyxl_is_date

from datamapperx import xlsx
from datamapperx.parsing import parse_excel_file
from datamapperx.profiles import DEFAULT_PROFILES
from datamapperx.xlsx import OpenpyxlReader, UnsupportedWorkbook, XlsxReader, read_workbook

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def openpyxl_rows(path: Path, sheet: str) -> dict[int, list]:
    """What openpyxl sees: {row number: values with trailing empties removed}, empty rows left out."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet]
        ws.reset_dimensions()
        rows = {}
        for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(values)
            while values and values[-1] is None:
                values.pop()
            if values:
                rows[number] = values
        return rows
    finally:
        wb.close()


def fast_rows(path: Path, sheet: str) -> dict[int, list]:
    with XlsxReader(path) as reader:
        rows = {}
        for number, values in reader.iter_rows(sheet):
            values = list(values)
            while values and values[-1] is None:
                values.pop()
            if values:
                rows[number] = values
        return rows


def trimmed(values: list) -> list:
    values = list(values)
    while values and values[-1] is None:
        values.pop()
    return values


def assert_same_as_openpyxl(path: Path) -> None:
    wb = load_workbook(path, read_only=True)
    names = [ws.title for ws in wb.worksheets]
    wb.close()
    with XlsxReader(path) as reader:
        assert reader.sheet_names == names
    for name in names:
        assert fast_rows(path, name) == openpyxl_rows(path, name), f"{path.name}: sheet {name!r} differs"


# ----------------------------------------------------------------------------- written by openpyxl
def make_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mixed & <odd> name"
    ws.append(["text", 1, 2.5, True, False, None, "é ü 漢字", "  padded  ", "=1+1"])
    ws["A3"] = "gap above, row 2 is missing"
    ws["E5"] = 1e-7
    ws["F5"] = 123456789012345678
    ws["G5"] = -0.0
    ws["H5"] = 1.5e300
    ws["B7"] = dt.datetime(2026, 10, 13, 10, 33, 12)
    ws["C7"] = dt.date(2020, 2, 29)
    ws["D7"] = dt.time(23, 59, 1)
    ws["E7"] = dt.timedelta(hours=30, minutes=5)
    ws["F7"] = dt.datetime(1900, 1, 1)
    ws["G7"] = 0.5
    ws["H7"] = "#N/A"
    ws["A9"] = "multi\nline"
    ws.merge_cells("A11:C11")
    ws["A11"] = "merged"
    ws.cell(row=13, column=40, value="far right")
    other = wb.create_sheet("Second")
    other["A1"] = "only here"
    other.sheet_state = "hidden"
    wb.create_sheet("Empty")
    wb.save(path)
    return path


def test_every_kind_of_cell_matches_openpyxl(tmp_path):
    assert_same_as_openpyxl(make_workbook(tmp_path / "mixed.xlsx"))


def test_the_1904_date_system_is_respected(tmp_path):
    wb = Workbook()
    wb.epoch = dt.datetime(1904, 1, 1)
    wb.active["A1"] = dt.datetime(2026, 10, 13, 10, 33, 12)
    wb.active["A2"] = dt.datetime(1904, 1, 2)
    wb.save(tmp_path / "mac.xlsx")
    assert_same_as_openpyxl(tmp_path / "mac.xlsx")
    assert fast_rows(tmp_path / "mac.xlsx", "Sheet")[1] == [dt.datetime(2026, 10, 13, 10, 33, 12)]


def test_limits_on_rows_and_columns_match_openpyxl(tmp_path):
    path = make_workbook(tmp_path / "mixed.xlsx")
    name = "Mixed & <odd> name"
    with XlsxReader(path) as reader:
        assert [n for n, _ in reader.iter_rows(name, min_row=3, max_row=7)] == [3, 5, 7]
        assert reader.grid(name, 4, 3) == [["text", 1, 2.5], [], ["gap above, row 2 is missing"]]
        assert reader.cell(name, "C1") == 2.5 and reader.cell(name, "b1") == 1
        assert reader.cell(name, "ZZ999") is None and reader.cell(name, "A0") is None and reader.cell(name, "nonsense") is None
        assert reader.cell(name, "AN13") == "far right"
    with OpenpyxlReader(path) as slow, XlsxReader(path) as fast:
        for sheet in fast.sheet_names:
            # openpyxl pads every row to the width asked for; the fast reader leaves out the trailing empties.
            assert [trimmed(r) for r in fast.grid(sheet, 20, 12)] == [trimmed(r) for r in slow.grid(sheet, 20, 12)], sheet


# ----------------------------------------------------------------------------- hand-made XML that openpyxl never writes
def raw_workbook(path: Path, sheet_xml: str, shared: str | None = None, styles: str | None = None, workbook_pr: str = "") -> Path:
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            + (
                '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
                if shared else ""
            )
            + (
                '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                if styles else ""
            )
            + "</Types>"
        ),
        "_rels/.rels": f'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<?xml version="1.0"?><workbook xmlns="{main}" xmlns:r="{rel}">{workbook_pr}<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
            + (f'<Relationship Id="rId2" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/>' if shared else "")
            + (f'<Relationship Id="rId3" Type="{rel}/styles" Target="styles.xml"/>' if styles else "")
            + "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0"?><worksheet xmlns="{main}"><sheetData>{sheet_xml}</sheetData></worksheet>',
    }
    if shared:
        files["xl/sharedStrings.xml"] = f'<?xml version="1.0"?><sst xmlns="{main}">{shared}</sst>'
    if styles:
        files["xl/styles.xml"] = f'<?xml version="1.0"?><styleSheet xmlns="{main}">{styles}</styleSheet>'
    with zipfile.ZipFile(path, "w") as z:
        for name, content in files.items():
            z.writestr(name, content)
    return path


def test_inline_rich_and_escaped_strings_are_read_like_openpyxl(tmp_path):
    shared = (
        "<si><t>plain</t></si>"
        '<si><r><t>rich </t></r><r><rPr><b/></rPr><t>text</t></r></si>'
        '<si><t xml:space="preserve">  keeps spaces  </t></si>'
        "<si><t>x005F_hidden</t></si>"
        '<si><t>name</t><rPh sb="0" eb="1"><t>ignored phonetic</t></rPh></si>'
    )
    rows = (
        '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c><c r="D1" t="s"><v>3</v></c><c r="E1" t="s"><v>4</v></c></row>'
        '<row r="2"><c r="A2" t="inlineStr"><is><t>inline</t></is></c><c r="B2" t="inlineStr"><is><r><t>in</t></r><r><t>line rich</t></r></is></c>'
        '<c r="C2" t="str"><v>formula text</v></c><c r="D2" t="e"><v>#DIV/0!</v></c><c r="E2" t="b"><v>1</v></c></row>'
        '<row><c t="s"><v>0</v></c><c><v>7</v></c><c r="E3"><v>1.5E+3</v></c></row>'
        '<row r="9"><c r="A9"><v></v></c><c r="B9" t="inlineStr"/></row>'
    )
    path = raw_workbook(tmp_path / "raw.xlsx", rows, shared=shared)
    assert fast_rows(path, "Data") == openpyxl_rows(path, "Data")
    assert fast_rows(path, "Data")[1][1] == "rich text" and fast_rows(path, "Data")[1][3] == "hidden"


def test_custom_date_formats_and_durations_match_openpyxl(tmp_path):
    styles = (
        '<numFmts count="3"><numFmt numFmtId="164" formatCode="yyyy\\-mm\\-dd\\ hh:mm:ss"/>'
        '<numFmt numFmtId="165" formatCode="[h]:mm:ss"/><numFmt numFmtId="166" formatCode="0.00&quot;d&quot;"/></numFmts>'
        '<cellXfs count="5"><xf numFmtId="0"/><xf numFmtId="164"/><xf numFmtId="165"/><xf numFmtId="166"/><xf numFmtId="14"/></cellXfs>'
    )
    rows = (
        '<row r="1"><c r="A1" s="1"><v>46308.43958333333</v></c><c r="B1" s="2"><v>1.25</v></c>'
        '<c r="C1" s="3"><v>3.5</v></c><c r="D1" s="4"><v>43831</v></c><c r="E1" s="1"><v>1e20</v></c><c r="F1" s="0"><v>5</v></c></row>'
    )
    path = raw_workbook(tmp_path / "styles.xlsx", rows, styles=styles)
    assert fast_rows(path, "Data") == openpyxl_rows(path, "Data")
    assert fast_rows(path, "Data")[1][4] == "#VALUE!"  # a date outside the supported range, as openpyxl reports


def test_the_mac_date_system_in_a_raw_file(tmp_path):
    styles = '<cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs>'
    rows = '<row r="1"><c r="A1" s="1"><v>1000.5</v></c></row>'
    path = raw_workbook(tmp_path / "mac.xlsx", rows, styles=styles, workbook_pr='<workbookPr date1904="1"/>')
    assert fast_rows(path, "Data") == openpyxl_rows(path, "Data")


@pytest.mark.parametrize(
    "code",
    [
        "General", "0.00", "0.00E+00", "#,##0", "yyyy-mm-dd", "dd/mm/yyyy hh:mm", "h:mm:ss AM/PM", "[h]:mm:ss", "[mm]:ss",
        "[Red]0.00", '0.00"days"', '"day" 0', "0\\d", "_(* #,##0_)", "mmm-yy;@", "@", "", "[$-409]d-mmm", "0.0%",
        "[hh]:mm", "[ss].00", "hh:mm;[Red]hh:mm", "\\y0",
    ],
)
def test_date_format_rules_equal_openpyxl(code):
    assert xlsx.is_date_format(code) == openpyxl_is_date(code)
    assert xlsx.is_timedelta_format(code) == bool(is_timedelta_format(code))


def test_the_builtin_date_formats_are_the_ones_openpyxl_treats_as_dates():
    expected = {i for i, code in BUILTIN_FORMATS.items() if openpyxl_is_date(code)}
    assert set(xlsx._BUILTIN_DATE_FORMATS) == expected


# ----------------------------------------------------------------------------- falling back
def test_iso_date_cells_make_the_fast_reader_step_aside_and_openpyxl_takes_over(tmp_path):
    rows = '<row r="1"><c r="A1" t="d"><v>2026-10-13T10:33:12</v></c><c r="B1"><v>5</v></c></row>'
    path = raw_workbook(tmp_path / "iso.xlsx", rows)
    with XlsxReader(path) as reader, pytest.raises(UnsupportedWorkbook):
        list(reader.iter_rows("Data"))
    assert read_workbook(path, lambda reader: [values for _n, values in reader.iter_rows("Data")]) == [[dt.datetime(2026, 10, 13, 10, 33, 12), 5]]


def test_a_file_both_readers_reject_raises_the_error_after_trying_openpyxl(tmp_path):
    # A cell that points past the end of the shared strings: the fast reader gives up and openpyxl cannot read it either.
    path = raw_workbook(tmp_path / "bad.xlsx", '<row r="1"><c r="A1" t="s"><v>99</v></c></row>', shared="<si><t>only one</t></si>")
    with pytest.raises(UnsupportedWorkbook):
        fast_rows(path, "Data")
    used = []

    def work(reader):
        used.append(type(reader).__name__)
        return list(reader.iter_rows("Data"))

    with pytest.raises(IndexError):
        read_workbook(path, work)
    assert used == ["XlsxReader", "OpenpyxlReader"]


def test_not_a_workbook_and_locked_files_are_reported_not_swallowed(tmp_path, monkeypatch):
    junk = tmp_path / "junk.xlsx"
    junk.write_text("this is not a zip file", encoding="utf-8")
    with pytest.raises(zipfile.BadZipFile):
        read_workbook(junk, lambda reader: None)
    with pytest.raises(ValueError, match="not a valid Excel workbook"):
        parse_excel_file(junk, DEFAULT_PROFILES["Function Test Log"])

    def locked(*_args, **_kwargs):
        raise PermissionError("in use")

    monkeypatch.setattr(xlsx, "ZipFile", locked)
    with pytest.raises(PermissionError, match="CLOSE the file in Excel"):
        parse_excel_file(tmp_path / "anything.xlsx", DEFAULT_PROFILES["Function Test Log"])


def test_a_file_with_no_worksheets_is_explained(tmp_path):
    path = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets/></workbook>')
    assert read_workbook(path, lambda reader: reader.sheet_names) == []
    with pytest.raises(ValueError, match="has no worksheets"):
        parse_excel_file(path, DEFAULT_PROFILES["Function Test Log"])


# ----------------------------------------------------------------------------- start-up cost
def test_opening_the_window_does_not_import_the_heavy_libraries():
    code = (
        "import sys, datamapperx.ui.main_window, datamapperx.parsing, datamapperx.sampling, datamapperx.exporter;"
        "heavy = [m for m in ('openpyxl', 'numpy', 'urllib.request', 'http.client', 'ssl') if m in sys.modules];"
        "print(','.join(heavy))"
    )
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=Path(__file__).resolve().parent.parent, check=False)
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "", f"imported at start-up: {out.stdout.strip()}"


def test_the_export_still_works_after_a_lazy_import(tmp_path):
    from datamapperx.exporter import generate_excel
    from datamapperx.models import LogRecord, Measurement

    record = LogRecord(path=tmp_path / "a.log", sn="S1", entries={"R1": Measurement("R1", "1.5")})
    generate_excel({"S1": [record]}, ["R1"], tmp_path / "out.xlsx", 1)
    assert load_workbook(tmp_path / "out.xlsx")["Raw Data"]["C2"].value == 1.5


# ----------------------------------------------------------------------------- the real logs (local only)
needs_samples = pytest.mark.skipif(not SAMPLES.is_dir(), reason="the samples folder holds company data and is not published")


@needs_samples
@pytest.mark.parametrize("family", ["Function test logfile", "EOL logfile"])
def test_every_real_workbook_is_read_cell_for_cell_like_openpyxl(family):
    for path in sorted((SAMPLES / family).glob("*.xlsx")):
        assert_same_as_openpyxl(path)


@needs_samples
def test_real_workbooks_are_parsed_identically_by_both_readers():
    from datamapperx.parsing import _ExcelSettings, _read_excel

    for family, profile_name in (("Function test logfile", "Function Test Log"), ("EOL logfile", "EOL Test Log")):
        settings = _ExcelSettings.from_profile(DEFAULT_PROFILES[profile_name])
        for path in sorted((SAMPLES / family).glob("*.xlsx")):
            with XlsxReader(path) as fast_reader, OpenpyxlReader(path) as slow_reader:
                fast, slow = _read_excel(fast_reader, path, settings), _read_excel(slow_reader, path, settings)
            assert fast == slow and fast.sn and fast.entries, path.name


# ----------------------------------------------------------------------------- damaged files
def test_a_folder_with_damaged_workbooks_is_reported_file_by_file_and_never_fails(tmp_path):
    from datamapperx.scanner import scan_folder

    good = make_workbook(tmp_path / "VSE001 1.xlsx")
    raw = good.read_bytes()
    damaged = []
    for cut in (0, 10, len(raw) // 2, len(raw) - 1):
        path = tmp_path / f"VSE00{cut} cut.xlsx"
        path.write_bytes(raw[:cut])
        damaged.append(path.name)
    (tmp_path / "VSE777 noise.xlsx").write_bytes(bytes(range(256)) * 20)
    damaged.append("VSE777 noise.xlsx")

    result = scan_folder(tmp_path, dict(DEFAULT_PROFILES["Function Test Log"]))  # must not raise

    warned = {warning.split(":")[0] for warning in result.parse_warnings}
    assert set(damaged) <= warned and "VSE001 1.xlsx" not in warned
    assert all("is not a valid Excel workbook" in warning for warning in result.parse_warnings)  # said in plain words


def test_odd_text_logs_are_read_or_skipped_without_failing(tmp_path):
    from datamapperx.scanner import scan_folder

    (tmp_path / "VSE001_noise.log").write_bytes(bytes(range(256)) * 80)
    (tmp_path / "VSE002_wide.log").write_bytes(("{@BLOCK|" + "x" * 2_000_000 + "|00\n{@A-RES|0|+1.0E+00\n").encode())
    (tmp_path / "VSE003_repeated.log").write_bytes(b"{@BLOCK|" * 50_000)
    (tmp_path / "VSE004_utf16.log").write_bytes(b"\xff\xfe" + "{@BLOCK|R1|00\n{@A-RES|0|+1.0E+00\n".encode("utf-16-le"))
    result = scan_folder(tmp_path, dict(DEFAULT_PROFILES["ICT_I3070 Test Log"]))
    assert not result.parse_warnings and result.total_files == 4
    assert "R1" in result.components  # the UTF-16 file was read properly
