"""What the workspace gives a report: the study information that was entered, the items chosen, and nothing stale."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from datamapperx.stat_reports import ALL_SECTIONS, DEFAULT_SECTIONS, export_statistical_report, report_context
from datamapperx.statistics import (
    Dataset,
    Observation,
    StaleResultError,
    StatisticalOptions,
    Study,
    StudyInfo,
)


def grr_study(**kwargs):
    rows = [Observation(item, f"S{part}", value + shift, 0, 40, source=f"{item}-{part}-{trial}")
            for item, shift in (("Alpha", 0), ("Beta", 3))
            for part, values in enumerate(((10, 11, 10), (14, 15, 14), (18, 17, 18), (22, 23, 22)), 1)
            for trial, value in enumerate(values, 1)]
    dataset = Dataset.from_observations(rows, description="RawData workbook demo.xlsx")
    return Study(dataset, options=StatisticalOptions(method="grr", assign_trials=True), **kwargs)


def capability_study(**kwargs):
    rows = [Observation("Cap", f"S{i}", 50 + (i % 7) * 0.4, 40, 60, source=f"c{i}") for i in range(40)]
    return Study(Dataset.from_observations(rows, description="RawData workbook cap.xlsx"),
                 options=StatisticalOptions(method="capability"), **kwargs)


INFO = StudyInfo(gage_name="Bench DMM 7", date_of_study="2026-10-30", reported_by="T. Man", misc="first article")


def prepared(study):
    result_set = study.run_statistics()
    return result_set, study.evaluate(result_set, judge=False)


def test_the_context_carries_only_the_study_information_that_was_entered():
    study = grr_study(info=StudyInfo(gage_name="Bench DMM 7"))
    result_set, _ = prepared(study)
    context = report_context(study, result_set, created="2026-10-31")
    assert context.info == (("Gage name", "Bench DMM 7"),) and context.created == "2026-10-31" and context.sections == DEFAULT_SECTIONS
    full = grr_study(info=INFO)
    assert dict(report_context(full, prepared(full)[0]).info) == {"Gage name": "Bench DMM 7", "Date of study": "2026-10-30",
                                                                  "Reported by": "T. Man", "Misc": "first article"}
    nothing = grr_study()
    assert report_context(nothing, prepared(nothing)[0]).info == ()


def test_the_report_records_the_day_it_was_made_without_being_told():
    study = grr_study()
    context = report_context(study, prepared(study)[0])
    assert len(context.created) == 10 and context.created[4] == "-"


def test_a_partial_report_counts_the_items_it_includes():
    study = grr_study()
    result_set, _ = prepared(study)
    partial = report_context(study, result_set, included=["Alpha"])
    assert (partial.included, partial.total, partial.partial) == (1, 2, True)
    assert not report_context(study, result_set).partial
    assert report_context(study, result_set, sections=("observations", "charts", "unknown")).sections == ALL_SECTIONS


def test_a_stale_result_set_cannot_be_reported():
    study = grr_study()
    result_set, _ = prepared(study)
    study.design["operators"] = ["A", "B"]
    with pytest.raises(StaleResultError):
        report_context(study, result_set)


def test_the_workspace_path_gives_a_short_excel_report_for_the_chosen_items(tmp_path):
    study = grr_study(info=INFO)
    result_set, evaluated = prepared(study)
    context = report_context(study, result_set, included=["Alpha"], sections=ALL_SECTIONS, created="2026-10-31")
    results = [r for r in evaluated.results if r.label == "Alpha"]
    book = load_workbook(export_statistical_report(tmp_path / "partial.xlsx", results, context=context))
    assert book.sheetnames == ["Alpha", "Plot data 1", "Observations"]
    labels = {str(row[0].value): row[1].value for row in book["Alpha"].iter_rows(min_col=1, max_col=2) if row[0].value}
    assert labels["Gage name:"] == "Bench DMM 7" and labels["Reported by:"] == "T. Man" and labels["Misc:"] == "first article"
    book.close()


def test_the_workspace_path_gives_a_pdf_without_verdicts(tmp_path):
    pytest.importorskip("reportlab")
    pypdf = pytest.importorskip("pypdf")
    study = capability_study(info=INFO)
    result_set, evaluated = prepared(study)
    context = report_context(study, result_set, created="2026-10-31")
    path = export_statistical_report(tmp_path / "cap.pdf", list(evaluated.results), context=context)
    text = " ".join("\n".join(page.extract_text() for page in pypdf.PdfReader(path).pages).split())
    for expected in ("Process Capability Report for Cap", "Gage name: Bench DMM 7", "Misc: first article", "Created 2026-10-31"):
        assert expected in text, expected
    assert "Included" not in text and "PASS" not in text and "FAIL" not in text
