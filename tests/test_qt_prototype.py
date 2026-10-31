"""The Qt proof of concept: its widgets, its background runner and the window on real (synthetic) scans.

Skipped when PySide6 is not installed, so the normal install and CI are unaffected. Runs headless (Qt's ``offscreen``
platform): no window opens. The checks that matter most say that what the window shows and writes is what the
application's own modules produce.
"""
from __future__ import annotations

import os
import threading
import time

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from openpyxl import load_workbook
from PySide6.QtCore import QPoint, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QLabel, QStyleFactory

from datamapperx.exporter import generate_excel
from datamapperx.palette import DARK, LIGHT
from datamapperx.scanner import scan_folder
from datamapperx.settings import Settings
from datamapperx.stat_charts import chart_primitives, charts_for_result
from datamapperx.stat_reports import display_number, summary_rows
from datamapperx.statistics import AnalysisConfig, analyze, normalize_records
from datamapperx.ui_qt.main_window import ANALYSIS, REVIEW, SELECT, MainWindow
from datamapperx.ui_qt.models import TableModel
from datamapperx.ui_qt.tasks import TaskRunner
from datamapperx.ui_qt.theme import apply_theme, make_palette, resolve_mode
from datamapperx.ui_qt.widgets import ActionTable, ChartPanel
from datamapperx.workflow import limit_texts, preview_table

from .conftest import ict_block, write_ict_log


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def pump(app, condition, timeout=15.0):
    end = time.time() + timeout
    while not condition() and time.time() < end:
        app.processEvents()
        time.sleep(0.002)
    assert condition(), "timed out waiting for the window"


@pytest.fixture
def logs(tmp_path):
    """Four samples x three files: Alpha and Beta vary between samples, Gamma is missing from one file."""
    folder = tmp_path / "logs"
    folder.mkdir()
    for s in range(4):
        for t in range(3):
            blocks = [ict_block("Alpha", f"{10 + s + t * 0.01:+.5E}", "+2.0E+01", "+0.0E+00"),
                      ict_block("Beta", f"{5 + s * 2 - t * 0.02:+.5E}", "+3.0E+01", "+0.0E+00")]
            if not (s == 3 and t == 2):
                blocks.append(ict_block("Gamma", f"{1 + s * 0.5 + t * 0.001:+.5E}"))
            write_ict_log(folder / f"S{s}_{t}.log", f"VSE{s:03d}", blocks)
    return folder


@pytest.fixture
def window(app, logs):
    win = MainWindow(app, settings=Settings(theme="light", last_folder=str(logs), expected_samples="4", repeats="3"))
    win.profile_box.setCurrentText("ICT_I3070 Test Log")
    win.show()
    yield win
    win.close()


def scanned(app, win):
    win.scan_logs()
    pump(app, lambda: not win._busy and win.scan_result is not None)
    return win


# ----------------------------------------------------------------------------- theme
def test_the_palette_comes_from_the_shared_colour_tokens(app):
    for tokens in (LIGHT, DARK):
        palette = make_palette(tokens)
        assert palette.window().color().name() == tokens.bg and palette.base().color().name() == tokens.surface
        assert palette.highlight().color().name() == tokens.accent
    assert resolve_mode("light") == "light" and resolve_mode("dark") == "dark"
    assert apply_theme(app, "dark").mode == "dark"
    assert app.palette().window().color().name() == DARK.bg and "QFrame#card" in app.styleSheet()
    styles = QStyleFactory.keys()  # a list of style names
    assert "Fusion" in styles
    apply_theme(app, "light")
    assert app.palette().window().color().name() == LIGHT.bg


# ----------------------------------------------------------------------------- the table model
def test_the_model_serves_text_alignment_and_tints_for_visible_cells_only(app):
    model = TableModel(LIGHT)
    model.set_table(["A", "B"], [(1, 2.5), ("x", None)], tags=["invalid", ""], aligns=["w", "e"], cell_tags={(1, 0): "pass"})
    assert (model.rowCount(), model.columnCount()) == (2, 2)
    index = model.index
    assert model.data(index(0, 1)) == "2.5" and model.data(index(1, 1)) == "" and model.data(index(0, 0)) == "1"
    assert model.data(index(0, 0), Qt.ItemDataRole.BackgroundRole).name() == LIGHT.bad_bg
    assert model.data(index(1, 0), Qt.ItemDataRole.BackgroundRole).name() == LIGHT.ok_bg  # a cell tag wins over the row's
    assert model.data(index(1, 1), Qt.ItemDataRole.BackgroundRole) is None
    assert model.data(index(0, 1), Qt.ItemDataRole.TextAlignmentRole) & int(Qt.AlignmentFlag.AlignRight)
    assert model.headerData(1, Qt.Orientation.Horizontal) == "B"
    model.set_palette(DARK)
    assert model.data(index(0, 0), Qt.ItemDataRole.BackgroundRole).name() == DARK.bad_bg


# ----------------------------------------------------------------------------- background tasks
def test_results_are_delivered_on_the_ui_thread_and_the_work_runs_elsewhere(app):
    runner, seen = TaskRunner(), {}
    main = threading.get_ident()
    runner.run(lambda report: (report(1, 2), threading.get_ident())[1],
               on_progress=lambda *a: seen.setdefault("progress", (a, threading.get_ident())),
               on_success=lambda worker: seen.update(done=threading.get_ident(), worker=worker),
               on_error=lambda exc: seen.update(error=exc))
    pump(app, lambda: "done" in seen)
    assert seen["done"] == main and seen["worker"] != main
    assert seen["progress"] == ((1, 2), main) and not runner.busy


def test_the_gil_switch_interval_is_shortened_only_while_a_task_runs(app):
    import sys

    before, runner, seen = sys.getswitchinterval(), TaskRunner(), {}
    runner.run(lambda report: sys.getswitchinterval(), on_success=lambda inside: seen.update(inside=inside),
               on_error=lambda exc: seen.update(error=exc))
    pump(app, lambda: "inside" in seen or "error" in seen)
    assert seen["inside"] < before and sys.getswitchinterval() == before


def test_an_error_in_the_work_reaches_on_error_and_never_gets_lost(app):
    runner, seen = TaskRunner(), []
    runner.run(lambda report: 1 / 0, on_success=lambda _r: seen.append("ok"), on_error=seen.append)
    pump(app, lambda: bool(seen))
    assert isinstance(seen[0], ZeroDivisionError)


# ----------------------------------------------------------------------------- action list
def table_of_three(kind="add", **options):
    table = ActionTable(LIGHT, kind=kind, name_heading="Component", **options)
    table.resize(420, 300)
    table.set_items(["A", "B", "C"], values=[("1", "2")] * 3 if kind == "add" else None)
    table.show()
    return table


def test_each_click_on_a_button_acts_on_its_own_row_even_when_clicks_are_quick(app):
    table = table_of_three(value_headings=("Min", "Max"))
    pressed = []
    table.action.connect(pressed.append)
    button = lambda row: table.visualRect(table.model().index(row, 3)).center()  # noqa: E731
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=button(1))
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=button(2))  # straight after, on another row
    assert pressed == [1, 2]
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=table.visualRect(table.model().index(0, 0)).center())
    assert pressed == [1, 2]  # a click on the name is a selection, not an action
    assert table.selected_rows() == [0]
    table.close()


def test_enter_and_a_double_click_on_a_name_activate_the_selected_rows_but_not_a_double_click_on_a_button(app):
    table = table_of_three("remove", numbered=True)  # columns: No., Component, button
    activated, deleted = [], []
    table.activated_rows.connect(activated.append)
    table.delete_pressed.connect(lambda: deleted.append(True))
    table.select_rows([0, 2])
    table.doubleClicked.emit(table.model().index(0, 1))
    assert activated == [[0, 2]]
    table.doubleClicked.emit(table.model().index(0, 2))  # the button column
    assert activated == [[0, 2]]
    QTest.keyClick(table, Qt.Key.Key_Return)
    assert activated == [[0, 2], [0, 2]]
    QTest.keyClick(table, Qt.Key.Key_Delete)
    assert deleted == [True]
    table.close()


def test_an_empty_list_says_why(app):
    table = ActionTable(LIGHT, kind="add", name_heading="Component")
    table.resize(300, 200)
    table.set_items([], empty_text="Scan a folder to list\nits components here.")
    image = table.grab().toImage()
    muted = sum(1 for x in range(0, image.width(), 2) for y in range(60, image.height(), 2)
                if image.pixelColor(QPoint(x, y)).name() != image.pixelColor(QPoint(2, 100)).name())
    assert muted > 30  # the hint text was painted into the empty viewport


# ----------------------------------------------------------------------------- the window: what it shows is what the core produces
def test_scan_lists_components_with_their_limits_and_remembers_the_selection(app, window, logs):
    scanned(app, window)
    assert window.samples_value.text() == "4 / 4" and window.files_value.text() == "12" and window.components_value.text() == "3"
    assert window.available.items() == ["Alpha", "Beta", "Gamma"]
    first = window.available.table_model.row_text(0)
    assert first[:3] == ("Alpha", *limit_texts(window.limits, "Alpha")) and first[1] != "–"  # Alpha has limits, Gamma has none
    assert window.available.table_model.row_text(2)[1:3] == ("–", "–")
    window._add(["Beta", "Alpha", "Beta"])  # a name twice is added once, in the order given
    assert window.selected == ["Beta", "Alpha"] and window.available.items() == ["Gamma"]
    assert window.settings.component_sets["ICT_I3070 Test Log"] == ["Beta", "Alpha"]
    window.move_selected(-1)  # nothing is selected in the list: nothing moves
    assert window.selected == ["Beta", "Alpha"]
    window.selected_table.select_rows([1])
    window.move_selected(-1)
    assert window.selected == ["Alpha", "Beta"] and window.selected_table.selected_rows() == [0]


def test_the_review_table_is_the_table_the_core_says_will_be_exported(app, window, logs):
    scanned(app, window)
    window._add(["Alpha", "Gamma"])
    window._show_step(REVIEW)
    expected = preview_table(window.scan_result, ["Alpha", "Gamma"], 0, 3)
    shown = window.preview.table_model
    assert [shown.headerData(c, Qt.Orientation.Horizontal) for c in range(shown.columnCount())] == list(expected.headings)
    assert [tuple(shown.row_text(r)) for r in range(shown.rowCount())] == [tuple(str(v) for v in row) for row in expected.rows]
    assert shown.tags == expected.tags and expected.tags[-1] == "missing"  # the file without Gamma
    chips = [label.text() for label in window.chips.findChildren(QLabel)]
    assert "✓  4 of 4 samples" in chips and "✓  3 files per sample" in chips
    assert any(text.startswith("!  1 empty cell") for text in chips)


def test_the_limits_panel_follows_the_chosen_row(app, window, logs):
    scanned(app, window)
    window._add(["Alpha", "Beta"])
    window._show_step(REVIEW)
    window.preview.select_row(5)
    app.processEvents()
    assert window.limits_row.text().startswith("Row 6") and "VSE001" in window.limits_row.text()
    model = window.limits_table.table_model
    assert [model.row_text(r)[0] for r in range(model.rowCount())] == ["Alpha", "Beta"]
    assert model.row_text(0)[2:4] == ("+0.0E+00", "+2.0E+01")


def test_export_writes_the_same_workbook_as_the_core_exporter(app, window, logs, tmp_path, monkeypatch):
    scanned(app, window)
    window._add(["Alpha", "Beta"])
    target = tmp_path / "from_qt.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    window.export_excel()
    pump(app, lambda: not window._busy and target.exists())
    reference = tmp_path / "from_core.xlsx"
    generate_excel(scan_folder(logs, window.profiles["ICT_I3070 Test Log"]).grouped, ["Alpha", "Beta"], reference, 3)

    def cells(path):
        book = load_workbook(path)
        try:
            return [[c.value for c in row] for row in book["Raw Data"].iter_rows()]
        finally:
            book.close()

    assert cells(target) == cells(reference) and len(cells(target)) > 12


def test_the_analysis_page_shows_the_numbers_the_engine_computes(app, window, logs):
    scanned(app, window)
    window._add(["Alpha", "Beta"])
    window._show_step(ANALYSIS)
    window.analysis.run()
    pump(app, lambda: bool(window.analysis.results))
    scan = scan_folder(logs, window.profiles["ICT_I3070 Test Log"])
    engine = analyze(normalize_records(scan.grouped, ["Alpha", "Beta"], assign_trials=True),
                     AnalysisConfig(method="grr", expected_repeats=3))
    shown = [r.metrics for r in window.analysis.results]
    assert [m["percent_study_var"] for m in shown] == pytest.approx([r.metrics["percent_study_var"] for r in engine], rel=1e-12)
    assert [m["ndc"] for m in shown] == [r.metrics["ndc"] for r in engine]
    table = window.analysis.table.table_model
    headers, rows = summary_rows(window.analysis.results, judged=False)
    titles = [table.headerData(c, Qt.Orientation.Horizontal) for c in range(table.columnCount())]
    column = titles.index("%Study Var")  # the table text is the application's own formatting of the same numbers
    assert [table.row_text(r)[column] for r in range(table.rowCount())] == [display_number(row[headers.index("%Study Var")], 2) for row in rows]
    window.analysis.table.select_row(1)
    app.processEvents()
    assert len(window.analysis.charts._specs) == len(charts_for_result(window.analysis.results[1]))


def test_a_step_that_is_not_ready_stays_on_select(app, window):
    window._show_step(REVIEW)
    assert window.pages.currentIndex() == SELECT


# ----------------------------------------------------------------------------- charts
def test_every_chart_of_an_item_is_painted(app, window, logs):
    scanned(app, window)
    window._add(["Alpha"])
    window._show_step(ANALYSIS)
    window.analysis.run()
    pump(app, lambda: bool(window.analysis.results))
    specs = charts_for_result(window.analysis.results[0])
    assert len(specs) == 4 and all(chart_primitives(spec, 400, 220) for spec in specs)
    panel = ChartPanel()
    panel.resize(900, 700)
    panel.set_specs(specs)
    image = panel.grab().toImage()
    dark = sum(1 for x in range(0, image.width(), 3) for y in range(0, image.height(), 3) if image.pixelColor(QPoint(x, y)).lightness() < 200)
    assert dark > 200  # lines, bars and text were drawn on the white background
