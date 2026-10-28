from __future__ import annotations

import threading

import pytest
from openpyxl import load_workbook

from datamapperx.exporter import (
    generate_excel,
    iter_export_rows,
    set_cell_value,
    summarize_limits,
)
from datamapperx.scanner import scan_folder

from .conftest import ict_block, write_ft_workbook, write_ict_log


def make_ict_folder(root, sns=("VSE001", "VSE002"), repeats=2):
    for sn in sns:
        for run in range(1, repeats + 1):
            write_ict_log(
                root / f"{sn}_run{run}.log", sn,
                [ict_block("R1", f"+{run}.0E+00", "+9.0E+00", "+0.0E+00"), ict_block("C1", "+1.0E-06", kind="CAP")],
            )
    return root


# ----------------------------------------------------------------------------- scanner
def test_scan_groups_by_sn_and_sorts_files(tmp_path, ict_profile):
    make_ict_folder(tmp_path)
    result = scan_folder(tmp_path, ict_profile)

    assert list(result.grouped) == ["VSE001", "VSE002"]
    assert [r.filename for r in result.grouped["VSE001"]] == ["VSE001_run1.log", "VSE001_run2.log"]
    assert result.components == ["C1", "R1"]
    assert result.total_sns == 2 and result.total_files == 4
    assert result.sns_with_wrong_repeat_count(2) == []
    assert result.sns_with_wrong_repeat_count(3) == ["VSE001", "VSE002"]


def test_scan_walks_subfolders_and_reports_ignored_extensions(tmp_path, ict_profile):
    sub = tmp_path / "deep" / "er"
    sub.mkdir(parents=True)
    make_ict_folder(sub, sns=("VSE010",), repeats=1)
    (tmp_path / "notes.docx").write_text("x")
    (tmp_path / "~$lock.log").write_text("x")

    result = scan_folder(tmp_path, ict_profile)

    assert result.total_sns == 1
    assert result.diagnostics.ignored_extensions == {".docx": 1}
    assert result.diagnostics.accepted_file_count == 1


def test_scan_skips_files_without_sn_and_reports_them(tmp_path, ict_profile):
    (tmp_path / "anon.log").write_text(ict_block("R1", "+1.0E+00"))
    result = scan_folder(tmp_path, ict_profile)
    assert result.total_sns == 0
    assert result.diagnostics.skipped_without_sn == ["anon.log"]


def test_scan_reports_unreadable_excel_as_warning(tmp_path, ft_profile):
    (tmp_path / "VSE5.xlsx").write_text("not a zip")
    write_ft_workbook(tmp_path / "VSE6.xlsx", "VSE6", [("R1", 1, 0, 2, "V")])
    result = scan_folder(tmp_path, ft_profile)
    assert list(result.grouped) == ["VSE6"]
    assert len(result.parse_warnings) == 1 and "VSE5.xlsx" in result.parse_warnings[0]


def test_scan_exposes_duplicate_groups(tmp_path, ict_profile):
    write_ict_log(tmp_path / "VSE1.log", "VSE1", [ict_block("R1", "+1.0E+00"), ict_block("R1", "+2.0E+00")])
    result = scan_folder(tmp_path, ict_profile)
    assert len(result.duplicate_groups) == 1
    group = result.duplicate_groups[0]
    assert group.component == "R1" and len(group.options) == 2
    assert group.record is result.grouped["VSE1"][0]


def test_scan_reports_progress_and_honours_cancel(tmp_path, ict_profile):
    make_ict_folder(tmp_path, sns=("VSE001", "VSE002", "VSE003"), repeats=3)
    seen: list[tuple[int, int]] = []
    scan_folder(tmp_path, ict_profile, progress=lambda d, t: seen.append((d, t)))
    assert seen[0] == (0, 9) and seen[-1] == (9, 9)

    cancel = threading.Event()
    cancel.set()
    cancelled = scan_folder(tmp_path, ict_profile, cancel_event=cancel)
    assert cancelled.cancelled and cancelled.total_files == 0


def test_scan_rejects_invalid_profile(tmp_path, ict_profile):
    ict_profile["measure_regex"] = "(unclosed"
    with pytest.raises(ValueError, match="Invalid Regex"):
        scan_folder(tmp_path, ict_profile)


# ----------------------------------------------------------------------------- exporter
def test_export_layout_matches_the_v1_format(tmp_path, ict_profile):
    logs = tmp_path / "logs"
    logs.mkdir()
    make_ict_folder(logs)
    result = scan_folder(logs, ict_profile)
    out = tmp_path / "out.xlsx"

    generate_excel(result.grouped, ["R1", "C1"], out, expected_repeated=2)

    ws = load_workbook(out)["Raw Data"]
    assert [c.value for c in ws[1]][:8] == ["STT", "SN", "R1", "C1", None, "NAME", "MIN", "MAX"]
    assert [c.value for c in ws[2]][:4] == [1, "VSE001", 1, 1e-06]
    assert [c.value for c in ws[5]][:4] == [4, "VSE002", 2, 1e-06]
    # Summary block: name, min, max, unit.
    assert [ws.cell(2, col).value for col in range(6, 10)] == ["R1", 0.0, 9.0, "RES"]
    assert ws["D2"].number_format == "0.0E+00"
    assert ws.freeze_panes == "C2"
    assert ws.max_row == 5


def test_export_refuses_wrong_repeat_count(tmp_path, ict_profile):
    make_ict_folder(tmp_path)
    result = scan_folder(tmp_path, ict_profile)
    with pytest.raises(ValueError, match="exactly 3 files"):
        generate_excel(result.grouped, ["R1"], tmp_path / "o.xlsx", expected_repeated=3)


def test_export_leaves_no_temp_file_and_overwrites_cleanly(tmp_path, ict_profile):
    make_ict_folder(tmp_path)
    result = scan_folder(tmp_path, ict_profile)
    out = tmp_path / "out.xlsx"
    generate_excel(result.grouped, ["R1"], out, 2)
    generate_excel(result.grouped, ["C1"], out, 2)
    assert load_workbook(out)["Raw Data"]["C1"].value == "C1"
    assert not list(tmp_path.glob("*.tmp"))


def test_export_reports_locked_output(tmp_path, ict_profile, monkeypatch):
    make_ict_folder(tmp_path)
    result = scan_folder(tmp_path, ict_profile)

    def locked(*_a, **_k):
        raise PermissionError

    monkeypatch.setattr("datamapperx.exporter.os.replace", locked)
    with pytest.raises(PermissionError, match="open in Excel"):
        generate_excel(result.grouped, ["R1"], tmp_path / "out.xlsx", 2)
    assert not list(tmp_path.glob("*.tmp"))


def test_preview_rows_match_export_order(tmp_path, ict_profile):
    make_ict_folder(tmp_path)
    result = scan_folder(tmp_path, ict_profile)
    rows = list(iter_export_rows(result.grouped, ["C1", "R1"]))
    assert [r.index for r in rows] == [1, 2, 3, 4]
    assert [r.sn for r in rows] == ["VSE001", "VSE001", "VSE002", "VSE002"]
    assert rows[0].values == ("+1.0E-06", "+1.0E+00")


def test_missing_component_gives_blank_value(tmp_path, ict_profile):
    write_ict_log(tmp_path / "VSE1.log", "VSE1", [ict_block("R1", "+1.0E+00")])
    result = scan_folder(tmp_path, ict_profile)
    assert next(iter_export_rows(result.grouped, ["R1", "NOPE"])).values == ("+1.0E+00", "")
    assert summarize_limits(result.grouped, ["NOPE"])["NOPE"] == {"min": "", "max": "", "unit": ""}


# ----------------------------------------------------------------------------- cell writer
class FakeCell:
    value = None
    number_format = "General"
    data_type = "n"


@pytest.mark.parametrize(
    ("text", "value", "number_format"),
    [
        ("+1.00500E+03", 1005.0, "0.00000E+00"),
        ("2E-07", 2e-07, "0E+00"),
        ("12", 12, "General"),
        ("12.0", 12, "General"),
        ("3.75", 3.75, "General"),
        ("-0.5", -0.5, "General"),
    ],
)
def test_set_cell_value_numbers(text, value, number_format):
    cell = FakeCell()
    set_cell_value(cell, text)
    assert cell.value == value and type(cell.value) is type(value)
    assert cell.number_format == number_format


@pytest.mark.parametrize("text", ["nan", "inf", "-Infinity", "1_000", "PASS", "0x1F", "", "12345678901234567890"])
def test_set_cell_value_keeps_non_numbers_as_text(text):
    cell = FakeCell()
    set_cell_value(cell, text)
    assert cell.value == text


def test_formula_like_text_is_never_a_formula():
    cell = FakeCell()
    set_cell_value(cell, "=HYPERLINK(\"http://x\")")
    assert cell.data_type == "s"
