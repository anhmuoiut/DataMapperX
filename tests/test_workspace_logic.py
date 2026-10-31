"""Phase 1 logic: the study design, the RawData.xlsx importer, notes found in the data, statistics-only results, Ppk."""
from __future__ import annotations

import copy
import dataclasses

import pytest

from datamapperx.exporter import generate_excel
from datamapperx.profiles import DEFAULT_PROFILES
from datamapperx.scanner import scan_folder
from datamapperx.statistics import (
    NOT_EVALUATED,
    PASS,
    AnalysisConfig,
    Dataset,
    DesignError,
    EvaluationSettings,
    Observation,
    RawDataFormatError,
    Severity,
    StaleResultError,
    StatisticalOptions,
    Study,
    analyze,
    apply_design,
    auto_within_method,
    describe_design,
    detect_conditions,
    import_rawdata_xlsx,
    validate_config,
)
from datamapperx.statistics.design import operator_names, relevant_design
from datamapperx.statistics.evaluation import validate_evaluation_settings

from .baseline_support import write_rawdata_logs


def rows(samples=3, repeats=4, **extra):
    """``repeats`` readings for each sample, in the order a RawData worksheet would have them."""
    return [Observation("X", f"S{s}", 10 + s + 0.1 * r, 0, 30, source=f"r{s}-{r}", **extra)
            for s in range(1, samples + 1) for r in range(repeats)]


# ----------------------------------------------------------------------------- study design
def test_operators_are_assigned_by_row_order_within_each_sample():
    data = rows(samples=2, repeats=4)
    by_operator = apply_design(data, {"operators": ["A", "B"]}, "grr")
    assert [(o.operator, o.trial) for o in by_operator[:4]] == [("A", "1"), ("A", "2"), ("B", "1"), ("B", "2")]
    by_pass = apply_design(data, {"operators": ["A", "B"], "operator_order": "by_pass"}, "grr")
    assert [(o.operator, o.trial) for o in by_pass[:4]] == [("A", "1"), ("B", "1"), ("A", "2"), ("B", "2")]
    assert [o.value for o in by_operator] == [o.value for o in data]  # only who and which trial are added
    assert all(o.operator == "" for o in data)  # the dataset itself is never touched


def test_one_appraiser_is_the_default_and_changes_nothing():
    data = rows()
    assert apply_design(data, {}, "grr") == tuple(data)
    assert apply_design(data, {"operators": ["only one"]}, "grr") == tuple(data)
    assert operator_names({"operators": ["only one"]}) == ()
    assert "One appraiser" in describe_design({}, "grr")[0]


def test_a_design_that_does_not_fit_the_data_says_what_to_change():
    with pytest.raises(DesignError, match="multiple of 2"):
        apply_design(rows(samples=2, repeats=3), {"operators": ["A", "B"]}, "grr")
    with pytest.raises(DesignError, match="already identify operators"):
        apply_design(rows(operator="A"), {"operators": ["A", "B"]}, "grr")
    with pytest.raises(DesignError, match="different"):
        apply_design(rows(), {"operators": ["A", "A"]}, "grr")
    with pytest.raises(DesignError, match="not both"):
        apply_design(rows(slot="1"), {"operators": ["A", "B"], "nest_as_operator": True}, "grr")
    with pytest.raises(DesignError, match="Slot"):
        apply_design(rows(), {"nest_as_operator": True}, "grr")


def test_nest_as_operator_uses_the_slot_as_operator_and_says_so():
    data = [dataclasses.replace(o, slot=str(1 + i % 2)) for i, o in enumerate(rows(samples=3, repeats=4))]
    nested = apply_design(data, {"nest_as_operator": True}, "grr")
    assert {o.operator for o in nested} == {"1", "2"} and all(o.slot == "" for o in nested)
    assert "not between people" in describe_design({"nest_as_operator": True}, "grr")[0]
    conditions = detect_conditions([], "grr", "worksheet", {"nest_as_operator": True})
    assert [c.code for c in conditions] == ["nest_as_operator"] and conditions[0].severity is Severity.INFORMATION


def test_subgroup_from_sample_only_applies_to_capability():
    data = rows(samples=2, repeats=3)
    grouped = apply_design(data, {"subgroup_from_sample": True}, "capability")
    assert [o.subgroup for o in grouped] == ["S1"] * 3 + ["S2"] * 3
    assert apply_design(data, {"subgroup_from_sample": True}, "grr") == tuple(data)


def test_only_the_design_that_matters_to_the_method_makes_results_stale():
    data = Dataset.from_observations(rows(samples=3, repeats=4))
    grr = Study(data, options=StatisticalOptions(method="grr"))
    result_set = grr.run_statistics()
    grr.design["subgroup_from_sample"] = True  # a capability setting: no effect on a GRR
    assert not grr.is_stale(result_set)
    grr.design["operators"] = ["A", "B"]
    assert grr.stale_reasons(result_set) == ("design",)
    assert relevant_design({"operators": ["A", "B"]}, "capability") == {"subgroup_from_sample": False}


def test_a_study_with_operators_is_a_crossed_study_and_without_them_a_one_appraiser_study():
    data = Dataset.from_observations(rows(samples=5, repeats=4))
    study = Study(data, options=StatisticalOptions(method="grr", assign_trials=True))  # RawData has no Trial column
    one = study.run_statistics().statistics[0]
    assert one.valid and not one.facts["crossed"]
    study.design["operators"] = ["A", "B"]
    crossed = study.run_statistics()
    assert crossed.statistics[0].valid and crossed.statistics[0].facts["crossed"]
    assert crossed.id == "RS-0002" and crossed.inputs != study.result_sets[0].inputs
    assert study.is_stale(study.result_sets[0]) and not study.is_stale(crossed)


def test_changing_the_design_never_changes_the_dataset_fingerprint():
    study = Study(Dataset.from_observations(rows()), options=StatisticalOptions(method="grr"))
    before = study.dataset.fingerprint
    study.design["operators"] = ["A", "B"]
    study.run_statistics()
    assert study.dataset.fingerprint == before


# ----------------------------------------------------------------------------- RawData.xlsx importer
@pytest.fixture
def rawdata(tmp_path):
    logs = tmp_path / "logs"
    logs.mkdir()
    write_rawdata_logs(logs)
    result = scan_folder(logs, copy.deepcopy(DEFAULT_PROFILES["ICT_I3070 Test Log"]))
    path = tmp_path / "RawData.xlsx"
    generate_excel(result.grouped, ["R1", "Iin @ No  Load", "Long"], path, expected_repeated=2)
    return path


def test_importing_the_exported_workbook_reads_every_reading_in_worksheet_order(rawdata):
    imported = import_rawdata_xlsx(rawdata)
    dataset = imported.dataset
    assert imported.items == ("R1", "Iin @ No  Load", "Long") and imported.rows == 6 and imported.samples == 3
    assert dataset.order_basis == "worksheet" and len(dataset.observations) == 18
    first = [o for o in dataset.observations if o.test_item == "R1"]
    assert [o.sample_id for o in first] == ["VSE000123", "VSE000123", "VSE000124", "VSE000124", "VSE000125", "VSE000125"]
    assert first[0].source == "Raw Data!R2" and len({o.source for o in first}) == 6
    assert imported.warnings == ()


def test_limits_units_and_text_serials_come_from_the_table_beside_the_data(rawdata):
    by_item = {o.test_item: o for o in import_rawdata_xlsx(rawdata).dataset.observations}
    assert (by_item["R1"].lsl, by_item["R1"].usl) == (0, 900)
    assert by_item["Iin @ No  Load"].usl == pytest.approx(0.01) and by_item["Iin @ No  Load"].lsl == pytest.approx(-0.01)
    assert all(isinstance(o.sample_id, str) for o in import_rawdata_xlsx(rawdata).dataset.observations)


def test_the_limits_table_is_never_read_as_measurements(rawdata):
    imported = import_rawdata_xlsx(rawdata)
    assert set(imported.items) == {"R1", "Iin @ No  Load", "Long"}  # no NAME / MIN / MAX / UNIT column
    assert all(o.sample_id.startswith("VSE") for o in imported.dataset.observations)


def test_a_missing_reading_stays_missing_and_is_reported_by_the_analysis(rawdata):
    values = [o for o in import_rawdata_xlsx(rawdata).dataset.observations if o.test_item == "Iin @ No  Load"]
    assert [o.value for o in values].count("") == 1  # VSE000124's second file lacks the component
    result = analyze(values, AnalysisConfig(method="grr"))[0]
    assert not result.valid and any("missing" in issue.message.lower() or "numeric" in issue.message.lower() for issue in result.issues)


def test_an_importer_error_names_the_problem(tmp_path):
    from openpyxl import Workbook

    wrong = tmp_path / "wrong.xlsx"
    book = Workbook()
    book.active.append(["Name", "Value"])
    book.save(wrong)
    with pytest.raises(RawDataFormatError, match="STT and SN"):
        import_rawdata_xlsx(wrong)
    empty = tmp_path / "empty.xlsx"
    book = Workbook()
    book.active.append(["STT", "SN", "R1"])
    book.save(empty)
    with pytest.raises(RawDataFormatError, match="no data rows"):
        import_rawdata_xlsx(empty)
    duplicate = tmp_path / "duplicate.xlsx"
    book = Workbook()
    book.active.append(["STT", "SN", "R1", "R1"])
    book.active.append([1, "A", 1, 2])
    book.save(duplicate)
    with pytest.raises(RawDataFormatError, match="more than once"):
        import_rawdata_xlsx(duplicate)


def test_a_workbook_without_a_limits_table_imports_with_a_warning(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "plain.xlsx"
    book = Workbook()
    book.active.title = "Raw Data"
    book.active.append(["STT", "SN", "X"])
    for number in range(1, 4):
        book.active.append([number, f"00{number}", float(number)])
    book.save(path)
    imported = import_rawdata_xlsx(path)
    assert [o.sample_id for o in imported.dataset.observations] == ["001", "002", "003"]  # leading zeros survive
    assert imported.warnings and "no specification limits" in imported.warnings[0]


# ----------------------------------------------------------------------------- conditions found in the data
def capability_study(values, order_basis="worksheet", **kwargs):
    observations = [Observation("X", f"S{i}", v, 0, 100, source=f"c{i}") for i, v in enumerate(values)]
    return Study(Dataset.from_observations(observations, order_basis=order_basis), options=StatisticalOptions(method="capability"), **kwargs)


def codes(evaluated):
    return {c.code: c.severity for c in evaluated.record.validation.conditions}


def test_capability_conditions_follow_the_data_and_the_order_basis():
    worksheet = capability_study([50 + (i % 7) for i in range(30)])
    found = codes(worksheet.evaluate(worksheet.run_statistics()))
    assert found["acquisition_order"] is Severity.WARNING and found["sample_size"] is Severity.INFORMATION
    assert found["process_stability"] is Severity.WARNING and found["normality"] is Severity.WARNING
    acquired = capability_study([50 + (i % 7) for i in range(120)], order_basis="acquisition")
    found = codes(acquired.evaluate(acquired.run_statistics()))
    assert "acquisition_order" not in found and "sample_size" not in found


def test_repeated_readings_of_one_unit_are_a_warning_for_capability():
    observations = [Observation("X", f"S{i // 3}", 50 + (i % 5), 0, 100, source=f"c{i}") for i in range(30)]
    study = Study(Dataset.from_observations(observations), options=StatisticalOptions(method="capability"))
    found = codes(study.evaluate(study.run_statistics()))
    assert found["repeated_measurements"] is Severity.WARNING


def test_grr_notes_are_information_and_follow_minitabs_design_advice():
    study = Study(Dataset.from_observations(rows(samples=5, repeats=4)), options=StatisticalOptions(method="grr", assign_trials=True))
    evaluated = study.evaluate(study.run_statistics())
    found = codes(evaluated)
    assert set(found) == {"one_appraiser", "trial_inferred", "design_samples"} and set(found.values()) == {Severity.INFORMATION}
    messages = {c.code: c.message for c in evaluated.record.validation.conditions}
    assert "not the order of measurement" in messages["trial_inferred"] and "Minitab advises at least 10" in messages["design_samples"]
    assert not any("work instruction" in m.lower() for m in messages.values())
    study.design["operators"] = ["A", "B"]
    crossed = codes(study.evaluate(study.run_statistics()))
    assert "one_appraiser" not in crossed and crossed["design_operators"] is Severity.INFORMATION


def test_the_notes_of_a_capability_study_never_cite_a_work_instruction():
    study = capability_study([50 + (i % 7) for i in range(30)])
    evaluated = study.evaluate(study.run_statistics())
    assert not any("work instruction" in c.message.lower() for c in evaluated.record.validation.conditions)
    text = {c.code: c.message for c in evaluated.record.validation.conditions}
    assert "Pp and Ppk do not" in text["acquisition_order"] and "Xbar-S" in text["process_stability"] and "Box-Cox" in text["normality"]


# ----------------------------------------------------------------------------- results are shown at once, criteria are optional
def test_without_criteria_the_statuses_are_exactly_what_analyze_gives():
    study = capability_study([50 + (i % 7) for i in range(60)])
    result_set = study.run_statistics()
    plain = study.evaluate(result_set)
    assert [r.status for r in plain.results] == [r.status for r in analyze(
        list(study.dataset.observations), dataclasses.replace(study.options.to_config()))]


def test_a_capability_result_is_judged_at_once_and_the_notes_do_not_block_it():
    study = capability_study([50 + (i % 7) for i in range(60)])
    evaluated = study.evaluate(study.run_statistics())
    assert evaluated.results[0].status in (PASS, "FAIL") and evaluated.record.judged
    assert {"acquisition_order", "process_stability", "normality"} <= set(codes(evaluated))


def test_statistics_only_means_no_status_no_decision_metric_and_the_same_numbers():
    study = capability_study([50 + (i % 7) for i in range(60)])
    result_set = study.run_statistics()
    judged, bare = study.evaluate(result_set), study.evaluate(result_set, judge=False)
    assert bare.results[0].status == NOT_EVALUATED and bare.results[0].decision_metric == "" and not bare.record.judged
    assert bare.results[0].metrics == judged.results[0].metrics
    assert bare.record.criteria_snapshot == {"kind": "none"} and bare.record.criteria_version != judged.record.criteria_version
    assert "acceptance_conflicts" not in bare.results[0].details and "cpk_threshold" not in bare.results[0].details
    grr = Study(Dataset.from_observations(rows(samples=5, repeats=4)), options=StatisticalOptions(method="grr", assign_trials=True))
    flat = grr.evaluate(grr.run_statistics(), judge=False)
    assert flat.results[0].valid and flat.results[0].status == NOT_EVALUATED and "ndc" in flat.results[0].metrics


def test_an_invalid_item_has_no_status_with_or_without_criteria():
    study = capability_study([50, "abc", 52, 53])
    result_set = study.run_statistics()
    judged, bare = study.evaluate(result_set), study.evaluate(result_set, judge=False)
    assert [r.status for r in judged.results] == [r.status for r in bare.results] and not bare.results[0].valid


def test_within_sigma_follows_the_subgroup_size():
    def data(subgroups, item="X"):
        return [Observation(item, f"S{i}", float(i), 0, 99, subgroup=g, source=f"{item}{i}") for i, g in enumerate(subgroups)]

    assert auto_within_method(data(["", "", "", ""])) == "individuals"  # no subgroups
    assert auto_within_method(data(["a", "b", "c", "d"])) == "individuals"  # subgroup size 1
    assert auto_within_method(data(["a", "a", "b", "b"])) == "pooled"
    assert auto_within_method(data(["a", "a", "b", ""])) == "individuals"  # a reading with no subgroup
    assert auto_within_method(data(["a", "a", "b"])) == "individuals"  # one subgroup of size 1
    assert auto_within_method(data(["a", "a", "b", "b"], "X") + data(["a", "b", "b"], "Y")) == "individuals"  # per item
    assert auto_within_method([]) == "individuals"


def test_stale_results_cannot_be_required_after_the_design_changes():
    study = capability_study([50 + (i % 7) for i in range(30)])
    result_set = study.run_statistics()
    study.design["subgroup_from_sample"] = True
    with pytest.raises(StaleResultError):
        study.require_current(result_set)


# ----------------------------------------------------------------------------- Ppk as the decision index
def decision_study(metric):
    study = capability_study([50 + (i % 7) * 0.4 for i in range(60)], evaluation=EvaluationSettings(decision_metric=metric, cpk_threshold=1.33))
    return study, study.evaluate(study.run_statistics())


def test_the_decision_index_can_be_cpk_or_ppk_and_defaults_to_cpk():
    for metric in ("cpk", "ppk"):
        _, evaluated = decision_study(metric)
        result = evaluated.results[0]
        assert result.decision_metric == metric
        expected = PASS if result.metrics[metric] > 1.33 else "FAIL"
        assert result.status == expected
    default = capability_study([50 + (i % 7) * 0.4 for i in range(60)])
    assert default.evaluate(default.run_statistics()).results[0].decision_metric == "cpk"


def test_ppk_can_decide_a_different_verdict_than_cpk():
    values = [50 + (i % 7) * 0.4 + (0.9 if i % 2 else 0) * (i // 10) for i in range(60)]  # a drifting process
    study = capability_study(values)
    result_set = study.run_statistics()
    cpk = study.evaluate(result_set, EvaluationSettings(decision_metric="cpk", cpk_threshold=1.33)).results[0]
    ppk = study.evaluate(result_set, EvaluationSettings(decision_metric="ppk", cpk_threshold=1.33)).results[0]
    assert cpk.metrics == ppk.metrics and cpk.decision_metric == "cpk" and ppk.decision_metric == "ppk"
    assert cpk.metrics["cpk"] != ppk.metrics["ppk"]


def test_a_grr_study_cannot_use_a_capability_index_to_decide():
    with pytest.raises(ValueError, match="supported GRR acceptance metric"):
        validate_evaluation_settings(EvaluationSettings(decision_metric="cpk"), "grr")
    validate_evaluation_settings(EvaluationSettings(decision_metric="ppk"), "capability")
    with pytest.raises(ValueError, match="supported GRR acceptance metric"):
        validate_config(AnalysisConfig(method="grr", decision_metric="ppk"))
    validate_config(AnalysisConfig(method="capability", decision_metric="ppk"))


# ----------------------------------------------------------------------------- running on a worker thread
def test_a_plan_keeps_the_inputs_it_was_made_from_even_if_the_study_is_edited_before_it_runs():
    study = Study(Dataset.from_observations(rows(samples=5, repeats=4)), options=StatisticalOptions(method="grr", assign_trials=True))
    plan = study.plan_statistics()
    study.design["operators"] = ["A", "B"]  # edited while the plan waits for its worker thread
    result_set = study.adopt(plan.execute())
    assert result_set.id == "RS-0001" and study.is_stale(result_set) and study.result_sets == [result_set]
    assert result_set.statistics[0].valid and not result_set.statistics[0].facts["crossed"]
    assert study.adopt(study.plan_statistics().execute()).id == "RS-0002"


def test_a_study_that_cannot_be_analysed_fails_when_the_plan_is_made_not_on_the_worker():
    study = Study(Dataset.from_observations(rows(samples=2, repeats=3)), options=StatisticalOptions(method="grr"),
                  design={"operators": ["A", "B"]})
    with pytest.raises(DesignError, match="multiple of 2"):
        study.plan_statistics()
    study.options = StatisticalOptions(method="grr", study_factor=-1)
    with pytest.raises(ValueError, match="study_factor"):
        study.plan_statistics()
