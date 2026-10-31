"""Reports are modelled on Minitab's Gage R&R and Process Capability reports and print what the analysis found."""

from __future__ import annotations

import copy
import math
import re

import pytest
from openpyxl import load_workbook

from datamapperx.stat_charts import ChartSeries, ChartSpec, chart_primitives, charts_for_result, draw_tk_chart
from datamapperx.stat_reports import (
    ReportContext,
    capability_indices,
    capability_ppm,
    capability_tables,
    display_number,
    export_statistical_report,
    format_value,
    grr_tables,
    item_label,
    session_text,
    summary_rows,
)
from datamapperx.statistics import AnalysisConfig, Observation, analyze

INFO = ReportContext(info=(("Gage name", "FIXTURE-1163"), ("Date of study", "10-22-2026"), ("Reported by", "Thao Le")),
                     included=1, total=1)


def capability(name="C3", count=40, low=0, high=6):
    return analyze([Observation(name, str(i), 2.5 + (i % 7) * 0.4, low, high, source=f"log{i}.txt") for i in range(count)],
                   AnalysisConfig())[0]


def grr(item="DCDC5W, Meas Vin", low=0, high=20):
    rows = [Observation(item, str(part), value, low, high, trial=str(trial))
            for part, values in enumerate(((1, 2, 3.2), (4, 5.1, 6), (7, 8, 9.3), (10, 10.4, 12)), 1)
            for trial, value in enumerate(values, 1)]
    return analyze(rows, AnalysisConfig(method="grr"))[0]


def crossed_grr(shift=1.0, retained=False):
    """Three samples measured twice by operators A and B; B reads `shift` higher. ``retained``: B also disagrees on S3."""
    rows = []
    for operator, offset in (("A", 0.0), ("B", shift)):
        for sample, centre in {"S1": 2, "S2": 5, "S3": 8}.items():
            bump = 3.0 if (sample, operator) == ("S3", "B") and retained else 0.0
            for trial, value in enumerate((centre - (0.2 if retained else 1), centre + (0.2 if retained else 1)), 1):
                rows.append(Observation("Crossed item", sample, value + offset + bump, 0, 20, trial=str(trial), operator=operator))
    return analyze(rows, AnalysisConfig(method="grr"))[0]


def text_of(path):
    pypdf = pytest.importorskip("pypdf")
    return "\n".join(page.extract_text() for page in pypdf.PdfReader(path).pages)


def grr_with_serials(count):
    serials = tuple(f"VSEC0024{index:05d}" for index in range(1, count + 1))
    rows = [Observation("DCDC5W, Iin @ No Load", serial, .0068 + part * .00001 + delta, .006, .008,
                        trial=str(trial), source=f"{serial}_{trial}.log")
            for part, serial in enumerate(serials)
            for trial, delta in enumerate((-.00001, .000004, .000013), 1)]
    result, = analyze(rows, AnalysisConfig(method="grr"))
    assert result.valid
    return result, serials


# ----------------------------------------------------------------------------- Variable Gage R&R: the tables of Minitab's report
def test_one_way_report_has_the_tables_minitab_prints_and_nothing_else():
    result = grr()
    tables = grr_tables(result)
    assert [t.title for t in tables] == ["One-Way ANOVA Table", "Variance Components", "Gage Evaluation"]
    anova, components, evaluation = tables
    assert anova.headers == ("Source", "DF", "SS", "MS", "F", "P") and [r[0] for r in anova.rows] == ["Sample ID", "Repeatability", "Total"]
    assert [r[0].strip() for r in components.rows] == ["Total Gage R&R", "Repeatability", "Part-To-Part", "Total Variation"]
    assert components.headers[2].startswith("%Contribution")
    assert evaluation.headers[1:] == ("StdDev (SD)", "Study Var\n(6 × SD)", "%Study Var\n(%SV)", "%Tolerance\n(SV/Toler)")


def test_the_numbers_in_the_tables_are_the_analysis_numbers():
    result = grr()
    m = result.metrics
    _, components, evaluation = grr_tables(result)
    by_source = {row[0].strip(): row for row in components.rows}
    assert by_source["Total Gage R&R"][1] == m["variance_grr"] and by_source["Part-To-Part"][1] == m["variance_part"]
    assert by_source["Total Gage R&R"][2] == pytest.approx(100 * m["variance_grr"] / m["variance_total"])
    assert sum(by_source[s][2] for s in ("Total Gage R&R", "Part-To-Part")) == pytest.approx(100)
    gage = {row[0].strip(): row for row in evaluation.rows}
    assert gage["Total Gage R&R"][1] == m["sigma_grr"] and gage["Total Gage R&R"][2] == pytest.approx(6 * m["sigma_grr"])
    assert gage["Total Gage R&R"][3] == pytest.approx(m["percent_study_var"]) and gage["Total Gage R&R"][4] == pytest.approx(m["percent_tolerance"])
    assert gage["Total Variation"][3] == pytest.approx(100)
    assert "Number of Distinct Categories = " + str(m["ndc"]) in session_text(result)


def test_without_two_limits_there_is_no_tolerance_column():
    result = grr(low=None, high=20)
    evaluation = grr_tables(result)[-1]
    assert evaluation.headers == ("Source", "StdDev (SD)", "Study Var\n(6 × SD)", "%Study Var\n(%SV)")
    assert all(len(row) == 4 for row in evaluation.rows)


def test_crossed_report_shows_the_two_way_tables_and_the_reproducibility_breakdown():
    removed = crossed_grr()
    titles = [t.title for t in grr_tables(removed)]
    assert titles == ["Two-Way ANOVA Table With Interaction", "Two-Way ANOVA Table Without Interaction", "Variance Components", "Gage Evaluation"]
    assert grr_tables(removed)[0].note == "α to remove interaction term = 0.05"
    sources = [r[0].strip() for r in grr_tables(removed)[2].rows]
    assert sources == ["Total Gage R&R", "Repeatability", "Reproducibility", "Operator", "Part-To-Part", "Total Variation"]
    retained = crossed_grr(retained=True)
    assert retained.details["interaction_removed"] is False
    assert [t.title for t in grr_tables(retained)][0] == "Two-Way ANOVA Table With Interaction" and len(grr_tables(retained)) == 3
    assert "Operator*Sample ID" in [r[0].strip() for r in grr_tables(retained)[1].rows]


# ----------------------------------------------------------------------------- Process Capability: the report of Minitab
def test_capability_indices_reproduce_the_figures_in_minitabs_process_capability_report():
    """The figures printed in 'CPK REPORT for ICT machine_MY47450178R' (Process Capability Report for I6 and for C3)."""
    i6 = capability_indices(0.000124192, 0.0001, 0.00015, within=2.76293e-08, overall=1.27074e-06)
    # The inputs are printed to six digits, so the last printed decimal of the largest figures can differ by one unit.
    assert (i6["ppl"], i6["ppu"], i6["cpl"], i6["cpu"]) == pytest.approx((6.35, 6.77, 291.86, 311.37), abs=0.011)
    assert round((0.00015 - 0.0001) / (6 * 1.27074e-06), 2) == 6.56 and round((0.00015 - 0.0001) / (6 * 2.76293e-08), 2) == 301.61
    c3 = capability_indices(7.99423e-11, -0.0001, 1e-10, within=3.63047e-13, overall=3.14699e-13)
    assert c3["ppu"] == pytest.approx(21.25, abs=0.01) and c3["cpu"] == pytest.approx(18.42, abs=0.01)
    assert c3["ppl"] == pytest.approx(105921193.36, rel=1e-5) and c3["cpl"] == pytest.approx(91815477.12, rel=1e-5)


def test_cpk_and_ppk_are_the_smaller_side_and_one_limit_gives_one_side():
    result = capability()
    indices = capability_indices(result.mean, result.lsl, result.usl, result.metrics["sigma_within"], result.metrics["sigma_overall"])
    assert min(indices["cpl"], indices["cpu"]) == pytest.approx(result.metrics["cpk"])
    assert min(indices["ppl"], indices["ppu"]) == pytest.approx(result.metrics["ppk"])
    one_sided = capability(low=None, high=6)
    assert one_sided.metrics["cp"] is None and one_sided.metrics["pp"] is None
    side = capability_indices(one_sided.mean, None, 6, one_sided.metrics["sigma_within"], one_sided.metrics["sigma_overall"])
    assert side["cpl"] is None and side["ppl"] is None and side["cpu"] == pytest.approx(one_sided.metrics["cpk"])


def test_ppm_observed_and_expected():
    values = [-2.0, -0.5, 0.0, 0.5, 2.0, 1.5, -1.5, 0.2, 0.3, 0.1]
    result = analyze([Observation("X", str(i), v, -1, 1, source=f"s{i}") for i, v in enumerate(values)], AnalysisConfig())[0]
    ppm = capability_ppm(result)
    assert ppm["observed"] == (pytest.approx(2e5), pytest.approx(2e5), pytest.approx(4e5))  # 2 of 10 below, 2 of 10 above
    lower = 1e6 * 0.5 * math.erfc((result.mean + 1) / (result.metrics["sigma_overall"] * math.sqrt(2)))
    assert ppm["overall"][0] == pytest.approx(lower) and ppm["overall"][2] == pytest.approx(ppm["overall"][0] + ppm["overall"][1])
    side = capability_ppm(capability(low=None, high=6))
    assert side["observed"][0] is None and side["overall"][0] is None and side["overall"][2] == side["overall"][1]


def test_capability_tables_are_the_ones_of_the_minitab_report():
    process, overall, within, performance = capability_tables(capability())
    assert [r[0] for r in process.rows] == ["LSL", "Target", "USL", "Sample Mean", "Sample N", "StDev(Overall)", "StDev(Within)"]
    assert [r[0] for r in overall.rows] == ["Pp", "PPL", "PPU", "Ppk", "Cpm"] and [r[0] for r in within.rows] == ["Cp", "CPL", "CPU", "Cpk"]
    assert performance.headers == ("", "Observed", "Expected Overall", "Expected Within")
    assert [r[0] for r in performance.rows] == ["PPM < LSL", "PPM > USL", "PPM Total"]
    assert dict(process.rows)["Target"] is None and dict(overall.rows)["Cpm"] is None  # printed as * without a target
    assert format_value(None, "idx") == "*" and format_value(7.99423e-11, "sig") == "7.99423E-11" and format_value(30, "sig") == "30"


def test_the_charts_are_in_the_order_of_minitabs_gage_report_and_use_the_analysis_estimates():
    from datamapperx.stat_reports import report_charts

    result = grr()
    assert [c.title for c in report_charts(result)] == ["Components of Variation", "R Chart", "Measurement by Sample ID", "Xbar Chart"]
    assert [c.title for c in report_charts(crossed_grr())][-2:] == ["Measurement by Operator", "Sample ID x Operator Interaction"]
    specs = charts_for_result(result)
    assert len(specs[3].series[0].points) == result.n
    assert specs[1].series[0].points == tuple((i + 1, part["range"]) for i, part in enumerate(result.details["part_summaries"]))
    cap = capability()
    histogram, = report_charts(cap)
    assert histogram.title == "" and not histogram.show_y and [s.label for s in histogram.series][1:] == ["Overall", "Within"]
    assert histogram.series[2].dash and [m[0] for m in histogram.markers] == ["LSL", "USL"]
    for spec in (*specs, *charts_for_result(cap)):
        for _, coordinates, _ in chart_primitives(spec):
            assert all(math.isfinite(value) for value in coordinates)


# ----------------------------------------------------------------------------- PDF
@pytest.mark.parametrize("factory", [grr, capability])
def test_pdf_has_the_results_and_none_of_the_internal_audit_content(tmp_path, factory):
    pytest.importorskip("reportlab")
    result = factory()
    original = copy.deepcopy(result)
    path = export_statistical_report(tmp_path / "analysis.pdf", [result], context=INFO)
    text = " ".join(text_of(path).split())
    assert result.test_item in text
    # the report is the Minitab report: no footer, no page numbers, nothing printed along the bottom of a page
    for footer in ("not an official", "Created", "DataMapperX", "Page 1"):
        assert footer not in text, footer
    for banned in ("PASS", "FAIL", "CAUTION", "Customer", "Decision", "Criteria", "Validation", "acknowledg", "Fingerprint", "stale",
                   "Methodology", "assumption", "Acceptance"):
        assert banned not in text, banned
    if result.method == "grr":
        for expected in ("Gage R&R for DCDC5W, Meas Vin", "One-Way ANOVA Table", "Variance Components", "%Contribution", "Gage Evaluation",
                         "%Tolerance", "Number of Distinct Categories", "Gage R&R (ANOVA) Report for", "Components of Variation", "R Chart",
                         "Xbar Chart"):
            assert expected in text, expected
    else:
        for expected in ("Process Capability Report for C3", "Process Data", "Overall Capability", "Potential (Within) Capability",
                         "Performance", "PPM < LSL", "PPM Total", "The actual process spread is represented by 6 sigma."):
            assert expected in text, expected
    assert result == original


def test_the_study_information_is_shown_only_when_entered_and_leaves_no_gap(tmp_path):
    pytest.importorskip("reportlab")
    result = grr()
    with_info = " ".join(text_of(export_statistical_report(tmp_path / "a.pdf", [result], context=INFO)).split())
    assert "Gage name: FIXTURE-1163" in with_info and "Date of study: 10-22-2026" in with_info and "Reported by: Thao Le" in with_info
    assert "Tolerance: 20" in with_info  # the process tolerance the percentages are computed against
    bare = " ".join(text_of(export_statistical_report(tmp_path / "b.pdf", [result])).split())
    assert "Gage name" not in bare and "Reported by" not in bare and "Misc" not in bare and "Tolerance: 20" in bare
    misc = ReportContext(info=(("Misc", "first article"),), included=1, total=1)
    assert "Misc: first article" in " ".join(text_of(export_statistical_report(tmp_path / "c.pdf", [result], context=misc)).split())


def test_a_partial_report_says_so_in_one_line_and_charts_can_be_left_out(tmp_path):
    pytest.importorskip("reportlab")
    pypdf = pytest.importorskip("pypdf")
    result = grr()
    partial = ReportContext(included=1, total=3, sections=())
    path = export_statistical_report(tmp_path / "partial.pdf", [result], context=partial)
    text = " ".join(text_of(path).split())
    assert "Included 1 of 3 test items." in text and "Gage R&R (ANOVA) Report" not in text
    assert len(pypdf.PdfReader(path).pages) == 1
    full = export_statistical_report(tmp_path / "full.pdf", [result], context=ReportContext(included=1, total=1))
    assert len(pypdf.PdfReader(full).pages) == 2 and "Included" not in text_of(full)


def test_each_test_item_has_its_own_pages_and_capability_is_landscape(tmp_path):
    pytest.importorskip("reportlab")
    pypdf = pytest.importorskip("pypdf")
    gage = export_statistical_report(tmp_path / "two.pdf", [grr("Item A"), grr("Item B")])
    pages = pypdf.PdfReader(gage).pages
    assert len(pages) == 4 and float(pages[0].mediabox.width) < float(pages[0].mediabox.height)  # portrait: tables, then graphs
    assert [("Gage R&R for Item A" in p.extract_text(), "Gage R&R for Item B" in p.extract_text()) for p in pages][::2] == [(True, False), (False, True)]
    cap = export_statistical_report(tmp_path / "cap.pdf", [capability("C3"), capability("C4")])
    pages = pypdf.PdfReader(cap).pages
    assert len(pages) == 2 and float(pages[0].mediabox.width) > float(pages[0].mediabox.height)
    assert "C3" in pages[0].extract_text() and "C4" in pages[1].extract_text() and "C4" not in pages[0].extract_text()


@pytest.mark.parametrize("count", [10, 30])
def test_all_four_grr_charts_and_every_full_sn_share_one_landscape_page(tmp_path, count):
    pypdf = pytest.importorskip("pypdf")
    result, serials = grr_with_serials(count)
    original = copy.deepcopy(result)
    path = export_statistical_report(tmp_path / f"all-{count}-serials.pdf", [result], context=INFO)
    pages = pypdf.PdfReader(path).pages
    assert len(pages) == 2
    assert float(pages[0].mediabox.width) < float(pages[0].mediabox.height)
    assert float(pages[1].mediabox.width) > float(pages[1].mediabox.height)
    text = pages[1].extract_text()
    for title in ("Components of Variation", "R Chart", "Measurement by Sample ID", "Xbar Chart"):
        assert title in text
    for serial in serials:
        assert text.count(serial) == 3, f"{serial} must label R, Xbar and Measurement by Sample ID"
    assert "FIXTURE-1163" in text and "Thao Le" in text and "10-22-2026" in text
    assert result == original


def test_crossed_report_retains_additional_graphs_after_the_four_chart_page(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    path = export_statistical_report(tmp_path / "crossed.pdf", [crossed_grr()], context=INFO)
    pages = pypdf.PdfReader(path).pages
    chart_pages = [page.extract_text() for page in pages if "Gage R&R (ANOVA) Report" in page.extract_text()]
    assert len(chart_pages) == 2
    assert all(title in chart_pages[0] for title in (
        "Components of Variation", "R Chart", "Measurement by Sample ID", "Xbar Chart"))
    assert "Measurement by Operator" in chart_pages[1] and "Sample ID x Operator Interaction" in chart_pages[1]


def test_long_notes_wrap_on_statistics_pages_and_do_not_split_the_four_charts(tmp_path):
    from dataclasses import replace

    pypdf = pytest.importorskip("pypdf")
    result, serials = grr_with_serials(30)
    notes = "Study notes " * 1000 + "END_OF_STUDY_NOTES"
    context = replace(INFO, info=(*INFO.info, ("Misc", notes)))
    path = export_statistical_report(tmp_path / "long-notes.pdf", [result], context=context)
    texts = [page.extract_text() for page in pypdf.PdfReader(path).pages]
    assert "END_OF_STUDY_NOTES" in "\n".join(texts)
    charts = [text for text in texts if "Components of Variation" in text]
    assert len(charts) == 1
    assert "See full notes on the statistics page." in charts[0]
    assert all(charts[0].count(serial) == 3 for serial in serials)
    assert all(title in charts[0] for title in ("R Chart", "Measurement by Sample ID", "Xbar Chart"))


def test_an_item_that_cannot_be_analysed_says_so_instead_of_printing_empty_tables(tmp_path):
    pytest.importorskip("reportlab")
    bad = analyze([Observation("Zero variance", str(i), 2, 0, 10) for i in range(3)], AnalysisConfig(method="grr"))[0]
    assert not bad.valid
    text = " ".join(text_of(export_statistical_report(tmp_path / "bad.pdf", [bad, grr()])).split())
    assert "Analysis not available" in text and "Gage R&R for Zero variance" in text
    assert session_text(bad).startswith("Zero variance") and "Analysis not available" in session_text(bad)


# ----------------------------------------------------------------------------- Excel
def test_excel_has_a_sheet_per_item_with_full_values_and_minitab_names(tmp_path):
    first, second = grr("Item A"), grr("Item B", high=30)
    path = export_statistical_report(tmp_path / "gage.xlsx", [first, second], context=INFO)
    book = load_workbook(path)
    assert book.sheetnames == ["Summary", "Item A", "Plot data 1", "Item B", "Plot data 2"]
    assert [c.value for c in book["Summary"][6]][:2] == ["Test Item", "%Contribution\n(Total Gage R&R)"]
    sheet = book["Item A"]
    cells = {row[0].value: row for row in sheet.iter_rows(min_col=1, max_col=6) if row[0].value}
    assert {"One-Way ANOVA Table", "Variance Components", "Gage Evaluation", "Gage name:", "Tolerance:"} <= set(cells)
    assert sheet["A1"].value == "Gage R&R for Item A" and len(sheet._charts) == 4
    total = next(row[1].value for row in sheet.iter_rows(min_col=1, max_col=2) if str(row[0].value).strip() == "Total Gage R&R")
    assert total == pytest.approx(first.metrics["variance_grr"], rel=1e-12)  # the whole value, only the display is formatted
    assert any(cell.number_format.startswith("0.0000000") for row in sheet.iter_rows() for cell in row if isinstance(cell.value, float))
    assert not any("PASS" in str(c.value) or "FAIL" in str(c.value) for ws in book for row in ws.iter_rows() for c in row)
    book.close()


def test_a_single_item_workbook_has_no_summary_and_capability_has_the_minitab_tables(tmp_path):
    result = capability("C3")
    book = load_workbook(export_statistical_report(tmp_path / "cap.xlsx", [result]))
    assert book.sheetnames == ["C3", "Plot data 1"]
    sheet = book["C3"]
    labels = [row[0].value for row in sheet.iter_rows(min_col=1, max_col=1) if row[0].value]
    for expected in ("Process Capability Report for C3", "Process Data", "Overall Capability", "Potential (Within) Capability", "Performance",
                     "PPM < LSL", "The actual process spread is represented by 6 sigma."):
        assert expected in labels, expected
    assert len(sheet._charts) == 1
    cpk = next(row[1].value for row in sheet.iter_rows(min_col=1, max_col=2) if row[0].value == "Cpk")
    assert cpk == pytest.approx(result.metrics["cpk"], rel=1e-12)
    book.close()


def test_excel_serial_axes_request_every_label_and_diagonal_full_text(tmp_path):
    from xml.etree import ElementTree
    from zipfile import ZipFile

    from openpyxl.utils import range_boundaries

    result, serials = grr_with_serials(30)
    path = export_statistical_report(tmp_path / "all-serials.xlsx", [result])
    workbook = load_workbook(path)
    ws = workbook[workbook.sheetnames[0]]
    data = workbook["Plot data 1"]
    assert len(ws._charts) == 4
    for chart in ws._charts[1:]:
        assert chart.x_axis.tickLblSkip == chart.x_axis.tickMarkSkip == 1
        assert chart.x_axis.txPr.bodyPr.rot == -2700000
        category_range = chart.series[0].cat.numRef.f.split("!")[-1]
        column, first, _, last = range_boundaries(category_range)
        assert tuple(data.cell(row, column).value for row in range(first, last + 1)) == serials
    workbook.close()
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(path) as archive:
        for number in (2, 3, 4):
            xml = ElementTree.fromstring(archive.read(f"xl/charts/chart{number}.xml"))
            assert xml.find(".//c:catAx/c:tickLblSkip", ns).get("val") == "1"
            assert xml.find(".//c:catAx/c:txPr/a:bodyPr", ns).get("rot") == "-2700000"


def test_the_observations_sheet_is_optional_and_charts_can_be_left_out(tmp_path):
    result = grr()
    chosen = ReportContext(included=1, total=1, sections=("observations",))
    book = load_workbook(export_statistical_report(tmp_path / "obs.xlsx", [result], context=chosen))
    assert book.sheetnames[-1] == "Observations" and "Plot data 1" not in book.sheetnames and not book[book.sheetnames[0]]._charts
    assert book["Observations"].max_row == len(result.observations) + 1
    plain = load_workbook(export_statistical_report(tmp_path / "plain.xlsx", [result]))
    assert "Observations" not in plain.sheetnames
    book.close()
    plain.close()


def test_excel_never_executes_log_strings_and_keeps_item_names_as_text(tmp_path):
    result = capability("=SUM(1,2)")
    book = load_workbook(export_statistical_report(tmp_path / "safe.xlsx", [result], context=ReportContext(
        info=(("Gage name", "=HYPERLINK(1)"),), included=1, total=1, sections=("charts", "observations"))))
    sheet = book[book.sheetnames[0]]
    assert sheet["A1"].value == "Process Capability Report for =SUM(1,2)" and sheet["B2"].value == "=HYPERLINK(1)"
    assert sheet["B2"].data_type == "s" and book["Observations"]["A2"].data_type == "s"
    assert not any(cell.data_type == "f" for ws in book for row in ws for cell in row)
    book.close()


def test_sheet_names_are_valid_and_unique(tmp_path):
    names = ["A/B: C?" * 6, "A/B: C?" * 6]
    results = [capability(names[0]), capability(names[1])]
    book = load_workbook(export_statistical_report(tmp_path / "names.xlsx", results))
    sheets = [n for n in book.sheetnames if n not in ("Summary",) and not n.startswith("Plot data")]
    assert len(sheets) == 2 and len(set(s.casefold() for s in sheets)) == 2 and all(len(s) <= 31 and not set("/:?") & set(s) for s in sheets)
    book.close()


def test_an_item_with_a_tester_or_slot_is_named_without_the_operator_list():
    result = crossed_grr()
    assert "Operators=" in result.system and item_label(result) == "Crossed item"


# ----------------------------------------------------------------------------- numbers are written out, as in Minitab's reports
SCIENTIFIC = re.compile(r"\d[eE][-+]?\d")  # 1.84897E-06, 5e-08


def tiny_grr():
    """Readings of a few millivolts that differ by nanovolts: every variance is far below 1e-5, the tolerance is 0.00005."""
    rows = [Observation("Tiny", f"S{part}", 0.0068 + part * 1e-8 + delta, 0.00678, 0.00683, trial=str(trial))
            for part in range(1, 11) for trial, delta in enumerate((-1e-9, 4e-10, 1.3e-9), 1)]
    result, = analyze(rows, AnalysisConfig(method="grr"))
    assert result.valid and 0 < result.metrics["variance_grr"] < 1e-12
    return result


def test_tables_are_in_fixed_decimals_however_small_the_numbers():
    assert format_value(1.84897e-06, "var") == "0.0000018" and format_value(0.0, "var") == "0.0000000"
    assert format_value(1.84897e-06, "var", 4) == "0.000001849"  # the report precision still sets the digits
    assert format_value(7.99423e-11, "sig") == "7.99423E-11"  # the capability Process Data keeps Minitab's own notation
    text = session_text(tiny_grr())
    assert not SCIENTIFIC.search(text) and "0.0000000" in text
    assert not SCIENTIFIC.search(session_text(tiny_grr(), 4))


def test_pdf_and_excel_of_tiny_numbers_have_no_scientific_notation_and_keep_the_full_values(tmp_path):
    pytest.importorskip("reportlab")
    result = tiny_grr()
    pdf = text_of(export_statistical_report(tmp_path / "tiny.pdf", [result], context=INFO))
    assert not SCIENTIFIC.search(pdf.replace("FIXTURE-1163", "")) and "Tolerance: 0.00005" in " ".join(pdf.split())
    sheet = load_workbook(export_statistical_report(tmp_path / "tiny.xlsx", [result], context=INFO))[result.test_item]
    floats = [cell for row in sheet.iter_rows() for cell in row if isinstance(cell.value, float)]
    assert floats and not any("E" in cell.number_format.upper() for cell in floats)
    total = next(row[1].value for row in sheet.iter_rows(min_col=1, max_col=2) if str(row[0].value).strip() == "Total Gage R&R")
    assert total == pytest.approx(result.metrics["variance_grr"], rel=1e-12)  # the cell holds the whole value; only its display is fixed


def limit_labels(title, lower, centre, upper, width=650, height=280):
    series = tuple(ChartSeries(name, ((1, value), (10, value))) for name, value in (("UCL", upper), ("CENTER", centre), ("LCL", lower)))
    spec = ChartSpec(title, "Part", "Sample Range", series)
    return [style["text"] for kind, _, style in chart_primitives(spec, width, height) if kind == "text" and "=" in style["text"]]


def test_control_limits_are_written_with_six_decimals_like_minitabs_charts():
    assert limit_labels("R Chart", 0, 0.001250, 0.003217) == ["UCL=0.003217", "Rbar=0.001250", "LCL=0"]
    assert limit_labels("Xbar Chart", 1.840833, 1.842112, 1.843390) == ["UCL=1.843390", "Mean=1.842112", "LCL=1.840833"]


def test_control_limits_that_are_tiny_or_close_together_get_more_decimals_instead_of_scientific_notation():
    assert limit_labels("R Chart", 0, 2.3e-08, 5.92077e-08) == ["UCL=0.0000000592", "Rbar=0.0000000230", "LCL=0"]
    close = limit_labels("Xbar Chart", 0.00684498, 0.006845, 0.00684503)
    assert len(set(close)) == 3 and not any(SCIENTIFIC.search(label) for label in close)


def test_charts_of_tiny_numbers_have_no_scientific_notation_and_the_scale_labels_fit():
    from datamapperx.stat_charts import _text_units

    for result in (tiny_grr(), capability(low=0, high=6)):
        for spec in charts_for_result(result):
            for _, (x, _), style in (p for p in chart_primitives(spec, 375, 218) if p[0] == "text"):
                assert not SCIENTIFIC.search(style["text"]), style["text"]
                if style["anchor"] == "e" and not style.get("category"):  # a label of the vertical scale
                    assert 8 * _text_units(style["text"]) <= x, style["text"]
    tiny = analyze([Observation("C3", str(i), 8e-11 + (i % 7) * 1e-13, -1e-4, 1e-10, source=f"log{i}.txt") for i in range(30)], AnalysisConfig())[0]
    texts = [style["text"] for kind, _, style in chart_primitives(charts_for_result(tiny)[0]) if kind == "text"]
    assert not any(SCIENTIFIC.search(text) for text in texts)
    ticks = [text for text in texts if re.fullmatch(r"\d+\.\d+", text)]  # the scale of the histogram, around 6E-11
    assert len(ticks) >= 2 and len(ticks) == len(set(ticks))  # neighbouring ticks are told apart, not rounded into one


# ----------------------------------------------------------------------------- the application's own tables and plumbing
def test_summary_rows_for_the_application_table():
    headers, rows = summary_rows([grr()], judged=False)
    assert headers == ("Test Item", "System", "%SV/Tol", "%Study Var", "NDC") and len(rows) == 1
    with pytest.raises(ValueError, match="one analysis method"):
        summary_rows([capability(), grr()])


def test_tk_uses_portable_chart_geometry():
    calls = []

    class Canvas:
        def winfo_width(self):
            return 600

        def winfo_height(self):
            return 280

        def __getattr__(self, name):
            return lambda *args, **kwargs: calls.append((name, args, kwargs))

    draw_tk_chart(Canvas(), charts_for_result(capability())[0])
    assert calls[0][:2] == ("delete", ("all",))
    assert any(name == "create_rectangle" for name, *_ in calls)
    assert any(name == "create_line" and kwargs.get("dash") for name, _, kwargs in calls)  # the Within curve and the limits are dashed
    assert all(isinstance(kwargs["font"][1], int) for name, _, kwargs in calls if name == "create_text")


def test_invalid_dataset_is_not_plotted_as_a_silently_reduced_subset():
    result, = analyze([Observation("X", "1", "bad", 0, 6), Observation("X", "2", 3, 0, 6)])
    chart, = charts_for_result(result)
    assert not chart.series and "Analysis unavailable" in chart.note


def test_export_failure_preserves_previous_file(tmp_path, monkeypatch):
    import datamapperx.stat_reports as reports

    destination = tmp_path / "existing.pdf"
    destination.write_bytes(b"original")

    def fail(*args):
        raise RuntimeError("Report failed")

    monkeypatch.setattr(reports, "_export_pdf", fail)
    with pytest.raises(RuntimeError, match="Report failed"):
        export_statistical_report(destination, [capability()])
    assert destination.read_bytes() == b"original"
    assert not list(tmp_path.glob(".datamapperx-stat-*"))


def test_export_rejects_empty_mixed_unsupported_and_invalid_precision(tmp_path):
    with pytest.raises(ValueError, match="Run an analysis"):
        export_statistical_report(tmp_path / "empty.xlsx", [])
    with pytest.raises(ValueError, match="one analysis method"):
        export_statistical_report(tmp_path / "mixed.xlsx", [capability(), grr()])
    with pytest.raises(ValueError, match="support .pdf and .xlsx"):
        export_statistical_report(tmp_path / "bad.csv", [capability()])
    with pytest.raises(ValueError, match="precision"):
        export_statistical_report(tmp_path / "bad.xlsx", [capability()], decimals=-1)
    assert display_number(None) == display_number(float("inf")) == "N/A"


def test_packaged_smoke_check_exercises_both_methods_and_report_formats(tmp_path):
    from datamapperx.selftest import _check_statistical_reports

    folder = tmp_path / "reports"
    _check_statistical_reports(folder)
    assert {path.name for path in folder.iterdir()} == {"capability.pdf", "capability.xlsx", "grr.pdf", "grr.xlsx"}


def test_packaged_smoke_check_fails_when_pdf_dependency_is_missing(tmp_path, monkeypatch):
    import datamapperx.stat_reports as reports
    from datamapperx.selftest import _check_statistical_reports

    def missing_dependency(*args):
        raise ModuleNotFoundError("No module named 'reportlab'", name="reportlab")

    monkeypatch.setattr(reports, "_export_pdf", missing_dependency)
    with pytest.raises(ModuleNotFoundError, match="reportlab"):
        _check_statistical_reports(tmp_path / "reports")
